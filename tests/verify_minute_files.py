# -*- coding: utf-8 -*-
"""Minute/hour-level scheduling precision: file round-tripping. Excel export writes a real date+time serial (a fractional
day) for a minute-mode leaf task's own Start/Finish/Actual/Baseline cells, formatted "dd.mm.yyyy hh:mm"; duration and
variance cells switch to pre-formatted text ("4 hrs", "+30 min") since Excel's number-format codes have no clean way to
render a raw minute count the way the day-mode "day"/"days" trick does. MS Project XML (MSPDI) export writes each task's
REAL startTime/endTime/actualStartTime/actualFinishTime/constraintTime instead of the fixed 08:00/17:00-style literals it
always used, and a minute-precise <Duration> (PT{h}H{m}M0S); import goes the other way and also RE-ACTIVATES minute mode
on the plan when the file actually looks minute-precise (MinutesPerDay != 480, or some task's own time isn't exactly the
calendar's own default open/close) — an ordinary MS Project file, where every task simply starts/ends at the calendar's
own boundary, still imports as an untouched day-mode plan, byte-for-byte as before this existed."""
import os, re
from playwright.sync_api import sync_playwright
import openpyxl
URL = os.environ.get("MILESTONE_URL", "http://127.0.0.1:8937/milestone.html")
errors, results = [], []
def check(name, cond, detail=""):
    results.append(bool(cond)); print(("PASS  " if cond else "FAIL  ") + name + (f"   [{str(detail)[:400]}]" if not cond and detail else ""))

with sync_playwright() as p:
    b = p.chromium.launch(headless=True)
    ctx = b.new_context(viewport={"width": 1300, "height": 800}, accept_downloads=True); ctx.add_init_script("delete window.showOpenFilePicker; delete window.showSaveFilePicker; delete window.showDirectoryPicker")
    pg = ctx.new_page(); pg.on("pageerror", lambda e: errors.append(str(e))); pg.on("console", lambda m: errors.append(m.text) if m.type in ("error", "warning") else None)
    pg.goto(URL); pg.wait_for_selector("#addTaskBtn"); pg.evaluate("() => localStorage.clear()"); pg.reload(); pg.wait_for_selector("#addTaskBtn")
    ev = pg.evaluate

    # ---------------------------------------------------------------- Excel export: date+time cells, minute-unit duration/variance/baseline text
    ev("""() => {
        project.timeUnit = 'minute';
        const t = { id: genId(), name: 'A', parentId: null, order: 0, startDate: '2026-09-07', endDate: '2026-09-07', startTime: '09:00', endTime: '11:00',
          progress: 0, milestone: false, color: null, predecessors: [], collapsed: false, updatedAt: 1, constraintType: 'ASAP', constraintDate: null,
          taskMode: 'auto', resource: '', actualStart: null, actualFinish: null };
        tasks.length = 0; tasks.push(t); normalizeData(); save();
        applyBaselineChange(0, 'all', false);
        t.startTime = '09:30'; t.endTime = '11:30'; normalizeData(); save(); render();
        for (const c of ['baselineStart','baselineFinish','baselineDuration','startVariance','finishVariance','durationVariance']) colHidden.delete(c);
        render();
    }""")
    pg.click("#dataMenuBtn"); pg.click("#excelExportItem"); pg.wait_for_selector("#excelModalBg.open")
    with pg.expect_download() as d: pg.click("#excelExportBtn")
    d.value.save_as("minute_export.xlsx")
    wb = openpyxl.load_workbook("minute_export.xlsx"); ws = wb["Tasks"]
    header = {c.value: c.column for c in ws[4]}
    row = 5
    def cellat(name): return ws.cell(row=row, column=header[name])
    check("Excel Start is a real datetime (09:30), formatted with hh:mm", cellat("Start").value.hour == 9 and cellat("Start").value.minute == 30 and "hh:mm" in cellat("Start").number_format, (cellat("Start").value, cellat("Start").number_format))
    check("Excel Finish is a real datetime (11:30)", cellat("Finish").value.hour == 11 and cellat("Finish").value.minute == 30)
    check("Excel Duration is pre-formatted minute-unit text ('2 hrs'), not a day-formatted number", cellat("Duration").value == "2 hrs", cellat("Duration").value)
    check("Excel Baseline Start/Finish keep the baseline's OWN time (09:00/11:00), not the moved task's", cellat("Baseline Start").value.hour == 9 and cellat("Baseline Start").value.minute == 0 and cellat("Baseline Finish").value.hour == 11 and cellat("Baseline Finish").value.minute == 0)
    check("Excel Baseline Duration is minute-unit text ('2 hrs')", cellat("Baseline Duration").value == "2 hrs", cellat("Baseline Duration").value)
    check("Excel Start/Finish Variance are minute-unit text ('+0.5 hrs' each — both ends moved by the same 30 minutes)", cellat("Start Variance").value == "+0.5 hrs" and cellat("Finish Variance").value == "+0.5 hrs", (cellat("Start Variance").value, cellat("Finish Variance").value))

    # ---------------------------------------------------------------- Excel export stays exactly day-mode for a day-mode plan (regression)
    ev("""() => {
        delete project.timeUnit; delete project.workHours;
        const t = tasks[0]; t.startTime = t.endTime = null; t.baselines[0] = [t.baselines[0][0], t.baselines[0][1]];
        normalizeData(); save(); render();
    }""")
    pg.click("#dataMenuBtn"); pg.click("#excelExportItem"); pg.wait_for_selector("#excelModalBg.open")
    with pg.expect_download() as d: pg.click("#excelExportBtn")
    d.value.save_as("day_export.xlsx")
    wb2 = openpyxl.load_workbook("day_export.xlsx"); ws2 = wb2["Tasks"]
    header2 = {c.value: c.column for c in ws2[4]}
    startCell2 = ws2.cell(row=5, column=header2["Start"])
    check("a day-mode plan's Excel export keeps the plain dd.mm.yyyy format, no time shown", "hh:mm" not in startCell2.number_format and startCell2.value.hour == 0 and startCell2.value.minute == 0, startCell2.number_format)
    durCell2 = ws2.cell(row=5, column=header2["Duration"])
    check("...and Duration is still a real number with the day-based format", isinstance(durCell2.value, (int, float)) and "day" in durCell2.number_format, (durCell2.value, durCell2.number_format))

    # ---------------------------------------------------------------- MSPDI export: real per-task times, minute-precise duration, baseline times
    # A fresh task + baseline of its own — not the one the Excel section above already stripped down to a plain
    # [start, finish] baseline (day-mode regression step).
    xml = ev("""() => {
        tasks.length = 0; delete project.baselines; delete project.compareBaseline;
        project.timeUnit = 'minute'; project.workHours = { start: '09:00', end: '18:00', breaks: [{start:'12:30', end:'13:15'}] };
        const t = { id: genId(), name: 'A', parentId: null, order: 0, startDate: '2026-09-07', endDate: '2026-09-07', startTime: '09:00', endTime: '11:00',
          progress: 0, milestone: false, color: null, predecessors: [], collapsed: false, updatedAt: 1, constraintType: 'ASAP', constraintDate: null,
          taskMode: 'auto', resource: '', actualStart: null, actualFinish: null };
        tasks.push(t); normalizeData(); save();
        applyBaselineChange(0, 'all', false);
        t.startTime = '10:00'; t.endTime = '12:00'; normalizeData(); save();
        return buildMspdi();
    }""")
    check("MSPDI Start carries the task's real time (10:00), not a fixed calendar default", "<Start>2026-09-07T10:00:00</Start>" in xml, xml[:300])
    check("...Finish too (12:00)", "<Finish>2026-09-07T12:00:00</Finish>" in xml)
    check("...Duration is minute-precise (PT2H0M0S for a 2h task), not the day-based PT*8H formula", "<Duration>PT2H0M0S</Duration>" in xml)
    check("...the Baseline block keeps the BASELINE's own time (09:00/11:00), separate from the task's current time", "<Start>2026-09-07T09:00:00</Start>" in xml and "<Finish>2026-09-07T11:00:00</Finish>" in xml)
    check("...DefaultStartTime/DefaultFinishTime/MinutesPerDay reflect the real custom working hours (540 min day minus a 45-min break = 495)", "<DefaultStartTime>09:00:00</DefaultStartTime>" in xml and "<MinutesPerDay>495</MinutesPerDay>" in xml, re.search(r"<MinutesPerDay>\d+</MinutesPerDay>", xml))

    # ---------------------------------------------------------------- MSPDI import: re-activates minute mode when the file is genuinely minute-precise
    r = ev("(xml) => { const parsed = parseMspdi(xml); return { project: parsed.project, task: { startTime: parsed.tasks[0].startTime, endTime: parsed.tasks[0].endTime, baselines: parsed.tasks[0].baselines }, warnings: parsed.warnings }; }", xml)
    check("parseMspdi() detects the real times and switches the parsed project to minute mode", r["project"].get("timeUnit") == "minute", r["project"])
    check("...and recovers the real working hours (start/end/break) from the calendar's own <WorkingTimes>, not just the defaults", r["project"].get("workHours") == {"start": "09:00", "end": "18:00", "breaks": [{"start": "12:30", "end": "13:15"}]}, r["project"].get("workHours"))
    check("...the task's own time round-trips exactly", r["task"]["startTime"] == "10:00" and r["task"]["endTime"] == "12:00", r["task"])
    check("...the baseline's time round-trips too, separately from the task's current time", r["task"]["baselines"]["0"] == ["2026-09-07", "2026-09-07", "09:00", "11:00"], r["task"]["baselines"])
    check("...and a warning explains the switch", any("Hours & minutes" in w for w in r["warnings"]), r["warnings"])

    # ---------------------------------------------------------------- an ORDINARY MS Project file (every time at the calendar's own boundary) stays day-mode
    xml_day = ev("""() => {
        delete project.timeUnit; delete project.workHours;
        const t = tasks[0]; t.startTime = t.endTime = null; normalizeData(); save();
        return buildMspdi();
    }""")
    r2 = ev("(xml) => { const parsed = parseMspdi(xml); return { timeUnit: parsed.project.timeUnit, startTime: parsed.tasks[0].startTime, warnings: parsed.warnings }; }", xml_day)
    check("an ordinary file (every task at the calendar's default 08:00/17:00) is NOT treated as minute-precise", r2["timeUnit"] is None and not r2["warnings"], r2)

    # ---------------------------------------------------------------- the real import UI: "import calendar too" gates whether minute mode is actually applied
    pg.evaluate("() => localStorage.clear()"); pg.reload(); pg.wait_for_selector("#addTaskBtn")
    pg.evaluate("(xml) => openForeignImport(xml, 'test.xml')", xml)
    pg.wait_for_selector("#foreignImportModalBg.open")
    check("the calendar checkbox is offered and checked by default", pg.is_visible("#foreignImportCalRow") and pg.is_checked("#foreignImportCal"))
    pg.click("#foreignImportGo"); pg.wait_for_timeout(200)
    r3 = ev("() => ({ timeUnit: project.timeUnit, workHours: project.workHours, task: { startTime: tasks[0].startTime, endTime: tasks[0].endTime } })")
    check("importing with the calendar checked actually switches the live plan to minute mode", r3["timeUnit"] == "minute" and r3["workHours"]["start"] == "09:00" and r3["task"]["startTime"] == "10:00", r3)

    pg.evaluate("() => localStorage.clear()"); pg.reload(); pg.wait_for_selector("#addTaskBtn")
    pg.evaluate("(xml) => openForeignImport(xml, 'test.xml')", xml)
    pg.wait_for_selector("#foreignImportModalBg.open")
    pg.uncheck("#foreignImportCal")
    pg.click("#foreignImportGo"); pg.wait_for_timeout(200)
    r4 = ev("() => ({ timeUnit: project.timeUnit, task: { startTime: tasks[0].startTime, endTime: tasks[0].endTime } })")
    check("...unchecking it leaves the plan in day mode and drops the speculative per-task times", r4["timeUnit"] is None and r4["task"]["startTime"] is None and r4["task"]["endTime"] is None, r4)

    check("no console errors or page errors across the whole run", not errors, errors[:5])
    n_ok, n_all = sum(results), len(results)
    print(f"\n{n_ok}/{n_all} checks passed")
    b.close()
    raise SystemExit(0 if n_ok == n_all else 1)
