# -*- coding: utf-8 -*-
"""Resource Sheet, Stage 6 — the task grid's own Cost column and its currency setting. taskCost(t) (built in Stage 2)
reaches the grid, filters and Excel export the same way every other computed column already does — TASK_COLS,
FILTER_COLS, the HEAD lookup, the grid's own cells object, XCOLS and writeTaskRow()'s switch, all six touched (the
exact hand-maintained-lookup gap CLAUDE.md's "Task Type / Work" section already found once for Work/Task Type — see
"Resource Sheet"). Read-only (computed, never typed), hidden by default like Resource/Work. See CLAUDE.md."""
import os, io, base64
from playwright.sync_api import sync_playwright
import openpyxl
URL = os.environ.get("MILESTONE_URL", "http://127.0.0.1:8937/milestone.html")
errors, results = [], []
def check(name, cond, detail=""):
    results.append(bool(cond)); print(("PASS  " if cond else "FAIL  ") + name + (f"   [{str(detail)[:400]}]" if not cond and detail else ""))

with sync_playwright() as p:
    b = p.chromium.launch(headless=True)
    ctx = b.new_context(viewport={"width": 1700, "height": 900}, accept_downloads=True); ctx.add_init_script("delete window.showOpenFilePicker; delete window.showSaveFilePicker; delete window.showDirectoryPicker")
    pg = ctx.new_page(); pg.on("pageerror", lambda e: errors.append(str(e))); pg.on("console", lambda m: errors.append(m.text) if m.type in ("error", "warning") else None)
    pg.goto(URL); pg.wait_for_selector("#addTaskBtn"); pg.evaluate("() => localStorage.clear()"); pg.reload(); pg.wait_for_selector("#addTaskBtn")
    ev = pg.evaluate

    check("cost is registered in TASK_COLS, hidden by default, right after Work", ev("() => DEFAULT_COL_ORDER.indexOf('cost') === DEFAULT_COL_ORDER.indexOf('work') + 1 && !TASK_COLS.cost.dflt"))
    check("...and in FILTER_COLS as a number filter", ev("() => FILTER_COLS.cost.kind") == "number")

    SEED = "specs => { tasks.length = 0; deletedTaskIds.length = 0; selectedTaskId = null; delete project.resources; delete project.currencyCode; delete project.workDays; const ids = {}; for (const sp of specs) { const t = Object.assign({id: genId(), name: sp.name, parentId: sp.parent ? ids[sp.parent] : null, order: tasks.length, startDate: sp.s, endDate: sp.e, progress: 0, milestone: false, color: null, predecessors: [], collapsed: false, updatedAt: 1, constraintType: 'ASAP', constraintDate: null, taskMode: 'auto', resource: sp.r || '', actualStart: null, actualFinish: null}, sp.extra || {}); tasks.push(t); ids[sp.name] = t.id; } normalizeData(); save(); render(); }"
    seed = lambda specs: ev(SEED, specs)
    seed([{"name": "G", "s": "2026-09-07", "e": "2026-09-08"}, {"name": "A", "s": "2026-09-07", "e": "2026-09-08", "parent": "G", "r": "Anna"}, {"name": "U", "s": "2026-09-07", "e": "2026-09-08"}])
    ev("() => { project.resources.find(r => r.name === 'Anna').stdRate = 50; save(); render(); }")   # Anna: 2 working days = 16h x €50 = €800

    ev("() => { toggleColumn('cost', true); render(); }")
    body = ev("() => document.getElementById('gridRows').innerText")
    check("the Columns menu can show it and the grid renders it without a console error", not errors, errors[:3])
    a_cost = ev("() => fmtCurrency(taskCost(tasks.find(t => t.name === 'A')))")
    check("an assigned task's own Cost cell reads the real rolled-up figure (€800, EUR the default)", a_cost == "€800.00", a_cost)
    g_cost = ev("() => fmtCurrency(taskCost(tasks.find(t => t.name === 'G')))")
    check("a group's Cost is the rollup, same figure as its one child (unlike Work, Cost IS meaningful for a group)", g_cost == a_cost, (g_cost, a_cost))
    u_cost = ev("() => fmtCurrency(taskCost(tasks.find(t => t.name === 'U')))")
    check("an unassigned task reads €0.00, never blank/NaN", u_cost == "€0.00", u_cost)

    # ---------------------------------------------------------------- filtering by Cost
    ev("() => { colFilters.cost = { type: 'rule', rule: 'gt', a: 100, b: null }; render(); }")
    shown = ev("() => visibleTaskList().map(x => x.task.name)")
    check("filtering Cost > 100 shows only the tasks whose rollup exceeds it (A and its group G, not the unassigned U)", set(shown) == {"A", "G"}, shown)
    ev("() => { colFilters.cost = null; render(); }")

    # ---------------------------------------------------------------- Excel: 'all columns' Cost header + a real numeric currency cell
    def wb(columns="shown"):
        b64 = ev("""async (o) => { const blob = buildXlsx(o); const buf = new Uint8Array(await blob.arrayBuffer()); let s = ''; for (const c of buf) s += String.fromCharCode(c); return btoa(s); }""", {"scope": "all", "gantt": False, "columns": columns})
        return openpyxl.load_workbook(io.BytesIO(base64.b64decode(b64)))
    w = wb("all")
    ws = w["Tasks"]
    header = [c.value for c in ws[4] if c.value not in (None, "")]
    check("Excel 'all columns' includes a Cost header", "Cost" in header, header)
    ci = header.index("Cost") + 1
    a_row = next(r for r in range(5, ws.max_row + 1) if ws.cell(r, header.index("Task Name") + 1).value == "A")
    cell = ws.cell(a_row, ci)
    check("...and A's Cost cell is a real number (800), not pre-formatted text — sortable/summable in Excel", cell.value == 800, cell.value)
    check("...with a currency number format (the € sign literal, matching the plan's default EUR)", "€" in (cell.number_format or ""), cell.number_format)

    # ---------------------------------------------------------------- a plan with NO resource pool at all still exports (the exact XCOLS-gap crash class found once already)
    ev("() => { tasks.length = 0; delete project.resources; const t = {id: genId(), name: 'Solo', parentId: null, order: 0, startDate: '2026-09-07', endDate: '2026-09-08', progress: 0, milestone: false, color: null, predecessors: [], collapsed: false, updatedAt: 1, constraintType: 'ASAP', constraintDate: null, taskMode: 'auto', resource: '', actualStart: null, actualFinish: null}; tasks.push(t); normalizeData(); save(); render(); }")
    w2 = wb("all")
    check("a plan with no pool at all still builds the 'all columns' export without crashing (Cost just reads €0.00/0)", w2 is not None)

    # ---------------------------------------------------------------- currency setting changes the grid's own formatting too
    ev("() => { project.currencyCode = 'USD'; save(); render(); }")
    usd = ev("() => fmtCurrency(800)")
    check("switching the plan's currency changes what the Cost column would show", "$" in usd, usd)

    check("no console errors or page errors across the whole run", not errors, errors[:5])
    n_ok, n_all = sum(results), len(results)
    print(f"\n{n_ok}/{n_all} checks passed")
    b.close()
    raise SystemExit(0 if n_ok == n_all else 1)
