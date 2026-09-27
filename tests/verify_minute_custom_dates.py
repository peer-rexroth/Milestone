# -*- coding: utf-8 -*-
"""Minute-mode scheduling precision: custom DATE columns carry a time of day too. A date field's time is stored beside it
(task.custom.<id>Time, e.g. date1Time) — only in a minute-mode plan (a day-mode plan drops it, like every task time). The list
edits it in the typed "dd.mm.yyyy hh:mm" box (a date alone clears the time, a time alone keeps the date, a blank clears both), the
column widens like the task date columns, the task dialog's Custom tab gets a time input beside the date, Excel writes a real
date+time cell (date-only when there is no time), copied rows and the MS Project notes carry it, the Edit-tasks dialog sets a
plain date (clearing the time), and the Custom tab's badge still counts fields, not their time inputs."""
import os, re, tempfile
from playwright.sync_api import sync_playwright
import openpyxl
URL = os.environ.get("MILESTONE_URL", "http://127.0.0.1:8937/milestone.html")
errors, results = [], []
def check(name, cond, detail=""):
    results.append(bool(cond)); print(("PASS  " if cond else "FAIL  ") + name + (f"   [{str(detail)[:400]}]" if not cond and detail else ""))
SEED = re.search(r'SEED = """(.*?)"""', open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'verify_clone.py')).read(), re.S).group(1)
TMP = tempfile.mkdtemp()

with sync_playwright() as p:
    b = p.chromium.launch(headless=True)
    ctx = b.new_context(viewport={"width": 1700, "height": 900}, accept_downloads=True); ctx.add_init_script("delete window.showOpenFilePicker; delete window.showSaveFilePicker; delete window.showDirectoryPicker")
    pg = ctx.new_page(); pg.on("pageerror", lambda e: errors.append(str(e))); pg.on("console", lambda m: errors.append(m.text) if m.type in ("error", "warning") else None)
    pg.goto(URL); pg.wait_for_selector("#addTaskBtn"); pg.evaluate("() => localStorage.clear()"); pg.reload(); pg.wait_for_selector("#addTaskBtn")
    ev = pg.evaluate
    tid = lambda n: ev("n => tasks.find(t => t.name === n).id", n)
    cust = lambda n: ev("n => { const c = tasks.find(t => t.name === n).custom || {}; return [c.date1 || null, c.date1Time || null]; }", n)
    toast = lambda: ev("() => document.getElementById('toastMsg').textContent")
    def plan(minute=True, value=None):
        ev("m => { if (m) { project.timeUnit = 'minute'; normalizeData(); } else { delete project.timeUnit; delete project.workHours; delete project.lagUnit; } }", minute)
        ev(SEED, [{"name": "A", "s": "2026-09-07", "e": "2026-09-07"}, {"name": "B", "s": "2026-09-08", "e": "2026-09-08"}])
        ev("v => { project.fieldNames = { date1: 'Review' }; for (const t of tasks) { if (v && t.name === 'A') t.custom = v; } colHidden.delete('date1'); normalizeData(); save(); render(); }", value)
    def cell(name):
        idx = ev("() => [...document.querySelectorAll('#gridHeader .col-filter-btn')].map(b => b.dataset.col).indexOf('date1')")
        return pg.locator(f".grid-row[data-id='{tid(name)}']").locator(":scope > div").nth(idx + 1)
    def edit(name, text):
        cell(name).click(); pg.wait_for_selector(".inline-edit"); pg.fill(".inline-edit", text); pg.keyboard.press("Enter"); pg.wait_for_timeout(150)

    # ---------------------------------------------------------------- storage
    plan(True, {"date1": "2026-09-07", "date1Time": "14:30"})
    check("a minute-mode plan keeps a date field's time (date1Time) and the grid shows date and time", cust("A") == ["2026-09-07", "14:30"] and cell("A").inner_text() == "07.09.2026 14:30", (cust("A"), cell("A").inner_text()))
    ev("() => { tasks.find(t => t.name === 'B').custom = { date1Time: '10:00', date1: 'garbage' }; normalizeData(); }")
    check("a time without a valid date is dropped", ev("() => tasks.find(t => t.name === 'B').custom") is None)
    ev("() => { tasks.find(t => t.name === 'B').custom = { date1: '2026-09-08', date1Time: '25:99' }; normalizeData(); }")
    check("an invalid time is dropped, the date stays", cust("B") == ["2026-09-08", None])
    ev("() => { tasks.find(t => t.name === 'B').custom = { date1: '2026-09-08', date1Time: '09:15' }; normalizeData(); }")
    w = ev("() => getComputedStyle(document.getElementById('main')).getPropertyValue('--task-cols')").split()
    check("the custom Date column is 150px wide in a minute-mode plan (104px in a day-mode one)", "150px" in w, w)

    # ---------------------------------------------------------------- the list editor
    plan(True)
    cell("A").click(); pg.wait_for_selector(".inline-edit")
    check("an empty date cell opens the typed date-time box", pg.get_attribute(".inline-edit", "type") == "text" and pg.get_attribute(".inline-edit", "placeholder") == "dd.mm.yyyy hh:mm" and pg.input_value(".inline-edit") == "")
    pg.keyboard.press("Escape")
    edit("A", "07.09.2026 14:30")
    check("'07.09.2026 14:30' stores the date and its time", cust("A") == ["2026-09-07", "14:30"], cust("A"))
    cell("A").click(); pg.wait_for_selector(".inline-edit")
    check("the editor shows both (there is no default time for a custom date — only the one you gave it)", pg.input_value(".inline-edit") == "07.09.2026 14:30")
    pg.keyboard.press("Escape")
    edit("A", "16:00")
    check("a time alone keeps the date", cust("A") == ["2026-09-07", "16:00"], cust("A"))
    edit("A", "09.09.2026")
    check("a date alone clears the time", cust("A") == ["2026-09-09", None], cust("A"))
    cell("A").click(); pg.wait_for_selector(".inline-edit")
    check("...and the editor then shows just the date", pg.input_value(".inline-edit") == "09.09.2026")
    pg.keyboard.press("Escape")
    edit("A", "10.09.2026 8:05")
    check("a single-digit hour is fine", cust("A") == ["2026-09-10", "08:05"], cust("A"))
    edit("A", "nonsense")
    check("nonsense is refused and changes nothing", cust("A") == ["2026-09-10", "08:05"] and "isn't a date" in toast(), toast())
    edit("A", "")
    check("clearing the cell clears date and time", cust("A") == [None, None])
    edit("B", "14:00")
    check("a time alone on an empty date cell is refused (there is no date to keep)", cust("B") == [None, None])
    edit("A", "07.09.2026 14:30")
    cell("A").click(); pg.wait_for_selector(".inline-edit")
    pg.evaluate("() => { const d = document.querySelector('.inline-date-hidden'); d.value = '2026-09-11'; d.dispatchEvent(new Event('change', { bubbles: true })); }"); pg.wait_for_timeout(150)
    check("picking a date with the calendar button keeps the typed time", cust("A") == ["2026-09-11", "14:30"], cust("A"))

    # ---------------------------------------------------------------- the task dialog
    pg.click(f".grid-row[data-id='{tid('A')}'] .icon-btn[title=Edit]"); pg.wait_for_selector("#taskModalBg.open")
    pg.click("#taskTabBtnCustom")
    check("the Custom tab shows a time input beside the date field in a minute-mode plan", pg.locator("#cf_date1_time").count() == 1 and ev("() => document.getElementById('cf_date1_time').value") == "14:30" and ev("() => document.getElementById('cf_date1').value") == "2026-09-11")
    check("...and the badge counts the field once, not its time as well", pg.inner_text("#taskCustomBadge") == "1", pg.inner_text("#taskCustomBadge"))
    pg.fill("#cf_date1_time", "17:45"); pg.dispatch_event("#cf_date1_time", "change")
    pg.click("#taskModalBg .modal-footer button.btn-primary"); pg.wait_for_timeout(200)
    check("saving the dialog stores the new time", cust("A") == ["2026-09-11", "17:45"], cust("A"))
    pg.click(f".grid-row[data-id='{tid('A')}'] .icon-btn[title=Edit]"); pg.wait_for_selector("#taskModalBg.open"); pg.click("#taskTabBtnCustom")
    pg.fill("#cf_date1_time", ""); pg.click("#taskModalBg .modal-footer button.btn-primary"); pg.wait_for_timeout(200)
    check("emptying the time input keeps a date-only value", cust("A") == ["2026-09-11", None], cust("A"))

    # ---------------------------------------------------------------- Excel, copied rows, MS Project
    ev("() => { const t = tasks.find(x => x.name === 'A'); t.custom = { date1: '2026-09-11', date1Time: '17:45' }; const b = tasks.find(x => x.name === 'B'); b.custom = { date1: '2026-09-12' }; save(); render(); }")
    pg.click("#dataMenuBtn"); pg.click("#excelExportItem"); pg.wait_for_selector("#excelModalBg.open")
    with pg.expect_download() as d: pg.click("#excelExportBtn")
    path = os.path.join(TMP, "cd.xlsx"); d.value.save_as(path)
    ws = openpyxl.load_workbook(path)["Tasks"]
    hdr = {c.value: c.column for c in ws[4]}
    ca, cb = ws.cell(row=5, column=hdr["Review"]), ws.cell(row=6, column=hdr["Review"])
    check("Excel: a custom date with a time is a real datetime cell formatted with hh:mm", ca.value.hour == 17 and ca.value.minute == 45 and "hh:mm" in ca.number_format, (ca.value, ca.number_format))
    check("...and one without a time stays a plain date", cb.value.hour == 0 and "hh:mm" not in cb.number_format, (cb.value, cb.number_format))
    ev("() => setSelection(tasks.map(t => t.id))")
    tsv = ev("() => buildClip().tsv")
    check("copied rows carry the time ('11.09.2026 17:45') and the date-only value ('12.09.2026')", "11.09.2026 17:45" in tsv and "12.09.2026" in tsv and "12.09.2026 " not in tsv.replace("12.09.2026\t", "12.09.2026").replace("12.09.2026\n", "12.09.2026"), tsv)
    xml = ev("() => buildMspdi()")
    check("the MS Project notes carry it too", "Review: 11.09.2026 17:45" in xml and "Review: 12.09.2026" in xml)

    # ---------------------------------------------------------------- bulk
    ev("() => { setSelection([tasks.find(t => t.name === 'A').id]); openBulkModal(); }"); pg.wait_for_timeout(200)
    pg.check("#bulkOn-date1"); pg.fill("#bulkVal-date1", "2026-10-01")
    pg.click("#bulkModalBg .modal-footer button.btn-primary"); pg.wait_for_timeout(200)
    check("the Edit-tasks dialog sets a plain date and clears the old time", cust("A") == ["2026-10-01", None], cust("A"))

    # ---------------------------------------------------------------- back to days
    ev("() => { tasks.find(t => t.name === 'A').custom = { date1: '2026-09-11', date1Time: '17:45' }; normalizeData(); }")
    pg.click("#scheduleMenuBtn"); pg.click("#planPrecisionItem"); pg.wait_for_selector("#precisionModalBg.open"); pg.click("#precisionDay"); pg.click("#precisionModalBg .modal-footer button.btn-primary"); pg.wait_for_timeout(200)
    check("switching to days with a time still set warns first", pg.locator("#confirmModalBg.open").count() == 1 and "lose its time of day" in pg.inner_text("#confirmModalBody"), pg.inner_text("#confirmModalBody") if pg.locator("#confirmModalBg.open").count() else None)
    pg.click("#confirmModalActionBtn"); pg.wait_for_timeout(200)
    check("switching the plan back to days drops the time and keeps the date", cust("A") == ["2026-09-11", None], cust("A"))
    check("a day-mode plan's cell shows the date only, in a 104px column", cell("A").inner_text() == "11.09.2026" and "150px" not in ev("() => getComputedStyle(document.getElementById('main')).getPropertyValue('--task-cols')"))
    cell("A").click(); pg.wait_for_selector(".inline-edit")
    check("...and its editor is the native date input, as before", pg.get_attribute(".inline-edit", "type") == "date")
    pg.keyboard.press("Escape")
    pg.click(f".grid-row[data-id='{tid('A')}'] .icon-btn[title=Edit]"); pg.wait_for_selector("#taskModalBg.open"); pg.click("#taskTabBtnCustom")
    check("...and the dialog has no time input", pg.locator("#cf_date1_time").count() == 0)
    pg.keyboard.press("Escape")

    check("no console errors or page errors across the whole run", not errors, errors[:5])
    n_ok, n_all = sum(results), len(results)
    print(f"\n{n_ok}/{n_all} checks passed")
    b.close()
    raise SystemExit(0 if n_ok == n_all else 1)
