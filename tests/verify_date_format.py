# -*- coding: utf-8 -*-
"""Date and time format (Plan settings → Date & time; a per-device preference): four 10-character date styles and a 24- or 12-hour clock —
display, typed entry in both orders, the typed boxes and the list's editors (the native date input is gone), prefs, columns that still FIT in every
combination (cells, open editors, dialogs, Task Form, Gantt ticks), Print, Excel, CSV round trip, the settings tab. See "Date and time format" in CLAUDE.md."""
import os, re, tempfile, json
from playwright.sync_api import sync_playwright
import openpyxl
URL = os.environ.get("MILESTONE_URL", "http://127.0.0.1:8937/milestone.html")
errors, results = [], []
def check(name, cond, detail=""):
    results.append(bool(cond)); print(("PASS  " if cond else "FAIL  ") + name + (f"   [{str(detail)[:600]}]" if not cond and detail else ""))
DATES = ["dd.mm.yyyy", "dd/mm/yyyy", "mm/dd/yyyy", "yyyy-mm-dd"]
SEED = """(minute) => { historyCoalesceMs = 0; project.workDays = [0,1,2,3,4,5,6]; delete project.fieldNames; delete project.timeUnit; delete project.workHours; tasks.length = 0; let o = 0;
  if (minute) { project.timeUnit = 'minute'; }
  const mk = (id, name, s, e, x) => Object.assign({ id, name, parentId: null, order: o++, startDate: s, endDate: e, progress: 0, milestone: false, color: null, predecessors: [], collapsed: false, updatedAt: 1 }, x || {});
  tasks.push(mk('a', 'Widest dates', '2030-08-28', '2030-12-30', Object.assign({ actualStart: '2030-08-28', actualFinish: '2030-12-30', deadline: '2030-12-30', baselines: { 0: ['2030-08-28', '2030-12-30'] }, custom: { date1: '2030-08-28' } },
       minute ? { startTime: '12:30', endTime: '17:00', actualStartTime: '12:30', actualFinishTime: '17:00' } : {})));
  tasks.push(mk('b', 'Second', '2031-01-02', '2031-01-03', { taskMode: 'manual' }));
  project.fieldNames = { date1: 'Due' };
  project.baselines = { 0: { setAt: '2030-01-01' } }; project.compareBaseline = 0;
  for (const c of ['actualStart', 'actualFinish', 'deadline', 'baselineStart', 'baselineFinish', 'date1']) colHidden.delete(c);
  normalizeData(); save(); setSelection([]); currentView = 'tasks'; applyDisplayFormat('dd.mm.yyyy', '24h'); render(); resetHistory(); }"""
with sync_playwright() as p:
    b = p.chromium.launch()
    ctx = b.new_context(viewport={"width": 1700, "height": 900}, accept_downloads=True); ctx.add_init_script("delete window.showOpenFilePicker; delete window.showSaveFilePicker; delete window.showDirectoryPicker")
    pg = ctx.new_page(); pg.on("pageerror", lambda e: errors.append(str(e))); pg.on("console", lambda m: errors.append(m.text) if m.type == "error" else None)
    pg.goto(URL); pg.wait_for_selector("#undoBtn"); pg.evaluate("() => localStorage.clear()"); pg.reload(); pg.wait_for_selector("#undoBtn")
    ev = pg.evaluate

    # ================================================== defaults, helpers
    check("a fresh profile reads exactly as before: dd.mm.yyyy and the 24-hour clock", ev("() => [dateFormat, timeFormat, fmtDate('2030-08-28'), fmtTime('14:30'), document.documentElement.classList.contains('time12')]") == ["dd.mm.yyyy", "24h", "28.08.2030", "14:30", False])
    for df, want in zip(DATES, ["28.08.2030", "28/08/2030", "08/28/2030", "2030-08-28"]):
        check(f"{df}: 28 Aug 2030 reads {want}, ten characters", ev("(f) => { dateFormat = f; const r = fmtDate('2030-08-28'); dateFormat = 'dd.mm.yyyy'; return r; }", df) == want and len(want) == 10)
    check("short dates (week ticks, yearly holidays): 28.08. / 28/08 / 08/28 / 08-28", ev("() => DATE_FORMATS.map(f => { dateFormat = f.id; return fmtDateShort('2030-08-28'); }).concat([(dateFormat = 'dd.mm.yyyy', 'x')])").__getitem__(slice(0, 4)) == ["28.08.", "28/08", "08/28", "08-28"])
    check("12-hour times: 00:05 → 12:05 AM, 09:30 → 9:30 AM, 12:00 → 12:00 PM, 14:30 → 2:30 PM, 23:59 → 11:59 PM; 24-hour is unchanged", ev("() => { timeFormat = '12h'; const r = ['00:05', '09:30', '12:00', '14:30', '23:59'].map(fmtTime); timeFormat = '24h'; return r.concat([fmtTime('14:30')]); }") == ["12:05 AM", "9:30 AM", "12:00 PM", "2:30 PM", "11:59 PM", "14:30"])
    check("hour labels of the Hours scale in the 12-hour clock: 12a, 1 … 11, 12p, 1 … 11 (20px wide cells)", ev("() => { timeFormat = '12h'; const r = [0, 1, 11, 12, 13, 23].map(fmtHourLabel); timeFormat = '24h'; return r; }") == ["12a", "1", "11", "12p", "1", "11"])
    # ---------------------------------------------------------------- parsing
    P = lambda df, s: ev("([f, s]) => { dateFormat = f; const r = parseUserDate(s); dateFormat = 'dd.mm.yyyy'; return r; }", [df, s])
    check("ISO is always understood, in every format (what the calendar button gives)", all(P(df, "2030-08-28") == "2030-08-28" for df in DATES))
    check("dd.mm.yyyy and dd/mm/yyyy read day first; mm/dd/yyyy reads month first", P("dd.mm.yyyy", "07.09.2030") == "2030-09-07" and P("dd/mm/yyyy", "07/09/2030") == "2030-09-07" and P("mm/dd/yyyy", "07/09/2030") == "2030-07-09")
    check("any separator is fine ( . / - ), the order comes from the format", P("mm/dd/yyyy", "12-24-2030") == "2030-12-24" and P("dd.mm.yyyy", "24/12/2030") == "2030-12-24")
    check("when the chosen order is impossible the other is tried: '24/12/2030' in the US format is 24 December; '12/24/2030' in the day-first formats too", P("mm/dd/yyyy", "24/12/2030") == "2030-12-24" and P("dd/mm/yyyy", "12/24/2030") == "2030-12-24")
    check("impossible dates are refused in either order (31.02., month 13, junk)", all(P(df, x) is None for df in DATES for x in ["31/02/2030", "13/13/2030", "ab.cd.2030", "5/5/30", ""]))
    B = lambda df, s: ev("([f, s]) => { dateFormat = f; const r = parseBoxDate(s); dateFormat = 'dd.mm.yyyy'; return r; }", [df, s])
    check("a two-digit year and a missing year follow the format too (24.12.30 / 12/24/30 / 12/24 = this year)", B("dd.mm.yyyy", "24.12.30") == "2030-12-24" and B("mm/dd/yyyy", "12/24/30") == "2030-12-24" and B("mm/dd/yyyy", "12/24") == ev("() => todayStr().slice(0, 4)") + "-12-24" and B("dd/mm/yyyy", "24/12") == ev("() => todayStr().slice(0, 4)") + "-12-24")
    T = lambda s: ev("(s) => parseBoxTime(s)", s)
    check("typed times: 24-hour and 12-hour both read, as 24-hour storage (14:30, 2:30pm, 2:30 PM, 12 am, 9, 930)", [T(x) for x in ["14:30", "2:30pm", "2:30 PM", "12am", "12 PM", "9", "930", "11:59 pm"]] == ["14:30", "14:30", "14:30", "00:00", "12:00", "09:00", "09:30", "23:59"])
    DT = lambda df, s: ev("([f, s]) => { dateFormat = f; const r = parseUserDateTime(s); dateFormat = 'dd.mm.yyyy'; return r; }", [df, s])
    check("a date and a time together: '12/24/2030 2:30 PM', '24.12.2030 14:30', '2030-12-24T14:30'", DT("mm/dd/yyyy", "12/24/2030 2:30 PM") == {"iso": "2030-12-24", "time": "14:30"} and DT("dd.mm.yyyy", "24.12.2030 14:30") == {"iso": "2030-12-24", "time": "14:30"} and DT("yyyy-mm-dd", "2030-12-24T14:30") == {"iso": "2030-12-24", "time": "14:30"})
    check("a date alone keeps the time (time null), a time alone keeps the date (iso null); a bare year is not a time; junk is refused", DT("mm/dd/yyyy", "12/24/2030") == {"iso": "2030-12-24", "time": None} and DT("dd.mm.yyyy", "2:30 PM") == {"iso": None, "time": "14:30"} and DT("dd.mm.yyyy", "2030") is None and DT("dd.mm.yyyy", "TBD") is None)
    check("break lists in the 12-hour clock: '12:00 PM-1:00 PM, 3:00 PM-3:15 PM' reads like '12:00-13:00, 15:00-15:15'; shown back in the chosen clock", ev("() => { const a = parseBreaksText('12:00 PM-1:00 PM, 3:00 PM-3:15 PM'), b = parseBreaksText('12:00-13:00, 15:00-15:15'); timeFormat = '12h'; const t = breaksText(a); timeFormat = '24h'; return [JSON.stringify(a) === JSON.stringify(b), t]; }") == [True, "12:00 PM-1:00 PM, 3:00 PM-3:15 PM"])

    # ================================================== the plan itself never changes
    ev(SEED, False); pg.wait_for_timeout(250)
    plan0 = ev("() => JSON.stringify(syncPayload())")
    ev("() => applyDisplayFormat('mm/dd/yyyy', '12h')"); pg.wait_for_timeout(200)
    check("the plan's data is byte-identical after a format change (dates stay ISO, times 24-hour: only the display changed)", ev("() => JSON.stringify(syncPayload())") == plan0)
    check("the list shows it: Start of the first task reads 08/28/2030", "08/28/2030" in pg.inner_text("#gridRows .grid-row[data-id='a']"))
    prefs = ev("() => JSON.parse(localStorage.getItem('milestone-prefs'))")
    check("it is a per-device preference (milestone-prefs: df, tf), not part of the plan", prefs.get("df") == "mm/dd/yyyy" and prefs.get("tf") == "12h" and "df" not in json.loads(plan0).get("project", {}))
    pg.reload(); pg.wait_for_selector("#undoBtn"); pg.wait_for_timeout(300)
    check("...and it survives a reload (including the html.time12 class)", ev("() => [dateFormat, timeFormat, document.documentElement.classList.contains('time12')]") == ["mm/dd/yyyy", "12h", True])
    ev("() => { const d = JSON.parse(localStorage.getItem('milestone-prefs')); d.df = 'banana'; d.tf = '99h'; localStorage.setItem('milestone-prefs', JSON.stringify(d)); }"); pg.reload(); pg.wait_for_selector("#undoBtn")
    check("a damaged preference falls back to the default", ev("() => [dateFormat, timeFormat]") == ["dd.mm.yyyy", "24h"])
    ev("() => { const d = JSON.parse(localStorage.getItem('milestone-prefs')); delete d.df; delete d.tf; localStorage.setItem('milestone-prefs', JSON.stringify(d)); }"); pg.reload(); pg.wait_for_selector("#undoBtn")
    check("an absent preference (every existing install) is the default — nothing changes for them", ev("() => [dateFormat, timeFormat]") == ["dd.mm.yyyy", "24h"])

    # ================================================== columns that FIT — every combination, day mode and Hours & minutes
    FIT = """() => { const cols = visibleTaskCols(), idx = k => cols.indexOf(k) + 1, out = [];
      for (const row of document.querySelectorAll('#gridRows .grid-row')) for (const k of ['start', 'end', 'actualStart', 'actualFinish', 'deadline', 'baselineStart', 'baselineFinish', 'date1']) {
        const c = row.children[idx(k)]; if (!c) continue; const e = c.querySelector('.grid-cell-dim') || c; if (c.scrollWidth > c.clientWidth + 0.5 || e.scrollWidth > e.clientWidth + 0.5) out.push(k + ':' + c.textContent.trim() + ':' + c.scrollWidth + '>' + c.clientWidth); }
      return out; }"""
    EDIT = """() => { const cols = visibleTaskCols(), out = []; const fields = [['a', 'start'], ['a', 'finish'], ['a', 'actualStart'], ['a', 'actualFinish'], ['a', 'deadline'], ['a', 'date1'], ['b', 'start']];
      for (const [id, f] of fields) { startInlineEdit(id, f); const row = document.querySelector('#gridRows .grid-row[data-id="' + id + '"]'), w = row.querySelector('.inline-wrap'); if (!w) { out.push(f + ': no editor'); cancelInlineEdit(); continue; }
        const i = w.querySelector('input[type=text]'), r = w.getBoundingClientRect(), cell = [...row.children].find(c => c === w), nx = w.nextElementSibling, prev = w.previousElementSibling;
        const room = nx ? nx.getBoundingClientRect().left - r.right : 99, gap = r.left - prev.getBoundingClientRect().right;
        if (i.scrollWidth > i.clientWidth + 1) out.push(f + ': text clipped ' + i.scrollWidth + '>' + i.clientWidth + ' [' + i.value + ']'); if (room < -2.5) out.push(f + ': spills ' + room); if (gap < -2.5) out.push(f + ': spills left ' + gap);
        cancelInlineEdit(); }
      return out; }"""
    for minute in (False, True):
        for tf in ("24h", "12h"):
            for df in DATES:
                ev(SEED, minute); ev("([d, t]) => applyDisplayFormat(d, t)", [df, tf]); pg.wait_for_timeout(120)
                tag = f"{'Hours & minutes' if minute else 'day mode'} · {df} · {tf}"
                bad = ev(FIT)
                check(f"every date cell fits its column — {tag}", not bad, bad)
                bad = ev(EDIT)
                check(f"every open date editor fits its column and shows its whole text — {tag}", not bad, bad)
    # widths themselves
    ev(SEED, True); ev("() => applyDisplayFormat('dd.mm.yyyy', '24h')"); pg.wait_for_timeout(120)
    cw = lambda k: ev("(k) => { const i = visibleTaskCols().indexOf(k) + 1; return Math.round(document.querySelector('#gridRows .grid-row').children[i].getBoundingClientRect().width); }", k)
    check("Hours & minutes, 24-hour: Start and Finish are 138px, as always", cw("start") == 138 and cw("end") == 138, (cw("start"), cw("end")))
    ev("() => applyDisplayFormat('dd.mm.yyyy', '12h')"); pg.wait_for_timeout(120)
    check("...in the 12-hour clock they widen to 160px (and only the date-and-time columns)", cw("start") == 160 and cw("end") == 160 and cw("actualStart") == 160 and cw("duration") == ev("() => 104"), (cw("start"), cw("duration")))
    ev("() => { colWidths.start = 200; save(); render(); }"); pg.wait_for_timeout(100)
    ev("() => applyDisplayFormat('dd.mm.yyyy', '24h')"); pg.wait_for_timeout(120)
    check("a width you gave a date column is dropped when the clock changes (it belonged to the other one): back to the 138px default", cw("start") == 138 and ev("() => colWidths.start") is None, (cw("start"), ev("() => colWidths.start")))
    ev(SEED, False); ev("() => applyDisplayFormat('mm/dd/yyyy', '24h')"); pg.wait_for_timeout(120)
    check("day mode: a date column stays 104px in every date format (all ten characters long)", all(ev("(f) => { applyDisplayFormat(f, '24h'); return Math.round(document.querySelector('#gridRows .grid-row').children[visibleTaskCols().indexOf('start') + 1].getBoundingClientRect().width); }", f) == 104 for f in DATES))
    ev("() => applyDisplayFormat('dd.mm.yyyy', '24h')")

    # ================================================== typed entry in the list
    ev(SEED, False); ev("() => applyDisplayFormat('mm/dd/yyyy', '24h')"); pg.wait_for_timeout(150)
    ev("() => startInlineEdit('b', 'start')"); pg.wait_for_timeout(150)
    check("the editor of a Manual task offers the format as its hint ('mm/dd/yyyy or text')", pg.get_attribute(".inline-wrap > input[type=text]", "placeholder") == "mm/dd/yyyy or text")
    pg.fill(".inline-wrap > input[type=text]", "12/24/2030"); pg.keyboard.press("Enter"); pg.wait_for_timeout(200)
    check("typing 12/24/2030 in the US format sets 24 December 2030", ev("() => tasks.find(t => t.id === 'b').startDate") == "2030-12-24")
    ev("() => startInlineEdit('a', 'start')"); pg.wait_for_timeout(150)
    check("an Auto task's editor says just the format (no 'or text')", pg.get_attribute(".inline-wrap > input[type=text]", "placeholder") == "mm/dd/yyyy" and pg.locator("#gridRows input[type=date].inline-edit").count() == 0)
    pg.fill(".inline-wrap > input[type=text]", "9/1/2031"); pg.keyboard.press("Enter"); pg.wait_for_timeout(200)
    check("...and 9/1/2031 is 1 September (month first)", ev("() => tasks.find(t => t.id === 'a').startDate") == "2031-09-01")
    ev("() => startInlineEdit('a', 'start')"); pg.wait_for_timeout(120); pg.fill(".inline-wrap > input[type=text]", "banana"); pg.keyboard.press("Enter"); pg.wait_for_timeout(200)
    toast = ev("() => document.getElementById('toastMsg').textContent")
    check("an Auto task refuses text and says how to write a date in the chosen format", "isn't a date" in toast and "12/24/2026" in toast and ev("() => tasks.find(t => t.id === 'a').startDate") == "2031-09-01", toast)
    ev("() => startInlineEdit('a', 'deadline')"); pg.wait_for_timeout(120)
    check("the Deadline editor is the typed box as well (no native date input anywhere in the list)", pg.locator("#gridRows .inline-wrap").count() == 1 and pg.locator("#gridRows input[type=date].inline-edit").count() == 0)
    pg.fill(".inline-wrap > input[type=text]", "11/30/2031"); pg.keyboard.press("Enter"); pg.wait_for_timeout(200)
    check("...Deadline 11/30/2031 is saved as 30 November", ev("() => tasks.find(t => t.id === 'a').deadline") == "2031-11-30")
    ev("() => startInlineEdit('a', 'start')"); pg.wait_for_timeout(120)
    pg.evaluate("() => { const w = document.querySelector('.inline-wrap'); const n = w.querySelector('input[type=date]'); n.value = '2031-10-15'; n.dispatchEvent(new Event('change', { bubbles: true })); }"); pg.wait_for_timeout(250)
    check("the calendar button's ISO value is read in every format", ev("() => tasks.find(t => t.id === 'a').startDate") == "2031-10-15", ev("() => tasks.find(t => t.id === 'a').startDate"))

    # ================================================== Hours & minutes: typed date + time in the list
    ev(SEED, True); ev("() => applyDisplayFormat('mm/dd/yyyy', '12h')"); pg.wait_for_timeout(150)
    txt = pg.inner_text("#gridRows .grid-row[data-id='a']")
    check("the list shows a date and a 12-hour time: '08/28/2030 12:30 PM' and the finish '12/30/2030 5:00 PM'", "08/28/2030 12:30 PM" in txt and "12/30/2030 5:00 PM" in txt, txt)
    ev("() => startInlineEdit('a', 'start')"); pg.wait_for_timeout(120)
    check("the editor opens with the same text and the hint 'mm/dd/yyyy h:mm AM/PM'", pg.input_value(".inline-wrap > input[type=text]") == "08/28/2030 12:30 PM" and pg.get_attribute(".inline-wrap > input[type=text]", "placeholder") == "mm/dd/yyyy h:mm AM/PM", pg.input_value(".inline-wrap > input[type=text]"))
    pg.fill(".inline-wrap > input[type=text]", "09/02/2030 9:15 AM"); pg.keyboard.press("Enter"); pg.wait_for_timeout(250)
    check("typing '09/02/2030 9:15 AM' stores 2030-09-02 and 09:15 (24-hour)", ev("() => { const t = tasks.find(x => x.id === 'a'); return [t.startDate, t.startTime]; }") == ["2030-09-02", "09:15"], ev("() => { const t = tasks.find(x => x.id === 'a'); return [t.startDate, t.startTime]; }"))
    ev("() => startInlineEdit('a', 'start')"); pg.wait_for_timeout(120); pg.fill(".inline-wrap > input[type=text]", "14:45"); pg.keyboard.press("Enter"); pg.wait_for_timeout(250)
    check("a time typed on its own in 24-hour form is read too (14:45), the date is kept", ev("() => { const t = tasks.find(x => x.id === 'a'); return [t.startDate, t.startTime]; }") == ["2030-09-02", "14:45"])
    # the calendar button keeps a 12-hour time
    ev("() => startInlineEdit('a', 'start')"); pg.wait_for_timeout(120)
    pg.evaluate("() => { const n = document.querySelector('.inline-wrap input[type=date]'); n.value = '2030-09-10'; n.dispatchEvent(new Event('change', { bubbles: true })); }"); pg.wait_for_timeout(250)
    check("picking a date with the calendar button keeps the typed 12-hour time", ev("() => { const t = tasks.find(x => x.id === 'a'); return [t.startDate, t.startTime]; }") == ["2030-09-10", "14:45"])

    # ================================================== dialogs: boxes show the format, keep their value, and fit
    ev("() => openTaskModal('a')"); pg.wait_for_timeout(300)
    vals = ev("() => ['taskStartInput', 'taskStartTimeInput', 'taskEndInput', 'taskEndTimeInput'].map(i => document.getElementById(i)._getRaw())")
    check("the task dialog's boxes show the format: date '09/10/2030', time '2:45 PM'", vals[0] == "09/10/2030" and vals[1] == "2:45 PM", vals)
    check("...their ISO / 24-hour values are untouched", ev("() => [document.getElementById('taskStartInput').value, document.getElementById('taskStartTimeInput').value]") == ["2030-09-10", "14:45"])
    check("...and their hints follow", ev("() => [taskStartInput.placeholder, taskStartTimeInput.placeholder]") == ["mm/dd/yyyy", "h:mm AM/PM"])
    fit = ev("() => [...document.querySelectorAll('#taskModalBg input.dt-date, #taskModalBg input.dt-time')].filter(e => e.offsetParent && e.scrollWidth > e.clientWidth + 1).map(e => e.id + ':' + e._getRaw() + ':' + e.scrollWidth + '>' + e.clientWidth)")
    check("every date and time box of the task dialog shows its whole text (12-hour clock, 'mm/dd/yyyy')", not fit, fit)
    wrap = ev("() => { const s = document.getElementById('taskStartInput').getBoundingClientRect(), t = document.getElementById('taskStartTimeInput').getBoundingClientRect(); return Math.abs(s.top - t.top) < 4; }")
    check("...and the time sits on the date's line (no wrapping)", wrap)
    ev("() => { applyDisplayFormat('dd.mm.yyyy', '24h'); }"); pg.wait_for_timeout(150)
    vals = ev("() => ['taskStartInput', 'taskStartTimeInput'].map(i => document.getElementById(i)._getRaw())")
    check("changing the format while the dialog is open rewrites its boxes at once ('10.09.2030', '14:45') without changing what they hold", vals == ["10.09.2030", "14:45"] and ev("() => [taskStartInput.value, taskStartTimeInput.value]") == ["2030-09-10", "14:45"], vals)
    ev("() => document.getElementById('taskModalBg').classList.remove('open')")
    for df, tf in (("yyyy-mm-dd", "12h"), ("mm/dd/yyyy", "12h")):
        ev("([d, t]) => applyDisplayFormat(d, t)", [df, tf]); ev("() => openTaskModal('a')"); pg.wait_for_timeout(250)
        fit = ev("() => [...document.querySelectorAll('#taskModalBg input.dt-date, #taskModalBg input.dt-time')].filter(e => e.offsetParent && e.scrollWidth > e.clientWidth + 1).map(e => e.id)")
        check(f"task dialog boxes fit ({df}, {tf})", not fit, fit); ev("() => document.getElementById('taskModalBg').classList.remove('open')")
    # precision dialog
    ev("() => applyDisplayFormat('dd.mm.yyyy', '12h')"); pg.click("#planSettingsBtn"); pg.click(".modal-bg.open [data-settings-tab='precision']"); pg.wait_for_selector("#precisionModalBg.open"); pg.wait_for_timeout(250)
    check("Scheduling precision shows the working hours in the 12-hour clock ('8:00 AM'–'5:00 PM') and they fit their boxes", ev("() => [whStart._getRaw(), whEnd._getRaw()]") == ["8:00 AM", "5:00 PM"] and ev("() => ['whStart', 'whEnd', 'whBreakFrom', 'whBreakTo'].every(i => { const e = document.getElementById(i); return e.scrollWidth <= e.clientWidth + 1; })"))
    check("...the Schedule-menu style summary reads in the clock too", "8:00 AM–5:00 PM" in ev("() => precisionSummary()"))
    pg.keyboard.press("Escape"); pg.wait_for_timeout(200)
    # task form
    ev("() => { showTaskForm = true; setSelection(['a']); render(); }"); pg.wait_for_timeout(300)
    check("the Task Form's Start / Finish boxes show the 12-hour times and fit (and its row doesn't overflow)", ev("() => [tfStartTime._getRaw(), tfStart._getRaw()]") == ["2:45 PM", "10.09.2030"] and ev("() => ['tfStart', 'tfStartTime', 'tfFinish', 'tfFinishTime'].every(i => { const e = document.getElementById(i); return e.scrollWidth <= e.clientWidth + 1; })") and ev("() => { const p = document.getElementById('taskFormPane'); return p.scrollWidth <= p.clientWidth + 1; }"))
    ev("() => { showTaskForm = false; render(); }")

    # ================================================== Gantt, Print
    ev("() => { applyDisplayFormat('mm/dd/yyyy', '12h'); currentView = 'gantt'; zoom = 'week'; render(); }"); pg.wait_for_timeout(300)
    ticks = pg.locator(".gantt-tick.zoom-week").all_inner_texts() if pg.locator(".gantt-tick.zoom-week").count() else pg.locator(".gantt-tick").all_inner_texts()
    check("Week ticks read 'CW n' over the Monday's date in the chosen order (mm/dd)", any(re.search(r"CW \d+\s*\n?\s*\d\d/\d\d", t) for t in ticks[:12]), ticks[:4])
    ev("() => { zoom = 'hour'; render(); }"); pg.wait_for_timeout(300)
    hl = pg.locator(".gantt-tick.zoom-hour-h").all_inner_texts()[:26]
    check("the Hours scale labels its cells 12a, 1, 2 … 11, 12p … in the 12-hour clock, none clipped", hl[0].strip() == "12a" and hl[12].strip() == "12p" and hl[1].strip() == "1" and ev("() => [...document.querySelectorAll('.gantt-tick.zoom-hour-h')].every(e => e.scrollWidth <= e.clientWidth + 1)"), hl[:14])
    day = pg.locator(".gantt-tick.zoom-hour-day").first.inner_text()
    check("...and the day cells show the date in the format ('Tue 09/…')", re.search(r"\d\d/\d\d/\d{4}", day), day)
    ev("() => { currentView = 'tasks'; zoom = 'fit'; render(); }")
    built = ev("() => { const o = Object.assign(readPrintOptions ? {} : {}, { paper: 'A4', orient: 'landscape', cols: ['id', 'name', 'start', 'end', 'duration'], range: 'all', scale: 'auto', baseline: false, critical: false, deps: false, today: false, shade: false, expandAll: false, title: 'T' }); const r = buildPrintPages(o); return r.pages[0]; }")
    check("Print draws the dates and times in the chosen format (09/10/2030 2:45 PM)", "09/10/2030 2:45 PM" in built, re.findall(r">([^<>]*\d{2}[/.]\d{2}[/.]\d{4}[^<>]*)<", built)[:4])
    ev("() => applyDisplayFormat('dd.mm.yyyy', '24h')")

    # ================================================== Excel and CSV
    def excel(cols="shown"):
        pg.click("#dataMenuBtn"); pg.click("#excelExportItem"); pg.wait_for_selector("#excelModalBg.open"); pg.check(f"input[name='excelCols'][value='{cols}']")
        with pg.expect_download() as d: pg.click("#excelExportBtn")
        path = os.path.join(tempfile.mkdtemp(), "x.xlsx"); d.value.save_as(path); return openpyxl.load_workbook(path)["Tasks"]
    ev(SEED, False); ev("() => applyDisplayFormat('mm/dd/yyyy', '24h')"); pg.wait_for_timeout(150)
    ws = excel(); hdr = {c.value: c.column for c in ws[4]}
    cell = ws.cell(row=5, column=hdr["Start"])
    check("Excel: a date cell is a real date with the chosen number format (mm\\/dd\\/yyyy)", cell.number_format == "mm\\/dd\\/yyyy" and str(cell.value)[:10] == "2030-08-28", (cell.number_format, cell.value))
    ev("() => applyDisplayFormat('yyyy-mm-dd', '24h')"); ws = excel(); cell = ws.cell(row=5, column=hdr["Start"])
    check("...yyyy\\-mm\\-dd for the ISO style", cell.number_format == "yyyy\\-mm\\-dd", cell.number_format)
    ev(SEED, True); ev("() => applyDisplayFormat('dd/mm/yyyy', '12h')"); pg.wait_for_timeout(150)
    ws = excel(); hdr = {c.value: c.column for c in ws[4]}; cell = ws.cell(row=5, column=hdr["Start"])
    check("Hours & minutes, 12-hour: the cell format is 'dd\\/mm\\/yyyy h:mm AM/PM' and the column is wide enough for it (>= 24)", cell.number_format == "dd\\/mm\\/yyyy h:mm AM/PM" and ws.column_dimensions[openpyxl.utils.get_column_letter(hdr["Start"])].width >= 24, (cell.number_format, ws.column_dimensions[openpyxl.utils.get_column_letter(hdr["Start"])].width))
    ev("() => applyDisplayFormat('dd.mm.yyyy', '24h')"); ws = excel(); cell = ws.cell(row=5, column=hdr["Start"])
    check("Hours & minutes, 24-hour: 'dd\\.mm\\.yyyy hh:mm' in a column >= 19 wide (before it was 13: Excel would have shown ####)", cell.number_format == "dd\\.mm\\.yyyy hh:mm" and ws.column_dimensions[openpyxl.utils.get_column_letter(hdr["Start"])].width >= 19)
    # CSV round trip with an ambiguous date in the US format
    ev(SEED, False); ev("() => { tasks.find(t => t.id === 'b').startDate = '2030-03-04'; tasks.find(t => t.id === 'b').endDate = '2030-03-05'; tasks.find(t => t.id === 'b').taskMode = 'auto'; save(); applyDisplayFormat('mm/dd/yyyy', '24h'); }")
    csv = ev("() => buildCsv({ scope: 'all', columns: 'shown', sep: ',' })")
    check("CSV export carries the chosen format ('03/04/2030' = 4 March in the US style)", "03/04/2030" in csv, csv[:300])
    res = ev("(txt) => { const r = tasksFromTable(parseDelimited(txt), { dateFormat: 'auto' }); return r.tasks.map(t => t.startDate); }", csv)
    check("...and Milestone's own import reads that ambiguous date back as 4 March (the chosen format settles it), not 3 April", res[1] == "2030-03-04", res)
    ev("() => applyDisplayFormat('dd/mm/yyyy', '24h')"); csv = ev("() => buildCsv({ scope: 'all', columns: 'shown', sep: ',' })")
    res = ev("(txt) => { const r = tasksFromTable(parseDelimited(txt), { dateFormat: 'auto' }); return r.tasks.map(t => t.startDate); }", csv)
    check("...in the day-first slash style '04/03/2030' is read as 4 March too", "04/03/2030" in csv and res[1] == "2030-03-04", res)

    # ================================================== the settings dialog (tiles, instant, real preview, "use my browser's format", "This device")
    ev(SEED, False); ev("() => applyDisplayFormat('dd.mm.yyyy', '24h')")
    pg.click("#planSettingsBtn"); pg.wait_for_selector(".modal-bg.open")
    tabs = [t.strip() for t in pg.locator(".modal-bg.open .settings-tab").all_inner_texts()]
    check("Plan settings: five tabs — Calendar, Days or hours, Scheduling rules, Custom fields, Formats (currency and the date and time format share the last one)", tabs == ["Calendar", "Days or hours", "Scheduling rules", "Custom fields", "Formats"] and pg.locator(".modal-bg.open .settings-tab-sep").count() == 0, tabs)
    check("...and the tabs fit the 680px dialog (nothing scrolls sideways)", ev("() => { const t = document.querySelector('.modal-bg.open .settings-tabs'); return t.scrollWidth <= t.clientWidth; }"))
    pg.click(".modal-bg.open [data-settings-tab='format']"); pg.wait_for_selector("#formatModalBg.open"); pg.wait_for_timeout(200)
    check("the dates are the tile labels (24.12.2026 big, dd.mm.yyyy small), four tiles, the current one selected", [x.split("\n")[0] for x in pg.locator("#fmtDateOpts .fmt-tile").all_inner_texts()] == ["24.12.2026", "24/12/2026", "12/24/2026", "2026-12-24"] and [x.split("\n")[1] for x in pg.locator("#fmtDateOpts .fmt-tile").all_inner_texts()] == DATES and pg.locator("#fmtDateOpts .fmt-tile[aria-checked=true]").get_attribute("data-df") == "dd.mm.yyyy")
    check("...two clock tiles (14:30 / 2:30 PM) with 24-hour selected", [x.split("\n")[0] for x in pg.locator("#fmtTimeOpts .fmt-tile").all_inner_texts()] == ["14:30", "2:30 PM"] and pg.locator("#fmtTimeOpts .fmt-tile[aria-checked=true]").get_attribute("data-tf") == "24h")
    check("scope badges say what each part is saved with: 'This plan' on the currency, 'This device' on the dates and times (tooltips explain)", [x.strip() for x in pg.locator("#formatModalBg .scope-badge").all_inner_texts()] == ["This plan", "This device"] and "not in the plan" in pg.get_attribute("#formatModalBg .device-badge", "title"))
    check("there is no Save and no Cancel — just Done — and the footer says changes apply as you click", pg.locator("#formatModalBg .modal-footer button").all_inner_texts() == ["Done"] and "as you click" in pg.inner_text("#formatModalBg .modal-footer"))
    pg.click("#fmtDateOpts .fmt-tile[data-df='mm/dd/yyyy']"); pg.wait_for_timeout(250)
    check("clicking a tile applies it at once: the format, the list behind the dialog, the selection ring", ev("() => dateFormat") == "mm/dd/yyyy" and "08/28/2030" in pg.inner_text("#gridRows .grid-row[data-id='a']") and pg.locator("#fmtDateOpts .fmt-tile[aria-checked=true]").get_attribute("data-df") == "mm/dd/yyyy")
    pg.click("#fmtTimeOpts .fmt-tile[data-tf='12h']"); pg.wait_for_timeout(250)
    check("...the same for the clock (and the date tile you chose stays)", ev("() => [dateFormat, timeFormat]") == ["mm/dd/yyyy", "12h"] and pg.locator("#fmtDateOpts .fmt-tile[aria-checked=true]").get_attribute("data-df") == "mm/dd/yyyy")
    check("it is saved as a preference at once (survives a reload without any Save)", ev("() => JSON.parse(localStorage.getItem('milestone-prefs')).df") == "mm/dd/yyyy" and ev("() => JSON.parse(localStorage.getItem('milestone-prefs')).tf") == "12h")
    prev = pg.inner_text("#fmtPreview")
    check("the preview is a real task of this plan: its name and its Start / Finish in the chosen format ('Widest dates  Start 08/28/2030  Finish 12/30/2030'; day mode shows no time)", "Widest dates" in prev and "08/28/2030" in prev and "12/30/2030" in prev and "PM" not in prev and "a task of this plan" in pg.inner_text("#formatModalBg .fmt-preview-label"), prev)
    check("...with a hint of how to type, and a caption under the clock tiles saying times belong to Hours & minutes plans", "12/24/2026" in pg.inner_text("#fmtPreviewHint") and "Hours & minutes" in pg.inner_text("#formatModalBg .fmt-cols .fmt-cap"), pg.inner_text("#fmtPreviewHint"))
    pg.click("#fmtDateOpts .fmt-tile[data-df='yyyy-mm-dd']"); pg.wait_for_timeout(200)
    check("the preview follows every click (ISO style now)", "2030-08-28" in pg.inner_text("#fmtPreview"), pg.inner_text("#fmtPreview"))
    # Hours & minutes plan: times in the preview; the empty plan: a sample
    ev(SEED, True); ev("() => applyDisplayFormat('dd.mm.yyyy', '12h')"); ev("() => renderFormatModal()"); pg.wait_for_timeout(150)
    prev = pg.inner_text("#fmtPreview")
    check("in an Hours & minutes plan the preview shows the times in the chosen clock ('28.08.2030 12:30 PM … 5:00 PM') and the hint says how to type them", "28.08.2030 12:30 PM" in prev and "30.12.2030 5:00 PM" in prev and "times like 2:30 PM" in pg.inner_text("#fmtPreviewHint"), prev)
    ev("() => { tasks.length = 0; renderFormatModal(); }"); pg.wait_for_timeout(100)
    check("with no task in the plan the preview is a sample, marked as one", "Sample task" in pg.inner_text("#fmtPreview") and "a sample" in pg.inner_text("#formatModalBg .fmt-preview-label"))
    ev(SEED, False); ev("() => { applyDisplayFormat('dd.mm.yyyy', '24h'); renderFormatModal(); }")
    # fits
    fit = ev("() => [...document.querySelectorAll('#formatModalBg .fmt-tile b, #formatModalBg .fmt-tile small, #formatModalBg #fmtSystemBtn, #formatModalBg .fmt-preview-row span')].filter(e => e.scrollWidth > e.clientWidth + 1).map(e => e.textContent.trim())")
    body = ev("() => { const b = document.querySelector('#formatModalBg .modal-body'); return [b.scrollWidth, b.clientWidth, b.scrollHeight, b.clientHeight]; }")
    check("every tile, the system button and the preview show their whole text, and the page needs no scrolling", not fit and body[0] <= body[1] and body[2] <= body[3], (fit, body))
    pg.set_viewport_size({"width": 720, "height": 700}); pg.wait_for_timeout(200)
    fit = ev("() => [...document.querySelectorAll('#formatModalBg .fmt-tile b, #formatModalBg .fmt-tile small')].filter(e => e.scrollWidth > e.clientWidth + 1).map(e => e.textContent.trim())")
    check("...also in a 720px-wide window (the dialog shrinks to 92%)", not fit, fit)
    pg.set_viewport_size({"width": 1700, "height": 900})
    # Escape closes straight away (nothing to discard) and the tab can be re-opened
    pg.keyboard.press("Escape"); pg.wait_for_timeout(200)
    check("Escape closes the page at once — no 'discard your changes?' question — and keeps what was chosen", not ev("() => document.getElementById('formatModalBg').classList.contains('open')") and not ev("() => document.getElementById('confirmModalBg').classList.contains('open')"))
    pg.click("#planSettingsBtn"); pg.click(".modal-bg.open [data-settings-tab='calendar']"); pg.wait_for_timeout(150); pg.click(".modal-bg.open [data-settings-tab='format']"); pg.wait_for_selector("#formatModalBg.open")
    check("switching to the tab from another one works (no question: nothing to lose)", pg.locator("#fmtDateOpts .fmt-tile[aria-checked=true]").count() == 1)
    pg.keyboard.press("Escape"); pg.wait_for_timeout(150)
    # use my browser's format — one context per locale
    for loc, df, tf, eg in (("en-US", "mm/dd/yyyy", "12h", "12/24/2026 · 12-hour"), ("de-DE", "dd.mm.yyyy", "24h", "24.12.2026 · 24-hour"), ("en-GB", "dd/mm/yyyy", "24h", "24/12/2026 · 24-hour"), ("sv-SE", "yyyy-mm-dd", "24h", "2026-12-24 · 24-hour")):
        c2 = b.new_context(viewport={"width": 1440, "height": 800}, locale=loc); c2.add_init_script("delete window.showOpenFilePicker; delete window.showSaveFilePicker; delete window.showDirectoryPicker")
        p2 = c2.new_page(); p2.goto(URL); p2.wait_for_selector("#undoBtn")
        det = p2.evaluate("() => detectSystemFormat()")
        p2.click("#planSettingsBtn"); p2.click(".modal-bg.open [data-settings-tab='format']"); p2.wait_for_selector("#formatModalBg.open"); p2.wait_for_timeout(150)
        same0 = p2.is_disabled("#fmtSystemBtn")
        check(f"{loc}: the browser's own format is detected as {df} / {tf}", det == {"df": df, "tf": tf}, det)
        if (df, tf) == ("dd.mm.yyyy", "24h"):
            check(f"{loc}: it already matches the default, so the button says 'Same as your browser' and is disabled", same0 and "Same as your browser" in p2.inner_text("#fmtSystemBtn"))
            p2.click("#fmtDateOpts .fmt-tile[data-df='mm/dd/yyyy']"); p2.wait_for_timeout(150)
        check(f"{loc}: the button offers it with an example ('Use my browser's format {eg}')", not p2.is_disabled("#fmtSystemBtn") and eg in p2.inner_text("#fmtSystemBtn"), p2.inner_text("#fmtSystemBtn"))
        p2.click("#fmtSystemBtn"); p2.wait_for_timeout(250)
        check(f"{loc}: one click applies it — format and clock — and the button turns into the quiet 'Same as your browser'", p2.evaluate("() => [dateFormat, timeFormat]") == [df, tf] and p2.is_disabled("#fmtSystemBtn") and "Same as your browser" in p2.inner_text("#fmtSystemBtn"), p2.evaluate("() => [dateFormat, timeFormat]"))
        c2.close()
    # help
    ev("() => openHelpModal('start')"); pg.fill("#helpSearch", "date format"); pg.wait_for_timeout(300)
    check("Help mentions the setting (searchable)", "Date & time" in pg.inner_text("#helpPane-search") or "date format" in pg.inner_text("#helpPane-search").lower(), pg.inner_text("#helpPane-search")[:200])
    ev("() => closeHelpModal()")
    b.close()
check("no console errors", not errors, errors)
print(f"\n{sum(results)}/{len(results)}"); raise SystemExit(0 if all(results) else 1)
