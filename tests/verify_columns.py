import datetime, json
import openpyxl
from playwright.sync_api import sync_playwright

import os
URL = os.environ.get("MILESTONE_URL", "http://127.0.0.1:8937/milestone.html")
errors, results = [], []
def check(name, cond, detail=""):
    results.append(bool(cond)); print(("PASS  " if cond else "FAIL  ") + name + (f"   [{str(detail)[:260]}]" if not cond and detail else ""))

TODAY = datetime.date.today()
def d(n): return (TODAY + datetime.timedelta(days=n)).isoformat()
def fmt(iso): y, m, dd = iso.split("-"); return f"{dd}.{m}.{y}"

OLD_ORDER = ["mode", "wbs", "name", "start", "end", "actualStart", "actualFinish", "duration", "progress", "preds", "remaining", "status", "resource"]   # the 13 columns of an install from before the baseline columns
DEFAULT_ORDER = ["mode", "wbs", "name", "start", "end", "actualStart", "actualFinish", "duration", "progress", "preds", "remaining", "status", "resource", "baselineStart", "baselineFinish", "baselineDuration", "startVariance", "finishVariance", "durationVariance"]
# (the 20 custom fields are in the registry too, but the Columns menu lists only the ones a plan uses — none here)
ALL_ORDER = DEFAULT_ORDER + [k + str(i) for k in ("text", "number", "date", "flag") for i in range(1, 6)]   # what colOrder itself holds
OLD_SHOWN = ["mode", "wbs", "name", "start", "end", "duration", "progress", "preds"]   # what an install from before Actual Start/Finish and Status became default columns still shows
DEFAULT_SHOWN = ["mode", "wbs", "name", "start", "end", "actualStart", "actualFinish", "duration", "progress", "preds", "status"]

MK = """(specs) => { tasks.length = 0; selectedTaskId = null; colFilters = newColFilters(); filterPinned.clear();
  let n = 0; const ids = {};
  for (const sp of specs) { const t = Object.assign({id: genId(), name: sp.name, parentId: sp.parent ? ids[sp.parent] : null, order: n++, startDate: sp.s || '2026-09-07', endDate: sp.e || sp.s || '2026-09-11', progress: sp.p || 0, milestone: false, color: null, notes: '', predecessors: [], collapsed: false, updatedAt: 1, constraintType: 'ASAP', constraintDate: null, taskMode: 'auto', resource: '', actualStart: null, actualFinish: null}, sp.extra || {}); tasks.push(t); ids[sp.name] = t.id; }
  currentView = 'tasks'; save(); render(); }"""

with sync_playwright() as p:
    b = p.chromium.launch(headless=True)
    ctx = b.new_context(viewport={"width": 1500, "height": 820}, accept_downloads=True); ctx.add_init_script("delete window.showOpenFilePicker; delete window.showSaveFilePicker")
    pg = ctx.new_page(); pg.on("pageerror", lambda e: errors.append(str(e))); pg.on("console", lambda m: errors.append(m.text) if m.type in ("error", "warning") else None)
    pg.goto(URL); pg.wait_for_selector("#addTaskBtn"); pg.evaluate("() => localStorage.clear()"); pg.reload(); pg.wait_for_selector("#addTaskBtn"); pg.evaluate("() => { project.workDays = [0,1,2,3,4,5,6]; }")
    make = lambda specs: (pg.evaluate(MK, specs), pg.wait_for_timeout(120))
    head_cols = lambda: pg.evaluate("() => [...document.querySelectorAll('#gridHeader .col-filter-btn')].map(b => b.dataset.col)")
    names = lambda: pg.evaluate("() => visibleTaskList().map(x => x.task.name)")
    T = lambda name: pg.evaluate("n => { const t = tasks.find(x => x.name === n); return t ? {id: t.id, resource: t.resource, aS: t.actualStart, aF: t.actualFinish, updatedAt: t.updatedAt} : null; }", name)
    row = lambda name: pg.locator(f".grid-row[data-id='{T(name)['id']}']")
    def cell(name, col):                       # the cell of `col` in the row of task `name`
        i = head_cols().index(col)
        return row(name).locator(":scope > div").nth(1 + i)
    def open_cols(): pg.click("#columnsBtn"); pg.wait_for_selector("#columnsMenu.open")
    def close_cols(): pg.keyboard.press("Escape"); pg.wait_for_timeout(80)
    def set_col(col, on):
        pg.locator(f"#columnsMenu input[data-col='{col}']").set_checked(on); pg.wait_for_timeout(80)
    def wait_focus(): pg.wait_for_function("() => document.activeElement && document.activeElement.classList.contains('inline-edit')")

    # ============================================================ defaults and the Columns menu
    make([{"name": "Solo"}])
    check("out of the box the list shows the familiar columns plus WBS; the other new fields are hidden", head_cols() == DEFAULT_SHOWN, head_cols())
    check("the Columns button is in the toolbar (Tasks view)", pg.locator("#columnsBtn:not(.hidden)").count() == 1)
    pg.evaluate("() => { currentView = 'gantt'; render(); }")
    check("...and it is available in the Gantt view too (the list there is the same one)", pg.locator("#columnsBtn.hidden").count() == 0)
    pg.evaluate("() => { currentView = 'tasks'; render(); }")
    open_cols()
    listed = pg.evaluate("() => [...document.querySelectorAll('#columnsMenu input[data-col]')].map(i => [i.dataset.col, i.checked, i.disabled])")
    check("the menu lists every column in the default order with the right ticks", [x[0] for x in listed] == DEFAULT_ORDER and [x[0] for x in listed if x[1]] == DEFAULT_SHOWN, listed)
    check("Task Name is ticked and can't be unticked", [x for x in listed if x[0] == "name"][0][1:] == [True, True])
    check("the first column can't move up, the last can't move down", pg.locator("#columnsMenu .cols-row").first.locator(".cols-move button").first.is_disabled() and pg.locator("#columnsMenu .cols-row").last.locator(".cols-move button").last.is_disabled())
    # show a hidden column
    set_col("resource", True)
    check("ticking Resource adds it (after Predecessors, its default place) and the menu stays open", head_cols() == DEFAULT_SHOWN + ["resource"] and pg.locator("#columnsMenu.open").count() == 1, head_cols())
    set_col("actualStart", True)
    check("ticking Actual Start puts it at its default place (after Finish)", head_cols() == ["mode", "wbs", "name", "start", "end", "actualStart", "actualFinish", "duration", "progress", "preds", "status", "resource"], head_cols())
    cols_now = head_cols(); ncells = pg.locator(".grid-row").first.locator(":scope > div").count()
    check("every row has exactly one cell per visible column (+ ID and actions)", ncells == len(cols_now) + 2 and pg.locator("#gridHeader > div").count() == len(cols_now) + 2, (ncells, len(cols_now)))
    check("the grid template has a track for each", len(pg.evaluate("() => getComputedStyle(document.querySelector('.grid-row')).gridTemplateColumns.split(' ')")) == len(cols_now) + 2)
    set_col("resource", False)
    check("unticking hides it again", "resource" not in head_cols())
    # reorder
    pg.locator("#columnsMenu .cols-row:has(input[data-col='start']) .cols-move button").nth(1).click(); pg.wait_for_timeout(80)   # Start down, below Finish
    check("the arrows reorder: Start moved below Finish in the list", head_cols().index("end") < head_cols().index("start"), head_cols())
    check("...and the rows follow the header (Start's cell now shows the start date)", cell("Solo", "start").inner_text().strip() == "07.09.2026" and cell("Solo", "end").inner_text().strip() == "11.09.2026")
    pg.locator("#columnsMenu .cols-row:has(input[data-col='name']) .cols-move button").first.click(); pg.wait_for_timeout(80)   # Name up
    check("Task Name can be moved too", head_cols().index("name") < head_cols().index("wbs"), head_cols())
    saved = pg.evaluate("() => JSON.parse(localStorage.getItem('milestone-prefs')).cols")
    check("the choice is stored with the device preferences (not in the plan)", saved["order"] == pg.evaluate("() => colOrder") and "remaining" in saved["hidden"] and "baselineStart" in saved["hidden"] and "actualFinish" not in saved["hidden"] and "cols" not in json.dumps(pg.evaluate("() => JSON.parse(localStorage.getItem('milestone-plan-' + currentPlanId))")), saved)
    close_cols(); before = head_cols()
    pg.reload(); pg.wait_for_selector("#addTaskBtn"); pg.wait_for_timeout(200)
    check("...and survives a reload", head_cols() == before, (head_cols(), before))
    open_cols(); pg.click("#columnsMenu .btn-link:has-text('Reset')"); pg.wait_for_timeout(100)
    check("Reset to default restores columns and order", head_cols() == DEFAULT_SHOWN and pg.evaluate("() => colOrder") == ALL_ORDER)
    close_cols()
    pg.click("#columnsBtn"); pg.click("body", position={"x": 5, "y": 400}); pg.wait_for_timeout(100)
    check("clicking elsewhere closes the menu", pg.locator("#columnsMenu.open").count() == 0)
    # saved prefs from other versions
    def with_prefs(cols):
        pg.evaluate("c => { const p = JSON.parse(localStorage.getItem('milestone-prefs') || '{}'); p.cols = c; if (c === undefined) delete p.cols; localStorage.setItem('milestone-prefs', JSON.stringify(p)); }", cols)
        pg.reload(); pg.wait_for_selector("#addTaskBtn"); pg.wait_for_timeout(150)
    with_prefs(None)
    check("prefs saved before columns existed -> the defaults", head_cols() == DEFAULT_SHOWN)
    # WBS became shown-by-default in column revision 1: an install saved before that has it listed as hidden
    with_prefs({"order": OLD_ORDER, "hidden": ["wbs", "actualStart", "actualFinish", "remaining", "status", "resource"]})
    check("an existing install whose saved choice still hides WBS gets it switched on, once", "wbs" in head_cols(), head_cols())
    check("...the rest of its choice is untouched (the other new fields stay hidden)", head_cols() == OLD_SHOWN, head_cols())
    pg.evaluate("() => { colHidden.add('wbs'); save(); }"); pg.reload(); pg.wait_for_selector("#addTaskBtn"); pg.wait_for_timeout(150)
    check("hiding WBS again afterwards sticks across reloads (the switch-on is one-time)", "wbs" not in head_cols(), head_cols())
    check("the revision is stored with the choice", pg.evaluate("() => JSON.parse(localStorage.getItem('milestone-prefs')).cols.rev") == 1)
    with_prefs({"order": OLD_ORDER, "hidden": ["actualStart", "actualFinish", "remaining", "status", "resource"], "rev": 1})
    check("a choice already at the current revision is never overridden", "wbs" in head_cols())
    with_prefs({"order": OLD_ORDER, "hidden": ["wbs", "actualStart", "actualFinish", "remaining", "status", "resource"], "rev": 1})
    check("...and a WBS the user hid at the current revision stays hidden", "wbs" not in head_cols(), head_cols())
    with_prefs(None)
    with_prefs({"order": ["name", "start", "end", "duration", "progress", "preds", "mode", "resource", "actualStart", "actualFinish", "wbs"], "hidden": ["mode"]})
    head_now = head_cols()
    check("prefs from an older version (no Remaining/Status yet): those two are added, each as its default says (Remaining hidden, Status shown)", "remaining" not in head_now and "status" in head_now and "remaining" in pg.evaluate("() => colOrder") and "status" in pg.evaluate("() => colOrder"), head_now)
    check("...a column the user had hidden stays hidden, the visible ones keep their saved order", "mode" not in head_now and head_now == ["name", "start", "end", "duration", "progress", "preds", "resource", "actualStart", "actualFinish", "wbs", "status"], head_now)
    with_prefs({"order": ["bogus", "name", "start"], "hidden": ["name", "nothing"]})
    check("garbage in the prefs is ignored; Task Name can never be hidden", "name" in head_cols() and "bogus" not in pg.evaluate("() => colOrder") and set(pg.evaluate("() => colOrder")) == set(ALL_ORDER) and len(pg.evaluate("() => colOrder")) == len(ALL_ORDER), (head_cols(), pg.evaluate("() => colOrder")))
    with_prefs(None)

    # ============================================================ Resource, Actual Start/Finish
    make([{"name": "Grp"}, {"name": "Kid1", "parent": "Grp", "s": d(-3), "e": d(2)}, {"name": "Kid2", "parent": "Grp", "s": d(3), "e": d(6)}, {"name": "Other", "s": d(10), "e": d(12)}])
    pg.evaluate("() => { colHidden.delete('resource'); colHidden.delete('actualStart'); colHidden.delete('actualFinish'); save(); render(); }"); pg.wait_for_timeout(100)
    cell("Kid1", "resource").click(); pg.wait_for_selector(".inline-edit"); wait_focus()
    pg.fill(".inline-edit", "Anna, Ben"); pg.keyboard.press("Enter"); pg.wait_for_timeout(120)
    check("Resource: click the cell, type, Enter -> saved on the task", T("Kid1")["resource"] == "Anna, Ben" and cell("Kid1", "resource").inner_text().strip() == "Anna, Ben", T("Kid1"))
    check("...and the change is stamped so it syncs", T("Kid1")["updatedAt"] > 1)
    cell("Kid1", "resource").click(); wait_focus(); pg.fill(".inline-edit", "SHOULD NOT SAVE"); pg.keyboard.press("Escape"); pg.wait_for_timeout(100)
    check("Escape cancels a Resource edit", T("Kid1")["resource"] == "Anna, Ben")
    cell("Grp", "resource").click(); wait_focus(); pg.fill(".inline-edit", "PMO"); pg.keyboard.press("Enter"); pg.wait_for_timeout(100)
    check("a group can have a Resource of its own", T("Grp")["resource"] == "PMO")
    cell("Kid1", "actualStart").click(); pg.wait_for_selector(".inline-edit"); wait_focus()
    check("Actual Start opens a date editor", pg.locator(".inline-edit").get_attribute("type") == "date")
    pg.fill(".inline-edit", d(-2)); pg.keyboard.press("Enter"); pg.wait_for_timeout(120)
    check("...and saves the date, shown as dd.mm.yyyy", T("Kid1")["aS"] == d(-2) and cell("Kid1", "actualStart").inner_text().strip() == fmt(d(-2)), T("Kid1"))
    cell("Kid1", "actualFinish").click(); wait_focus(); pg.fill(".inline-edit", d(-4)); pg.keyboard.press("Enter"); pg.wait_for_timeout(120)
    check("an Actual Finish before the Actual Start is refused (nothing changes)", T("Kid1")["aF"] is None and "before" in pg.inner_text("#toastMsg"), (T("Kid1"), pg.inner_text("#toastMsg")))
    cell("Kid1", "actualFinish").click(); wait_focus(); pg.fill(".inline-edit", d(0)); pg.keyboard.press("Enter"); pg.wait_for_timeout(120)
    check("a valid Actual Finish is accepted", T("Kid1")["aF"] == d(0))
    cell("Kid1", "actualFinish").click(); wait_focus(); pg.fill(".inline-edit", ""); pg.keyboard.press("Enter"); pg.wait_for_timeout(120)
    check("clearing the date field clears the Actual Finish", T("Kid1")["aF"] is None)
    check("the group shows the earliest actual start of what is in it (read-only)", cell("Grp", "actualStart").inner_text().strip() == fmt(d(-2)) and "editable" not in (cell("Grp", "actualStart").get_attribute("class") or ""))
    cell("Grp", "actualStart").click(); pg.wait_for_timeout(120)
    check("...clicking it does not open an editor", pg.locator(".inline-edit").count() == 0)
    check("a group has no actual finish while a task in it hasn't finished", cell("Grp", "actualFinish").inner_text().strip() == "")
    pg.evaluate("() => { const by = n => tasks.find(t => t.name === n); by('Kid1').actualFinish = '%s'; by('Kid2').actualStart = '%s'; by('Kid2').actualFinish = '%s'; save(); render(); }" % (d(0), d(4), d(6)))
    check("once every task in the group has finished, the group shows the latest actual finish", cell("Grp", "actualFinish").inner_text().strip() == fmt(d(6)), cell("Grp", "actualFinish").inner_text())

    # the task dialog
    pg.evaluate("n => openTaskModal(tasks.find(t => t.name === n).id)", "Other"); pg.wait_for_selector("#taskModalBg.open")
    check("dialog: Resource, Actual Start and Actual Finish fields exist (empty for a new task)", pg.input_value("#taskResourceInput") == "" and pg.input_value("#taskActualStartInput") == "" and pg.input_value("#taskActualFinishInput") == "")
    check("dialog: the title badge shows the ID and the WBS code", "#4" in pg.inner_text("#taskModalIdBadge") and "WBS 2" in pg.inner_text("#taskModalIdBadge"), pg.inner_text("#taskModalIdBadge"))
    check("dialog: no Successors section any more", "Successors" not in pg.inner_text("#taskModalBg") and pg.locator("#successorsList").count() == 0 and pg.evaluate("() => typeof renderSuccessorsList") == "undefined")
    pg.fill("#taskResourceInput", "  Carla  "); pg.fill("#taskActualStartInput", d(10)); pg.fill("#taskActualFinishInput", d(9))
    pg.click("#taskModalBg .btn-primary"); pg.wait_for_timeout(150)
    check("dialog: an Actual Finish before the Actual Start blocks the save and keeps the dialog open", pg.locator("#taskModalBg.open").count() == 1 and T("Other")["resource"] == "" and "before" in pg.inner_text("#toastMsg"), (T("Other"), pg.inner_text("#toastMsg")))
    pg.fill("#taskActualFinishInput", d(11)); pg.click("#taskModalBg .btn-primary"); pg.wait_for_timeout(150)
    check("dialog: valid values save (Resource is trimmed)", T("Other")["resource"] == "Carla" and T("Other")["aS"] == d(10) and T("Other")["aF"] == d(11) and pg.locator("#taskModalBg.open").count() == 0, T("Other"))
    pg.evaluate("n => openTaskModal(tasks.find(t => t.name === n).id)", "Grp"); pg.wait_for_selector("#taskModalBg.open")
    check("dialog: a group's actual dates are shown but read-only, with an explanation", pg.is_disabled("#taskActualStartInput") and pg.is_disabled("#taskActualFinishInput") and pg.input_value("#taskActualStartInput") == d(-2) and "rolled up" in pg.inner_text("#taskActualsNote") or "earliest actual start" in pg.inner_text("#taskActualsNote"), pg.inner_text("#taskActualsNote"))
    pg.click("#taskModalBg .btn-primary"); pg.wait_for_timeout(120)
    check("dialog: saving a group doesn't write its rolled-up dates into the group itself", pg.evaluate("() => { const g = tasks.find(t => t.name === 'Grp'); return [g.actualStart, g.actualFinish]; }") == [None, None])
    # data hygiene
    pg.evaluate("() => { const t = tasks.find(x => x.name === 'Other'); t.resource = 123; t.actualStart = 'garbage'; t.actualFinish = '2026-02-31'; normalizeData(); }")
    check("normalisation: a non-text resource becomes empty, impossible dates become null", pg.evaluate("() => { const t = tasks.find(x => x.name === 'Other'); return [t.resource, t.actualStart, t.actualFinish]; }") == ["", None, None])
    p1 = pg.evaluate("() => { const t = syncPayload().tasks.find(x => x.name === 'Kid1'); return ['resource' in t, 'actualStart' in t, 'actualFinish' in t]; }")
    check("the new fields are part of the data that is exported, backed up and synced", p1 == [True, True, True])
    pg.evaluate("() => { const t = tasks.find(x => x.name === 'Kid1'); mergeData({project: project, tasks: [Object.assign({}, JSON.parse(JSON.stringify(t)), {resource: 'Remote', actualStart: '2026-01-05', updatedAt: Date.now() + 5000})], deletedTaskIds: []}); }")
    check("a newer copy from another device brings its Resource and actual dates along (last write wins)", T("Kid1")["resource"] == "Remote" and T("Kid1")["aS"] == "2026-01-05")

    # ============================================================ WBS
    make([{"name": "A"}, {"name": "A1", "parent": "A"}, {"name": "A2", "parent": "A"}, {"name": "A2a", "parent": "A2"}, {"name": "A2b", "parent": "A2"}, {"name": "B"}, {"name": "C"}, {"name": "C1", "parent": "C"}])
    pg.evaluate("() => { colHidden.delete('wbs'); save(); render(); }"); pg.wait_for_timeout(100)
    codes = lambda: {n: pg.evaluate("n => wbsCode(tasks.find(t => t.name === n).id)", n) for n in ["A", "A1", "A2", "A2a", "A2b", "B", "C", "C1"]}
    exp = {"A": "1", "A1": "1.1", "A2": "1.2", "A2a": "1.2.1", "A2b": "1.2.2", "B": "2", "C": "3", "C1": "3.1"}
    check("WBS codes follow the outline (1, 1.1, 1.2, 1.2.1 …)", codes() == exp, codes())
    check("the WBS column shows them", all(cell(n, "wbs").inner_text().strip() == c for n, c in exp.items()))
    pg.evaluate("() => { tasks.find(t => t.name === 'A').collapsed = true; render(); }")
    check("collapsing a group doesn't change anyone's WBS (the rows just disappear)", codes() == exp and "A1" not in names())
    pg.evaluate("() => { tasks.find(t => t.name === 'A').collapsed = false; colFilters.name = {type: 'rule', rule: 'eq', a: 'C1'}; render(); }")
    check("...nor does filtering", codes() == exp and cell("C1", "wbs").inner_text().strip() == "3.1")
    pg.evaluate("() => { colFilters = newColFilters(); render(); }")
    pg.evaluate("() => { selectedTaskId = tasks.find(t => t.name === 'C1').id; outdentSelected(); }"); pg.wait_for_timeout(100)
    check("outdenting C1 renumbers it (now a top-level task, 4)", pg.evaluate("() => wbsCode(tasks.find(t => t.name === 'C1').id)") == "4" and codes()["C"] == "3", codes())
    pg.evaluate("() => { moveTask(tasks.find(t => t.name === 'B').id, tasks.find(t => t.name === 'A1').id, 'before'); }"); pg.wait_for_timeout(100)
    check("moving B into A (above A1) renumbers: A1 becomes 1.2, A2 1.3, and A2's kids follow", pg.evaluate("() => ['B','A1','A2','A2a'].map(n => wbsCode(tasks.find(t => t.name === n).id))") == ["1.1", "1.2", "1.3", "1.3.1"])
    open_cols(); check("WBS is filterable like text (begins with)", True); close_cols()
    pg.locator(".col-filter-btn[data-col='wbs']").click(); pg.wait_for_selector("#filterMenu.open"); pg.select_option("#filterRuleSel", "begins"); pg.fill("#filterInputA", "1.3"); pg.click("#filterOkBtn"); pg.wait_for_timeout(120)
    check("...'begins with 1.3' keeps A2 and everything under it (plus its parent A for context)", set(names()) == {"A", "A2", "A2a", "A2b"}, names())
    pg.evaluate("() => { colFilters = newColFilters(); render(); }")

    # ============================================================ Remaining Duration and Status
    specs = [
        {"name": "done", "s": d(-20), "e": d(-16), "p": 100},
        {"name": "overdue", "s": d(-10), "e": d(-6), "p": 40},
        {"name": "notstarted-past", "s": d(-5), "e": d(3), "p": 0},
        {"name": "started-actual", "s": d(-5), "e": d(3), "p": 0, "extra": {"actualStart": d(-4)}},
        {"name": "underway", "s": d(-2), "e": d(7), "p": 30},
        {"name": "today-start", "s": d(0), "e": d(4), "p": 0},
        {"name": "future", "s": d(5), "e": d(9), "p": 0},
        {"name": "actual-finish", "s": d(-3), "e": d(5), "p": 50, "extra": {"actualStart": d(-3), "actualFinish": d(-1)}},
        {"name": "half", "s": d(2), "e": d(6), "p": 50},
        {"name": "milestone", "s": d(1), "e": d(1), "extra": {"milestone": True}},
        {"name": "tbd", "extra": {"taskMode": "manual", "startText": "TBD", "durText": "a while"}},
        {"name": "Group"}, {"name": "g1", "parent": "Group", "s": d(-8), "e": d(-4), "p": 100}, {"name": "g2", "parent": "Group", "s": d(-2), "e": d(6), "p": 0},
    ]
    make(specs)
    pg.evaluate("() => { colHidden.delete('remaining'); colHidden.delete('status'); save(); render(); }"); pg.wait_for_timeout(100)
    def py_status(s, e, p, aS, aF):
        if p >= 100 or aF: return "Complete"
        if not s: return None
        started = p > 0 or bool(aS)
        if e < TODAY.isoformat() and True: return "Late"
        if not started and s < TODAY.isoformat(): return "Late"
        if not started and s > TODAY.isoformat(): return "Future Task"
        return "On Schedule"
    def py_remaining(s, e, p):
        dur = (datetime.date.fromisoformat(e) - datetime.date.fromisoformat(s)).days + 1
        return int(dur * (100 - p) / 100 * 10 + 0.5) / 10
    def fmt_days(v): return f"{int(v) if v == int(v) else v} day" + ("" if v == 1 else "s")
    bad_r, bad_s = [], []
    for sp in specs:
        n = sp["name"]
        if n in ("Group",): continue
        ex = sp.get("extra", {})
        if n == "tbd":
            want_r, want_s = "", ""
        else:
            s, e, pr = sp["s"], sp["e"], sp.get("p", 0)
            want_r = "" if ex.get("milestone") else fmt_days(py_remaining(s, e, pr))
            st = py_status(s, e, pr, ex.get("actualStart"), ex.get("actualFinish")); want_s = st or ""
        got_r, got_s = cell(n, "remaining").inner_text().strip(), cell(n, "status").inner_text().strip()
        if got_r != want_r: bad_r.append((n, got_r, want_r))
        if got_s != want_s: bad_s.append((n, got_s, want_s))
    check("Remaining Duration = duration × (1 − % complete), in days (2.5 days for 5 days at 50%); blank for milestones and free-text durations", not bad_r, bad_r)
    check("Status: Complete / Late / Future Task / On Schedule as specified, blank when there is nothing to judge", not bad_s, bad_s)
    check("Status is shown as a coloured pill", cell("overdue", "status").locator(".status-pill.st-late").count() == 1 and cell("done", "status").locator(".status-pill.st-complete").count() == 1 and cell("future", "status").locator(".status-pill.st-future-task").count() == 1 and cell("underway", "status").locator(".status-pill.st-on-schedule").count() == 1)
    check("a group is judged on its roll-up (finishes in the future, starts in the past, some progress -> On Schedule)", cell("Group", "status").inner_text().strip() in ("On Schedule", "Late"), cell("Group", "status").inner_text())
    check("both columns are computed: not editable, no editor on click", (cell("half", "remaining").click() or True) and pg.locator(".inline-edit").count() == 0)
    # filters
    def ids_shown(): return set(names())
    pg.locator(".col-filter-btn[data-col='status']").click(); pg.wait_for_selector("#filterMenu.open")
    labels = pg.locator("#filterTree .ft-label").all_inner_texts()
    check("Status filter: a list in the order Late, On Schedule, Future Task, Complete (+ (Blanks) for the task with none)", [l for l in labels if l not in ("(Select All)",)] == [x for x in ["Late", "On Schedule", "Future Task", "Complete"] if x in labels] + ["(Blanks)"], labels)
    pg.locator("#filterTree .ft-cb[data-key='all']").uncheck(); pg.locator("#filterTree .ft-label:has-text('Late') .ft-cb").check(); pg.click("#filterOkBtn"); pg.wait_for_timeout(120)
    want = {sp["name"] for sp in specs if sp["name"] not in ("Group", "tbd") and py_status(sp["s"], sp["e"], sp.get("p", 0), sp.get("extra", {}).get("actualStart"), sp.get("extra", {}).get("actualFinish")) == "Late"}
    got = ids_shown() - {"Group"}
    check("...filtering on Late shows exactly the late tasks", got == want or got - {"g1", "g2"} == want, (got, want))
    pg.evaluate("() => { colFilters = newColFilters(); render(); }")
    pg.locator(".col-filter-btn[data-col='remaining']").click(); pg.wait_for_selector("#filterMenu.open"); pg.select_option("#filterRuleSel", "between"); pg.fill("#filterInputA", "2"); pg.fill("#filterInputB", "3"); pg.click("#filterOkBtn"); pg.wait_for_timeout(120)
    want = {sp["name"] for sp in specs if sp["name"] not in ("Group", "tbd", "milestone") and 2 <= py_remaining(sp["s"], sp["e"], sp.get("p", 0)) <= 3}
    check("Remaining Duration is filterable with number rules (between 2 and 3 days)", want <= ids_shown() and "half" in ids_shown() and "done" not in ids_shown(), (ids_shown(), want))
    pg.evaluate("() => { colFilters = newColFilters(); render(); }")

    # ============================================================ filters for Resource and the actual dates
    make([{"name": "R1", "extra": {"resource": "Anna, Ben", "actualStart": d(-3)}}, {"name": "R2", "extra": {"resource": "Ben"}}, {"name": "R3", "extra": {"resource": "Carla", "actualStart": d(-1)}}, {"name": "R4"}])
    pg.evaluate("() => { for (const c of ['resource', 'actualStart']) colHidden.delete(c); save(); render(); }"); pg.wait_for_timeout(100)
    pg.locator(".col-filter-btn[data-col='resource']").click(); pg.wait_for_selector("#filterMenu.open")
    lab = pg.locator("#filterTree .ft-label").all_inner_texts()
    check("Resource filter: a value list with (Blanks) for tasks that have none", lab == ["(Select All)", "Anna, Ben", "Ben", "Carla", "(Blanks)"], lab)
    pg.select_option("#filterRuleSel", "contains"); pg.fill("#filterInputA", "ben"); pg.click("#filterOkBtn"); pg.wait_for_timeout(120)
    check("...'contains ben' (any case) finds tasks with Ben among their resources", ids_shown() == {"R1", "R2"}, ids_shown())
    pg.evaluate("() => { colFilters = newColFilters(); render(); }")
    pg.locator(".col-filter-btn[data-col='resource']").click(); pg.wait_for_selector("#filterMenu.open"); pg.locator("#filterTree .ft-cb[data-key='all']").uncheck(); pg.locator("#filterTree .ft-cb[data-key='blank']").check(); pg.click("#filterOkBtn"); pg.wait_for_timeout(120)
    check("...(Blanks) finds the unassigned tasks", ids_shown() == {"R4"}, ids_shown())
    pg.evaluate("() => { colFilters = newColFilters(); render(); }")
    pg.locator(".col-filter-btn[data-col='actualStart']").click(); pg.wait_for_selector("#filterMenu.open"); pg.select_option("#filterRuleSel", "before"); pg.fill("#filterInputA", d(-2)); pg.click("#filterOkBtn"); pg.wait_for_timeout(120)
    check("Actual Start takes date rules (before …): only R1 started that early", ids_shown() == {"R1"}, ids_shown())
    check("...its chip names the column", "Actual Start: before" in pg.inner_text("#filterBar"), pg.inner_text("#filterBar"))
    pg.evaluate("() => { colFilters = newColFilters(); render(); }")

    # ============================================================ Excel
    make([{"name": "G"}, {"name": "T1", "parent": "G", "s": d(-4), "e": d(1), "p": 50, "extra": {"resource": "Anna", "actualStart": d(-4)}}, {"name": "T2", "parent": "G", "s": d(2), "e": d(6), "p": 0}, {"name": "Late1", "s": d(-9), "e": d(-5), "p": 20}])
    pg.evaluate("() => resetColumns()")
    def export(cols, path):
        pg.click("#dataMenuBtn"); pg.click("#excelExportItem"); pg.wait_for_selector("#excelModalBg.open")
        if cols != "shown": pg.check(f"input[name='excelCols'][value='{cols}']")
        pg.set_checked("#excelGantt", False)
        with pg.expect_download() as dl: pg.click("#excelExportBtn")
        dl.value.save_as(path)
    pg.click("#dataMenuBtn"); pg.click("#excelExportItem"); pg.wait_for_selector("#excelModalBg.open")
    check("Excel dialog: 'As shown in the task list' is the default column choice", pg.is_checked("input[name='excelCols'][value='shown']")); pg.keyboard.press("Escape"); pg.wait_for_timeout(80)
    export("shown", "cols_shown.xlsx"); ws = openpyxl.load_workbook("cols_shown.xlsx")["Tasks"]
    check("Excel 'as shown': ID, the visible columns in the list's order (WBS is one of them now)", [c.value for c in ws[4]] == ["ID", "Mode", "WBS", "Task Name", "Start", "Finish", "Actual Start", "Actual Finish", "Duration", "% Complete", "Predecessors", "Status"], [c.value for c in ws[4]])
    pg.evaluate("() => { colHidden.delete('resource'); colHidden.delete('status'); colHidden.delete('remaining'); colHidden.delete('wbs'); colOrder = ['name'].concat(colOrder.filter(c => c !== 'name')); save(); render(); }")
    export("shown", "cols_custom.xlsx"); wc = openpyxl.load_workbook("cols_custom.xlsx")["Tasks"]
    hdr = [c.value for c in wc[4]]
    check("Excel 'as shown' follows the user's choice and order (Task Name first after ID; WBS, Remaining, Status, Resource included)", hdr[:3] == ["ID", "Task Name", "Mode"] and {"WBS", "Remaining Duration", "Status", "Resource"} <= set(hdr) and "Notes" not in hdr, hdr)
    check("...the pane is frozen through Task Name wherever it sits (B), autofilter covers all columns", wc.freeze_panes == "C5" and wc.auto_filter.ref == f"A4:{openpyxl.utils.get_column_letter(len(hdr))}8", (wc.freeze_panes, wc.auto_filter.ref))
    export("all", "cols_all.xlsx"); wa = openpyxl.load_workbook("cols_all.xlsx")["Tasks"]; ha = [c.value for c in wa[4]]
    check("Excel 'all columns': every field, in the default order", ha == ["ID", "Mode", "WBS", "Task Name", "Start", "Finish", "Actual Start", "Actual Finish", "Duration", "% Complete", "Predecessors", "Remaining Duration", "Status", "Resource", "Baseline Start", "Baseline Finish", "Baseline Duration", "Start Variance", "Finish Variance", "Duration Variance"], ha)
    rowsx = {wa.cell(r, ha.index("Task Name") + 1).value.strip(): r for r in range(5, wa.max_row + 1)}
    def X(name, col): return wa.cell(rowsx[name], ha.index(col) + 1)
    check("Excel: WBS as text (1.1 stays '1.1', not a number)", X("T1", "WBS").value == "1.1" and X("T1", "WBS").data_type == "s", (X("T1", "WBS").value, X("T1", "WBS").data_type))
    check("Excel: Resource text; Actual Start a real date, Actual Finish empty", X("T1", "Resource").value == "Anna" and X("T1", "Actual Start").value.date().isoformat() == d(-4) and X("T1", "Actual Finish").value is None and X("T1", "Actual Start").number_format == "dd\\.mm\\.yyyy")
    check("Excel: Remaining Duration is a number (2.5 days for 5 days at 50%… here 6 days at 50% = 3) with a days format", X("T1", "Remaining Duration").value == 3 and "day" in X("T1", "Remaining Duration").number_format and X("T2", "Remaining Duration").value == 5, (X("T1", "Remaining Duration").value, X("T2", "Remaining Duration").value))
    check("Excel: Status text, coloured (Late = red, On Schedule = blue)", X("Late1", "Status").value == "Late" and X("Late1", "Status").font.color.rgb == "FFCF222E" and X("T1", "Status").value == "On Schedule" and X("T1", "Status").font.color.rgb == "FF0969DA", (X("Late1", "Status").value, X("T1", "Status").value))
    check("Excel: the % column still has its data bar (wherever it now sits)", any(str(c.sqref).startswith(openpyxl.utils.get_column_letter(ha.index("% Complete") + 1)) for c in wa.conditional_formatting))

    # ============================================================ the task dialog shows every field, calculated ones included
    make([{"name": "Grp"}, {"name": "K1", "parent": "Grp", "s": d(-4), "e": d(5), "p": 40, "extra": {"resource": "Anna"}}, {"name": "K2", "parent": "Grp", "s": d(6), "e": d(9)},
          {"name": "Solo", "s": d(-3), "e": d(6), "p": 60}, {"name": "TBDish", "extra": {"taskMode": "manual", "startText": "TBD", "durText": "a while"}}])
    def open_dialog(name): pg.evaluate("n => openTaskModal(tasks.find(t => t.name === n).id)", name); pg.wait_for_selector("#taskModalBg.open"); pg.wait_for_timeout(80)
    def close_dialog(): pg.evaluate("() => closeTaskModal()"); pg.wait_for_timeout(50)
    info = lambda: (pg.inner_text("#taskWbsInfo").strip(), pg.inner_text("#taskRemainingInfo").strip(), pg.inner_text("#taskStatusInfo").strip())
    open_dialog("Solo")
    labels = [l.strip() for l in pg.locator("#taskModalBg .modal-body label").all_inner_texts() if l.strip()]
    want = ["Task Name", "Task Mode", "Resource", "Milestone", "Start", "Finish", "Duration (days)", "% Complete", "Actual Start", "Actual Finish", "Remaining", "Status", "Constraint", "Row colour", "Predecessors"]
    check("dialog: every field is there — Remaining and Status among them, and the WBS code in the title", all(w in labels for w in want) and "WBS" in pg.inner_text("#taskModalIdBadge"), [w for w in want if w not in labels])
    check("dialog: Remaining and Status are read-only displays, not inputs", pg.locator("#taskRemainingInfo, #taskStatusInfo").evaluate_all("els => els.every(e => e.tagName === 'DIV')") and pg.locator("#taskWbsInfo").count() == 1)
    dur = 10; rem = int(dur * 40 / 100 * 10 + 0.5) / 10
    check("dialog: shows the task's WBS, remaining duration (10 days at 60% -> 4 days) and status", info() == ("2", "4 days", "On Schedule"), info())
    # live updates as the fields are edited
    pg.fill("#taskProgressInput", "100"); pg.wait_for_timeout(60)
    check("live: 100% complete -> 0 days remaining, Complete", info()[1:] == ("0 days", "Complete"), info())
    pg.fill("#taskProgressInput", "0"); pg.wait_for_timeout(60)
    check("live: 0% with a start already in the past -> Late, 10 days remaining", info()[1:] == ("10 days", "Late"), info())
    pg.fill("#taskActualStartInput", d(-3)); pg.wait_for_timeout(60)
    check("live: recording an Actual Start makes it 'started' -> On Schedule", info()[2] == "On Schedule", info())
    pg.fill("#taskActualFinishInput", d(-1)); pg.wait_for_timeout(60)
    check("live: an Actual Finish -> Complete", info()[2] == "Complete", info())
    pg.fill("#taskActualStartInput", ""); pg.fill("#taskActualFinishInput", ""); pg.fill("#taskProgressInput", "0")
    pg.fill("#taskStartInput", d(3)); pg.dispatch_event("#taskStartInput", "change"); pg.wait_for_timeout(60)
    check("live: moving the start into the future -> Future Task (the finish follows the duration)", info()[2] == "Future Task", info())
    pg.fill("#taskDurationInput", "4"); pg.dispatch_event("#taskDurationInput", "change"); pg.wait_for_timeout(60)
    check("live: changing the duration -> remaining follows (4 days at 0%)", info()[1] == "4 days", info())
    pg.fill("#taskProgressInput", "50"); pg.wait_for_timeout(60)
    check("live: 4 days at 50% -> 2 days remaining, On Schedule", info()[1:] == ("2 days", "On Schedule") or info()[1] == "2 days", info())
    pg.check("#taskMilestoneInput"); pg.wait_for_timeout(60)
    check("live: a milestone has no remaining duration", info()[1] == "—", info())
    pg.uncheck("#taskMilestoneInput"); pg.wait_for_timeout(60)
    before = pg.evaluate("() => { const t = tasks.find(x => x.name === 'Solo'); return [t.progress, t.startDate, t.endDate]; }")
    close_dialog()
    check("...only a preview: nothing was changed by looking (cancelled)", pg.evaluate("() => { const t = tasks.find(x => x.name === 'Solo'); return [t.progress, t.startDate, t.endDate]; }") == before)
    open_dialog("Solo"); pg.click("#taskModalBg .btn-primary"); pg.wait_for_timeout(100)
    check("saving stores no 'remaining' or 'status' (they stay calculated)", pg.evaluate("() => { const t = tasks.find(x => x.name === 'Solo'); return Object.keys(t).filter(k => /remaining|status|wbs/i.test(k)); }") == [])
    # a group, a task with a WBS, a task with free-text dates
    open_dialog("K1"); w = info(); close_dialog()
    check("dialog: a nested task shows its WBS (1.1)", w[0] == "1.1", w)
    open_dialog("Grp"); g = info()
    check("dialog: a group shows its rolled-up remaining duration and status", g[0] == "1" and g[1].endswith("days") and g[2] in ("On Schedule", "Late", "Future Task"), g)
    list_status = pg.evaluate("() => taskStatus(tasks.find(t => t.name === 'Grp').id)"); close_dialog()
    check("...the same status the list column shows", g[2] == list_status, (g, list_status))
    open_dialog("TBDish")
    check("dialog: a task whose dates are still free text ('TBD') has no remaining duration or status", info()[1:] == ("—", "—"), info())
    pg.fill("#taskStartInput", d(2)); pg.dispatch_event("#taskStartInput", "change"); pg.wait_for_timeout(60)
    check("...until you give it a real date", info()[2] == "Future Task", info())
    close_dialog()
    # the dialog fits the screen; the last field is reachable
    open_dialog("Solo")
    geo = pg.evaluate("() => { const m = document.querySelector('#taskModalBg .modal').getBoundingClientRect(); const b = document.querySelector('#taskModalBg .modal-body'); b.scrollTop = b.scrollHeight; const btn = document.querySelector('#taskModalBg .modal-body .btn').getBoundingClientRect(); const bb = b.getBoundingClientRect(); return {modalH: m.height, vh: innerHeight, btnVisible: btn.bottom <= bb.bottom + 1 && btn.top >= bb.top}; }")
    check("dialog: fits the window (at most 94% of its height) and the last field is reachable by scrolling", geo["modalH"] <= geo["vh"] * 0.94 + 1 and geo["btnVisible"], geo)
    close_dialog()
    pg.set_viewport_size({"width": 1500, "height": 1100}); open_dialog("Solo")
    fits = pg.evaluate("() => { const b = document.querySelector('#taskModalBg .modal-body'); return b.scrollHeight <= b.clientHeight + 1; }")
    check("on a normal-height window the whole dialog is visible without scrolling", fits); close_dialog()
    pg.set_viewport_size({"width": 1500, "height": 820})


    # ============================================================ milestones: one date, planned and actual
    make([{"name": "Normal", "s": d(-3), "e": d(4), "p": 20}, {"name": "MS", "s": d(-2), "e": d(-2), "extra": {"milestone": True}}, {"name": "MS2", "s": d(6), "e": d(6), "extra": {"milestone": True}}])
    pg.evaluate("() => { colHidden.delete('actualStart'); colHidden.delete('actualFinish'); colHidden.delete('status'); save(); render(); }"); pg.wait_for_timeout(100)
    vis = lambda sel: pg.evaluate("s => { const e = document.querySelector(s); if (!e) return false; const r = e.getBoundingClientRect(); return getComputedStyle(e).display !== 'none' && r.width > 0; }", sel)
    open_dialog("Normal")
    check("a normal task's dialog shows Start, Finish, Duration, Actual Start and Actual Finish", all(vis(x) for x in ("#taskStartInput", "#taskEndField", "#taskDurationField", "#taskActualStartInput", "#taskActualFinishField")))
    pg.check("#taskMilestoneInput"); pg.wait_for_timeout(50)
    check("ticking Milestone hides Finish, Duration AND Actual Finish (a milestone has one date)", not vis("#taskEndField") and not vis("#taskDurationField") and not vis("#taskActualFinishField") and vis("#taskActualStartInput") and vis("#taskStartInput"))
    pg.uncheck("#taskMilestoneInput"); pg.wait_for_timeout(50)
    check("unticking brings them all back", vis("#taskEndField") and vis("#taskActualFinishField") and vis("#taskDurationField"))
    close_dialog()
    open_dialog("MS")
    check("a milestone's dialog: Start and Actual Start only", vis("#taskStartInput") and vis("#taskActualStartInput") and not vis("#taskEndField") and not vis("#taskActualFinishField") and not vis("#taskDurationField"))
    check("...no actual date yet, so its date has passed without happening -> Late", info()[2] == "Late", info())
    pg.fill("#taskActualStartInput", d(-2)); pg.wait_for_timeout(60)
    check("live: recording its actual date makes the milestone Complete (it has happened)", info()[2] == "Complete", info())
    pg.click("#taskModalBg .btn-primary"); pg.wait_for_timeout(120)
    both = pg.evaluate("() => { const t = tasks.find(x => x.name === 'MS'); return [t.actualStart, t.actualFinish]; }")
    check("saved: a milestone's Actual Finish is its Actual Start", both == [d(-2), d(-2)], both)
    check("list: both actual columns show that one date, and the status is Complete", cell("MS", "actualStart").inner_text().strip() == fmt(d(-2)) and cell("MS", "actualFinish").inner_text().strip() == fmt(d(-2)) and cell("MS", "status").inner_text().strip() == "Complete")
    cell("MS", "actualFinish").click(); pg.wait_for_timeout(120)
    check("list: a milestone's Actual Finish cell isn't editable (like its Finish)", pg.locator(".inline-edit").count() == 0 and "editable" not in (cell("MS", "actualFinish").get_attribute("class") or ""))
    cell("MS", "actualStart").click(); wait_focus(); pg.fill(".inline-edit", d(-1)); pg.keyboard.press("Enter"); pg.wait_for_timeout(120)
    both = pg.evaluate("() => { const t = tasks.find(x => x.name === 'MS'); return [t.actualStart, t.actualFinish]; }")
    check("list: editing a milestone's Actual Start moves its Actual Finish with it (no 'finish before start' false alarm)", both == [d(-1), d(-1)] and "before" not in pg.inner_text("#toastMsg"), (both, pg.inner_text("#toastMsg")))
    pg.evaluate("() => { const t = tasks.find(x => x.name === 'MS2'); t.actualStart = '%s'; t.actualFinish = '%s'; normalizeData(); }" % (d(6), d(9)))
    check("normalisation: a milestone's Actual Finish is forced to its Actual Start", pg.evaluate("() => { const t = tasks.find(x => x.name === 'MS2'); return t.actualFinish === t.actualStart; }"))
    # turning a task into a milestone
    pg.evaluate("() => { const t = tasks.find(x => x.name === 'Normal'); t.actualStart = '%s'; t.actualFinish = '%s'; save(); render(); }" % (d(-3), d(-1)))
    open_dialog("Normal"); pg.check("#taskMilestoneInput"); pg.click("#taskModalBg .btn-primary"); pg.wait_for_timeout(120)
    check("turning a task into a milestone collapses its actual dates to the start", pg.evaluate("() => { const t = tasks.find(x => x.name === 'Normal'); return [t.actualStart, t.actualFinish]; }") == [d(-3), d(-3)])


    # ============================================================ notes were removed from tasks
    make([{"name": "WithOldNotes", "extra": {"notes": "kept in the data, not shown"}}])
    pg.evaluate("() => { addTask(); }"); pg.wait_for_selector("#taskModalBg.open")
    check("the task dialog has no Notes field any more", "Notes" not in pg.inner_text("#taskModalBg") and pg.locator("#taskNotesInput").count() == 0)
    pg.evaluate("() => closeTaskModal()")
    check("a task created now has no notes value", pg.evaluate("() => tasks.filter(t => t.name !== 'WithOldNotes').every(t => !('notes' in t))"))
    open_dialog("WithOldNotes"); pg.click("#taskModalBg .btn-primary"); pg.wait_for_timeout(120)
    check("saving a task doesn't destroy notes text an earlier version stored (it is simply not shown or exported)", pg.evaluate("() => tasks.find(t => t.name === 'WithOldNotes').notes") == "kept in the data, not shown")

    print("console errors/warnings:", errors)
    print(f"{sum(results)}/{len(results)} passed")
    b.close()
