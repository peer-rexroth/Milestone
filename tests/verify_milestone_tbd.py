from playwright.sync_api import sync_playwright
import os, io, base64
import openpyxl
URL = os.environ.get("MILESTONE_URL", "http://127.0.0.1:8937/milestone.html")
errors, results = [], []
def check(name, cond, detail=""):
    results.append(bool(cond)); print(("PASS  " if cond else "FAIL  ") + name + (f"   [{str(detail)[:400]}]" if not cond and detail else ""))

SEED = """() => { tasks.length = 0; deletedTaskIds.length = 0; setSelection([]); delete project.workDays; delete project.holidays;
  const mk = (n, o, extra) => Object.assign({ id: genId() + n, name: n, parentId: null, order: o, startDate: '2026-09-14', endDate: '2026-09-14', progress: 0, milestone: false, color: null, predecessors: [], collapsed: false, updatedAt: 1, constraintType: 'ASAP', constraintDate: null, taskMode: 'auto', resource: '', actualStart: null, actualFinish: null }, extra || {});
  const m = mk('M', 0, { milestone: true, taskMode: 'manual' }), a = mk('A', 1, { milestone: true });
  const s = mk('S', 2, { startDate: '2026-09-15', endDate: '2026-09-16' }); s.predecessors = [{ id: m.id, type: 'FS', lag: 0 }];
  tasks.push(m, a, s); currentView = 'tasks'; normalizeData(); save(); render(); resetHistory(); }"""

with sync_playwright() as p:
    b = p.chromium.launch(headless=True)
    ctx = b.new_context(viewport={"width": 1600, "height": 900}); ctx.add_init_script("delete window.showOpenFilePicker; delete window.showSaveFilePicker")
    pg = ctx.new_page(); pg.on("pageerror", lambda e: errors.append(str(e))); pg.on("console", lambda m: errors.append(m.text) if m.type in ("error", "warning") else None)
    pg.goto(URL); pg.wait_for_selector("#addTaskBtn"); pg.evaluate("() => localStorage.clear()"); pg.reload(); pg.wait_for_selector("#addTaskBtn")
    ev = pg.evaluate
    ev("() => { historyCoalesceMs = 0; }")
    ev(SEED)
    tid = lambda n: ev("n => tasks.find(t => t.name === n).id", n)
    get = lambda n, k: ev("([n, k]) => tasks.find(t => t.name === n)[k]", [n, k])
    cell = lambda n, field: pg.locator(f".grid-row[data-id='{tid(n)}'] [onclick*=\"'{field}'\"]")
    finish_text = lambda n: ev("n => clipCell(tasks.find(t => t.name === n), 'end', 0)", n)

    check("the Finish cell of a MANUAL milestone is editable", cell("M", "finish").count() == 1)
    check("...an Auto milestone's Finish is not (its finish is its start), and neither is the Duration of either", cell("A", "finish").count() == 0 and cell("M", "duration").count() == 0 and cell("A", "duration").count() == 0)
    cell("M", "finish").click(); pg.wait_for_timeout(150)
    check("clicking it opens the text editor with the date (dd.mm.yyyy) and the calendar button", ev("() => !!document.querySelector('.inline-wrap input[type=text]')") and pg.input_value(".inline-wrap input[type=text]") == "14.09.2026" and pg.locator(".inline-wrap button").count() == 1, pg.input_value(".inline-wrap input[type=text]") if pg.locator(".inline-wrap input[type=text]").count() else None)
    pg.fill(".inline-wrap input[type=text]", "TBD"); pg.keyboard.press("Enter"); pg.wait_for_timeout(200)
    check("typing TBD sets the milestone's Finish to that text (the Start date stays)", get("M", "endText") == "TBD" and get("M", "startText") is None and get("M", "startDate") == "2026-09-14" and get("M", "milestone") is True, (get("M", "endText"), get("M", "startText")))
    check("...the cell shows TBD", "TBD" in cell("M", "finish").inner_text(), cell("M", "finish").inner_text())
    check("...it is a milestone still, now without a full date: unscheduled (no bar), ignored as a predecessor", ev("() => { const m = tasks.find(t => t.name === 'M'); return m.milestone && isUnscheduled(m) && !effectiveDates(m.id).start && constraintStart(tasks.find(t => t.name === 'S')) === null; }"))
    check("...copy and Excel see the text too", finish_text("M") == "TBD")
    ev("() => setView('gantt')"); pg.wait_for_timeout(250)
    check("in the Gantt view it shows the bracket at the known date, as a task with one date does", pg.locator(".gantt-bracket").count() >= 1, pg.locator(".gantt-bracket").count())
    ev("() => setView('tasks')"); pg.wait_for_timeout(200)
    pg.keyboard.press("Control+z"); pg.wait_for_timeout(150)
    check("Ctrl+Z takes the TBD back", get("M", "endText") is None and get("M", "endDate") == "2026-09-14")
    cell("M", "finish").click(); pg.wait_for_timeout(150); pg.fill(".inline-wrap input[type=text]", "TBD"); pg.keyboard.press("Enter"); pg.wait_for_timeout(200)
    pg.reload(); pg.wait_for_selector("#addTaskBtn")
    check("the text survives a reload (normalizing keeps it on a manual milestone)", get("M", "endText") == "TBD" and get("M", "milestone") is True)

    # a real date afterwards
    cell("M", "finish").click(); pg.wait_for_timeout(150); pg.fill(".inline-wrap input[type=text]", "21.09.2026"); pg.keyboard.press("Enter"); pg.wait_for_timeout(200)
    check("typing a real date into the milestone's Finish makes it THE date of the milestone: Start and Finish both 21.09., the text is gone", get("M", "startDate") == "2026-09-21" and get("M", "endDate") == "2026-09-21" and get("M", "endText") is None and get("M", "startText") is None, (get("M", "startDate"), get("M", "endDate")))
    check("...and it is scheduled again (a successor now follows it: S starts the working day after)", ev("() => { const s = tasks.find(t => t.name === 'S'); return constraintStart(s) === '2026-09-22'; }"), ev("() => constraintStart(tasks.find(t => t.name === 'S'))"))
    # start TBD
    cell("M", "start").click(); pg.wait_for_timeout(150); pg.fill(".inline-wrap input[type=text]", "TBD"); pg.keyboard.press("Enter"); pg.wait_for_timeout(200)
    check("Start = TBD on a milestone was already possible; its Finish cell then shows the same TBD instead of a stale date", get("M", "startText") == "TBD" and "TBD" in cell("M", "finish").inner_text(), cell("M", "finish").inner_text())
    cell("M", "start").click(); pg.wait_for_timeout(150); pg.fill(".inline-wrap input[type=text]", "22.09.2026"); pg.keyboard.press("Enter"); pg.wait_for_timeout(200)
    check("a date in Start clears every text again (both dates 22.09.)", get("M", "startText") is None and get("M", "endText") is None and get("M", "endDate") == "2026-09-22")
    # blank
    cell("M", "finish").click(); pg.wait_for_timeout(150); pg.fill(".inline-wrap input[type=text]", ""); pg.keyboard.press("Enter"); pg.wait_for_timeout(200)
    check("an empty Finish is kept as 'not set' text like on any manual task", get("M", "endText") == "" and ev("() => isUnscheduled(tasks.find(t => t.name === 'M'))"))
    # dialog
    cell("M", "finish").click(); pg.wait_for_timeout(150); pg.fill(".inline-wrap input[type=text]", "after go-live"); pg.keyboard.press("Enter"); pg.wait_for_timeout(200)
    ev("() => openTaskModal(tasks.find(t => t.name === 'M').id)"); pg.wait_for_timeout(300)
    note = pg.inner_text("#taskTextDatesNote")
    check("the task dialog names the free text ('Finish: \"after go-live\"')", 'Finish: "after go-live"' in note, note)
    pg.click("#taskModalBg .modal-footer .btn-primary"); pg.wait_for_timeout(250)
    check("saving the dialog without touching a date keeps the text", get("M", "endText") == "after go-live")
    ev("() => openTaskModal(tasks.find(t => t.name === 'M').id)"); pg.wait_for_timeout(300)
    pg.fill("#taskStartInput", "2026-09-30"); pg.click("#taskModalBg .modal-footer .btn-primary"); pg.wait_for_timeout(250)
    check("...changing the date in the dialog replaces it (milestone on 30.09., no text)", get("M", "endText") is None and get("M", "startDate") == "2026-09-30" and get("M", "endDate") == "2026-09-30", (get("M", "endText"), get("M", "startDate")))

    # switching modes / a milestone that is not manual
    cell("M", "finish").click(); pg.wait_for_timeout(150); pg.fill(".inline-wrap input[type=text]", "TBD"); pg.keyboard.press("Enter"); pg.wait_for_timeout(200)
    ev("() => setTaskMode(tasks.find(t => t.name === 'M').id, 'auto')"); pg.wait_for_timeout(150)
    check("switching the milestone to Auto drops the free text (it needs a real date again)", get("M", "endText") is None and get("M", "startText") is None and not ev("() => isUnscheduled(tasks.find(t => t.name === 'M'))"))
    ev("() => { tasks.find(t => t.name === 'A').endText = 'TBD'; normalizeData(); }")
    check("an Auto milestone never keeps Finish text (normalizing removes it)", get("A", "endText") is None)

    # Excel: the milestone's text and the WBS column
    b64 = ev("""async () => { const blob = buildXlsx({ scope: 'all', gantt: true, columns: 'shown' }); const buf = new Uint8Array(await blob.arrayBuffer()); let s = ''; for (const c of buf) s += String.fromCharCode(c); return btoa(s); }""")
    wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(b64)))
    ws = wb["Tasks"]; heads = {ws.cell(4, c).value: c for c in range(1, ws.max_column + 1)}
    from openpyxl.utils import get_column_letter as L
    def width(sheet, col_idx):
        for key, dim in sheet.column_dimensions.items():
            if dim.min <= col_idx <= dim.max: return dim.width
    check("the Excel Tasks sheet has a WBS column, wide enough for eight levels: 15 (the app's 104px column, 7px per unit)", "WBS" in heads and abs(width(ws, heads["WBS"]) - 15) < 0.01, width(ws, heads.get("WBS", 1)))
    gws = wb["Gantt"] if "Gantt" in wb.sheetnames else None
    gh = {gws.cell(3, c).value: c for c in range(1, 12)} if gws else {}
    check("...and the Gantt sheet's list part uses the same width", gws is not None and "WBS" in gh and abs(width(gws, gh["WBS"]) - 15) < 0.01, gh)
    ev("() => { const m = tasks.find(t => t.name === 'M'); setTaskMode(m.id, 'manual'); m.endText = 'TBD'; normalizeData(); save(); render(); }")
    b64 = ev("""async () => { const blob = buildXlsx({ scope: 'all', gantt: false, columns: 'shown' }); const buf = new Uint8Array(await blob.arrayBuffer()); let s = ''; for (const c of buf) s += String.fromCharCode(c); return btoa(s); }""")
    ws = openpyxl.load_workbook(io.BytesIO(base64.b64decode(b64)))["Tasks"]; heads = {ws.cell(4, c).value: c for c in range(1, ws.max_column + 1)}
    vals = {ws.cell(r, heads["Task Name"]).value: ws.cell(r, heads["Finish"]).value for r in range(5, 5 + 3)}
    check("a milestone's Finish text 'TBD' is exported as text in the Finish column", vals.get("◆ M") == "TBD", vals)
    # a deep outline fits the column
    ev("""() => { tasks.length = 0; let parent = null; const mk = (id, par, order) => ({ id, name: id, parentId: par, order, startDate: '2026-09-07', endDate: '2026-09-08', progress: 0, milestone: false, color: null, predecessors: [], collapsed: false, updatedAt: 1, constraintType: 'ASAP', constraintDate: null, taskMode: 'auto', resource: '', actualStart: null, actualFinish: null });
      for (let d = 0; d < 8; d++) { tasks.push(mk('f' + d, parent, 0)); const n = mk('n' + d, parent, 1); tasks.push(n); parent = n.id; } normalizeData(); save(); render(); }""")
    check("(the deepest code of that outline is 2.2.2.2.2.2.2.2 = 15 characters, and the column is 15 units wide)", ev("() => wbsCode(tasks.find(t => t.id === 'n7').id)") == "2.2.2.2.2.2.2.2")
    check("no console errors", not errors, errors[:5])
    print("console errors/warnings:", errors[:5]); print(f"{sum(results)}/{len(results)} passed"); b.close()
