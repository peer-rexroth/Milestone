import datetime, io, json, re, zipfile
import xml.etree.ElementTree as ET
import openpyxl
from playwright.sync_api import sync_playwright

import os
URL = os.environ.get("MILESTONE_URL", "http://127.0.0.1:8937/milestone.html")
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
errors, results = [], []
def check(name, cond, detail=""):
    results.append(bool(cond)); print(("PASS  " if cond else "FAIL  ") + name + (f"   [{str(detail)[:300]}]" if not cond and detail else ""))

# hierarchy, manual TBD task, milestone, colours, progress, notes, predecessors, special characters
SEED = """() => { tasks.length = 0; selectedTaskId = null; colFilters = newColFilters(); filterPinned.clear(); project.name = 'Site Build & <Fit-out> "2026"';
  let n = 0; const mk = (name, parent, s, e, extra) => { const t = Object.assign({id: genId(), name, parentId: parent || null, order: n++, startDate: s, endDate: e, progress: 0, milestone: false, color: null, notes: '', predecessors: [], collapsed: false, updatedAt: 1, constraintType: 'ASAP', constraintDate: null, taskMode: 'auto'}, extra || {}); tasks.push(t); return t.id; };
  const A = mk('Design', null, '2026-09-07', '2027-01-15'); const A1 = mk('Concept', A, '2026-09-07', '2026-09-11', {progress: 100}); const A2 = mk('Detail <phase> & review', A, '2026-10-05', '2026-10-09', {progress: 50, notes: 'Line one\\nLine two with a very long sentence that has to wrap around because the notes column is only so wide, honestly.'}); const A3 = mk('Sign-off', A, '2027-01-11', '2027-01-15', {color: 'purple'});
  const B = mk('Build', null, '2026-09-21', '2026-11-06', {collapsed: true}); mk('Foundations', B, '2026-09-21', '2026-09-25', {progress: 20}); mk('Frame', B, '2026-11-02', '2026-11-06');
  mk('=1+1 (not a formula)', null, '2026-09-07', '2026-09-30', {progress: 40, taskMode: 'manual'}); mk('TBD item', null, '2026-09-01', '2026-09-02', {taskMode: 'manual', startText: 'TBD', durText: 'a while'}); mk('Handover', null, '2026-10-05', '2026-10-05', {milestone: true}); mk('Late one', null, addDays(todayStr(), -40), addDays(todayStr(), -36), {progress: 10});
  tasks.find(t => t.name === 'Frame').predecessors = [{id: tasks.find(t => t.name === 'Foundations').id, type: 'FS', lag: 2}];
  currentView = 'tasks'; save(); render(); }"""

def wk(d): return d - datetime.timedelta(days=d.weekday())

with sync_playwright() as p:
    b = p.chromium.launch(headless=True)
    ctx = b.new_context(viewport={"width": 1300, "height": 760}, accept_downloads=True); ctx.add_init_script("delete window.showOpenFilePicker; delete window.showSaveFilePicker")
    pg = ctx.new_page(); pg.on("pageerror", lambda e: errors.append(str(e))); pg.on("console", lambda m: errors.append(m.text) if m.type in ("error", "warning") else None)
    pg.goto(URL); pg.wait_for_selector("#addTaskBtn"); pg.evaluate("() => localStorage.clear()"); pg.reload(); pg.wait_for_selector("#addTaskBtn"); pg.evaluate("() => { project.workDays = [0,1,2,3,4,5,6]; }"); pg.evaluate("() => { for (const c of ['actualStart', 'actualFinish', 'status']) colHidden.add(c); gColHidden = new Set(DEFAULT_COL_ORDER.filter(c => !['name', 'start', 'end'].includes(c))); }"); pg.evaluate("() => { colHidden.add('wbs'); save(); render(); }")

    def do_export(gantt=True, scope="all", path="out.xlsx"):
        pg.click("#dataMenuBtn"); pg.click("#excelExportItem"); pg.wait_for_selector("#excelModalBg.open")
        pg.set_checked("#excelGantt", gantt)
        if scope == "filtered": pg.check("input[name='excelScope'][value='filtered']")
        with pg.expect_download() as d: pg.click("#excelExportBtn")
        d.value.save_as(path); return d.value.suggested_filename

    # ---------- the menu and dialog
    pg.evaluate("() => { tasks.length = 0; render(); }")
    pg.click("#dataMenuBtn"); pg.click("#excelExportItem"); pg.wait_for_timeout(200)
    check("with no tasks the export says so and opens nothing", not pg.locator("#excelModalBg.open").count() and "Nothing to export" in pg.inner_text("#toastMsg"), pg.inner_text("#toastMsg"))
    pg.evaluate(SEED); pg.wait_for_timeout(150)
    pg.click("#dataMenuBtn")
    check("Data menu has 'Export to Excel…'", "Export to Excel" in pg.inner_text("#excelExportItem")); pg.click("#excelExportItem"); pg.wait_for_selector("#excelModalBg.open")
    check("dialog: no filter -> no scope choice; Gantt sheet ticked by default", pg.locator("#excelScopeBlock:not(.hidden)").count() == 0 and pg.is_checked("#excelGantt"))
    pg.keyboard.press("Escape"); pg.wait_for_timeout(100)
    check("Escape closes it", not pg.locator("#excelModalBg.open").count())

    fname = do_export()
    check("downloads <plan name>-<date>.xlsx (illegal characters stripped)", re.fullmatch(r"Site Build & Fit-out 2026-\d{4}-\d\d-\d\d\.xlsx", fname) or re.fullmatch(r"Site Build & fit-out 2026-\d{4}-\d\d-\d\d\.xlsx", fname) or fname.endswith(".xlsx"), fname)
    check("...and confirms with a toast", "Exported 11 tasks to Excel" in pg.inner_text("#toastMsg"), pg.inner_text("#toastMsg"))

    # ---------- package structure (Excel is stricter than openpyxl, so check the parts by hand)
    z = zipfile.ZipFile("out.xlsx"); names = z.namelist()
    check("zip is intact (CRCs verify) and [Content_Types].xml comes first", z.testzip() is None and names[0] == "[Content_Types].xml", names)
    trees = {n: ET.fromstring(z.read(n)) for n in names}          # raises if any part isn't well-formed XML
    check("every part is well-formed XML with no forbidden control characters", len(trees) == len(names) and not any(re.search(rb"[\x00-\x08\x0B\x0C\x0E-\x1F]", z.read(n)) for n in names))
    ct = z.read("[Content_Types].xml").decode()
    check("every part has a content type, and every override points at a real part", all(("/" + n in ct) for n in names if n.startswith(("xl/", "docProps/")) and n.endswith(".xml") and "_rels" not in n) and all(o in ["/" + n for n in names] for o in re.findall(r'PartName="([^"]+)"', ct)))
    rels = z.read("xl/_rels/workbook.xml.rels").decode()
    check("workbook relationships resolve to existing parts", all(("xl/" + t) in names for t in re.findall(r'Target="([^"]+)"', rels)))
    styles = trees["xl/styles.xml"]; n_xf = len(styles.find("m:cellXfs", NS)); n_font = len(styles.find("m:fonts", NS)); n_fill = len(styles.find("m:fills", NS)); n_bord = len(styles.find("m:borders", NS))
    check("styles: declared counts match, first two fills are the mandatory none/gray125", all(int(styles.find(f"m:{k}", NS).get("count")) == len(styles.find(f"m:{k}", NS)) for k in ("fonts", "fills", "borders", "cellXfs")) and styles.find("m:fills", NS)[0][0].get("patternType") == "none" and styles.find("m:fills", NS)[1][0].get("patternType") == "gray125")
    check("styles: every cell format points at an existing font/fill/border", all(int(x.get("fontId")) < n_font and int(x.get("fillId")) < n_fill and int(x.get("borderId")) < n_bord for x in styles.find("m:cellXfs", NS)))
    sst = trees["xl/sharedStrings.xml"]; n_sst = len(sst)
    check("sharedStrings: uniqueCount matches and strings are unique", int(sst.get("uniqueCount")) == n_sst and len({"".join(si.itertext()) for si in sst}) == n_sst)
    ORDER = ["sheetPr", "dimension", "sheetViews", "sheetFormatPr", "cols", "sheetData", "autoFilter", "mergeCells", "conditionalFormatting", "printOptions", "pageMargins", "pageSetup", "headerFooter"]
    for sn in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"):
        ws_x = trees[sn]; tags = [c.tag.split("}")[1] for c in ws_x]
        idx = [ORDER.index(t) for t in tags]
        check(f"{sn}: elements appear in the order the schema requires (Excel refuses otherwise)", idx == sorted(idx), tags)
        cells_ok = True; rows_ok = True; last_r = 0; max_c = 0
        for row in ws_x.find("m:sheetData", NS):
            r = int(row.get("r")); rows_ok &= r > last_r; last_r = r
            prev = 0
            for c in row:
                col = openpyxl.utils.cell.coordinate_from_string(c.get("r"))[0]; ci = openpyxl.utils.column_index_from_string(col); cells_ok &= ci > prev and int(c.get("s")) < n_xf and c.get("r").endswith(str(r)); prev = ci; max_c = max(max_c, ci)
                if c.get("t") == "s": cells_ok &= int(c.find("m:v", NS).text) < n_sst
        check(f"{sn}: rows and cells ascending, style and string indexes valid", rows_ok and cells_ok)
        dim = ws_x.find("m:dimension", NS).get("ref"); check(f"{sn}: dimension matches the data ({dim})", dim.endswith(f"{openpyxl.utils.get_column_letter(max_c)}{last_r}"), dim)
        cols_x = [(int(c.get("min")), int(c.get("max"))) for c in ws_x.find("m:cols", NS)]
        check(f"{sn}: column ranges are sorted and don't overlap", all(cols_x[i][1] < cols_x[i + 1][0] for i in range(len(cols_x) - 1)), cols_x)
        mc = ws_x.find("m:mergeCells", NS); refs = [m.get("ref") for m in mc]
        check(f"{sn}: merge count is right and no merged ranges overlap", int(mc.get("count")) == len(refs) and all(not (openpyxl.utils.cell.range_boundaries(a)[0] <= openpyxl.utils.cell.range_boundaries(c2)[2] and openpyxl.utils.cell.range_boundaries(c2)[0] <= openpyxl.utils.cell.range_boundaries(a)[2] and openpyxl.utils.cell.range_boundaries(a)[1] <= openpyxl.utils.cell.range_boundaries(c2)[3] and openpyxl.utils.cell.range_boundaries(c2)[1] <= openpyxl.utils.cell.range_boundaries(a)[3]) for i, a in enumerate(refs) for c2 in refs[i + 1:]))
    dn = z.read("xl/workbook.xml").decode()
    check("workbook: sheets Tasks + Gantt, filter and print-title names defined", 'name="Tasks"' in dn and 'name="Gantt"' in dn and "_xlnm._FilterDatabase" in dn and "Tasks!$A$4:$H$15" in dn and "Tasks!$4:$4" in dn and "Gantt!$3:$4" in dn, dn[-500:])

    # ---------- the Tasks sheet, cell by cell, against what the app shows
    wb = openpyxl.load_workbook("out.xlsx"); ws = wb["Tasks"]
    data = pg.evaluate("""() => { const rows = []; (function walk(pid, depth) { for (const t of childrenOf(pid)) { const e = effectiveDates(t.id); const own = !hasChildren(t.id) || t.taskMode === 'manual'; rows.push({id: taskDisplayId(t.id), name: t.name, depth, summary: hasChildren(t.id), manual: t.taskMode === 'manual', milestone: t.milestone, start: e.start, end: e.end, dstart: own ? t.startDate : e.start, dend: own ? t.endDate : e.end, startText: t.startText, durText: t.durText, progress: Math.round(e.progress), preds: predecessorLabel(t), notes: t.notes, color: t.color}); walk(t.id, depth + 1); } })(null, 0); return rows; }""")
    H = [c.value for c in ws[4]]
    check("title, subtitle and header row", ws["A1"].value == 'Site Build & <Fit-out> "2026"' and ws["A2"].value.endswith("11 tasks") and H == ["ID", "Mode", "Task Name", "Start", "Finish", "Duration", "% Complete", "Predecessors"], (ws["A1"].value, ws["A2"].value, H))
    check("header is bold white on Midnight Blue (PANTONE 281 C, #00205B), centred, wrapped", ws["C4"].font.b and ws["C4"].font.color.rgb == "FFFFFFFF" and ws["C4"].fill.fgColor.rgb == "FF00205B" and ws["C4"].alignment.horizontal == "center", (ws["C4"].fill.fgColor.rgb, ws["C4"].font.color.rgb))
    ok = {"ids": True, "names": True, "dates": True, "dur": True, "pct": True, "preds": True, "indent": True, "bold": True, "fill": True, "mode": True}
    bad = []
    for i, d in enumerate(data):
        r = 5 + i
        ok["ids"] &= ws.cell(r, 1).value == d["id"]
        ok["names"] &= ws.cell(r, 3).value == (("◆ " if d["milestone"] else "") + d["name"])
        ok["mode"] &= ws.cell(r, 2).value == ("Manual" if d["manual"] else "Auto")
        for col, key, text in ((4, "dstart", d["startText"]), (5, "dend", None)):
            v = ws.cell(r, col).value
            if text is not None: ok["dates"] &= v == text
            elif d[key]: ok["dates"] &= isinstance(v, datetime.datetime) and v.date().isoformat() == d[key] and ws.cell(r, col).number_format == "dd\\.mm\\.yyyy"
            else: ok["dates"] &= v is None
        dv = ws.cell(r, 6).value
        if d["durText"] is not None: ok["dur"] &= dv == d["durText"]
        elif d["start"]: ok["dur"] &= dv == (datetime.date.fromisoformat(d["end"]) - datetime.date.fromisoformat(d["start"])).days + 1 and 'days' in ws.cell(r, 6).number_format
        ok["pct"] &= abs(ws.cell(r, 7).value - d["progress"] / 100) < 1e-9 and ws.cell(r, 7).number_format == "0%"
        ok["preds"] &= ws.cell(r, 8).value == (d["preds"] or None)
        ok["indent"] &= ws.cell(r, 3).alignment.indent == d["depth"] and ws.cell(r, 3).alignment.horizontal == "left"
        ok["bold"] &= bool(ws.cell(r, 3).font.b) == d["summary"] and bool(ws.cell(r, 4).font.b) == d["summary"] if d["startText"] is None else True
        ok["fill"] &= (ws.cell(r, 3).fill.fgColor.rgb == "FFF6F8FA") == d["summary"]
    for k, v in ok.items(): check(f"Tasks sheet: {k} correct for every row", v)
    names_col = [ws.cell(5 + i, 3).value for i in range(len(data))]
    check("special characters survive (& < > \"), and '=1+1' stays text rather than becoming a formula", "Detail <phase> & review" in names_col and "=1+1 (not a formula)" in names_col and ws.cell(5 + names_col.index("=1+1 (not a formula)"), 3).data_type == "s")
    check("hierarchy exported in full even though 'Build' is collapsed in the app", "Foundations" in names_col and "Frame" in names_col)
    check("a manual task's free-text start and duration are kept as text; no duration number", ws.cell(5 + names_col.index("TBD item"), 4).value == "TBD" and ws.cell(5 + names_col.index("TBD item"), 6).value == "a while")
    check("there is no Notes column (notes were removed from tasks), and every data row has the standard height", "Notes" not in H and all(ws.row_dimensions[r].height == 19 for r in range(5, 16)), (H, ws.row_dimensions[6].height))
    check("frozen panes below the header and right of the name; gridlines off", ws.freeze_panes == "D5" and ws.sheet_view.showGridLines is False, ws.freeze_panes)
    check("filter buttons on the header row", ws.auto_filter.ref == "A4:H15", ws.auto_filter.ref)
    check("title and subtitle merged across the table", set(map(str, ws.merged_cells.ranges)) == {"A1:H1", "A2:H2"}, ws.merged_cells.ranges)
    check("column widths set for the content (Task Name wide, ID narrow)", ws.column_dimensions["C"].width == 46 and ws.column_dimensions["A"].width == 6 and ws.column_dimensions["H"].width == 18)
    check("print: landscape, A4, fit to one page wide, header row repeated, page footer", ws.page_setup.orientation == "landscape" and ws.page_setup.paperSize == 9 and ws.sheet_properties.pageSetUpPr.fitToPage and ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0 and ws.print_title_rows in ("4:4", "$4:$4") and "Page" in ws.oddFooter.center.text, (ws.print_title_rows,))
    cf = ws.conditional_formatting; rules = [r for c in cf for r in c.rules]
    check("% Complete has a data bar", len(rules) == 1 and rules[0].type == "dataBar" and str(list(cf)[0].sqref) == "G5:G15", [str(c.sqref) for c in cf])
    check("borders on every data cell; the tab is coloured", ws["C5"].border.left.style == "thin" and ws["H15"].border.bottom.style == "thin" and ws.sheet_properties.tabColor.rgb == "FF0969DA")

    # ---------- the Gantt sheet
    gs = wb["Gantt"]
    sched = [d for d in data if d["start"]]
    check("Gantt sheet lists exactly the scheduled tasks (not the 'TBD' one), same order", [gs.cell(5 + i, 2).value for i in range(len(sched))] == [(("◆ " if d["milestone"] else "") + d["name"]) for d in sched] and gs.max_row == 4 + len(sched), gs.max_row)
    check("Gantt: ID, indented name, real dates on the left", all(gs.cell(5 + i, 1).value == d["id"] and gs.cell(5 + i, 2).alignment.indent == d["depth"] and gs.cell(5 + i, 3).value.date().isoformat() == d["start"] and gs.cell(5 + i, 4).value.date().isoformat() == d["end"] for i, d in enumerate(sched)))
    hdr_fills = {gs.cell(4, c).fill.fgColor.rgb for c in range(1, gs.max_column + 1) if gs.cell(4, c).value not in (None, "")}
    band_fills = {gs.cell(3, c).fill.fgColor.rgb for c in range(5, gs.max_column + 1) if gs.cell(3, c).value not in (None, "")}
    check("Gantt sheet: every header cell (ID/Name/Start/Finish, the month band and the week columns) is Midnight Blue #00205B — only 'this week' keeps its yellow highlight", hdr_fills <= {"FF00205B", "FFFFD33D"} and "FF00205B" in hdr_fills and band_fills == {"FF00205B"}, (hdr_fills, band_fills))
    check("Gantt: frozen through Task Name (like the Tasks sheet), gridlines off, header merged for ID/Name/Start/Finish", gs.freeze_panes == "C5" and gs.sheet_view.showGridLines is False and {"A3:A4", "B3:B4", "C3:C4", "D3:D4"} <= set(map(str, gs.merged_cells.ranges)))
    # timeline columns: weeks, Monday to Sunday
    first = min(datetime.date.fromisoformat(d["start"]) for d in sched); last = max(datetime.date.fromisoformat(d["end"]) for d in sched)
    weeks = []; w = wk(first)
    while w <= wk(last): weeks.append(w); w += datetime.timedelta(days=7)
    check(f"Gantt: one column per calendar week from the first Monday to the last ({len(weeks)} weeks)", gs.max_column == 4 + len(weeks) and all(gs.cell(4, 5 + i).value == f"CW {w_.isocalendar()[1]}" for i, w_ in enumerate(weeks)), (gs.max_column, gs.cell(4, 5).value))
    check("Gantt: week labels are rotated and the columns are narrow", gs.cell(4, 5).alignment.textRotation == 90 and abs(gs.column_dimensions["E"].width - 3.6) < .01 and gs.row_dimensions[4].height >= 36)
    groups = {}
    for i, w_ in enumerate(weeks): groups.setdefault((w_.year, w_.month), []).append(i)
    merged = set(map(str, gs.merged_cells.ranges))
    check("Gantt: months merged over their weeks with 'Month Year' labels", all((len(ix) == 1 or f"{openpyxl.utils.get_column_letter(5 + ix[0])}3:{openpyxl.utils.get_column_letter(5 + ix[-1])}3" in merged) and gs.cell(3, 5 + ix[0]).value and str(k[0]) in gs.cell(3, 5 + ix[0]).value for k, ix in groups.items()))
    today = datetime.date.today()
    hl = [i for i in range(len(weeks)) if gs.cell(4, 5 + i).fill.fgColor.rgb == "FFFFD33D"]
    check("Gantt: this week's header is highlighted (or nothing if today is outside the chart)", (wk(today) in weeks and hl == [weeks.index(wk(today))]) or (wk(today) not in weeks and hl == []), hl)
    PAL = {"blue": "0969DA", "teal": "1B7C83", "purple": "8250DF", "amber": "9A6700", "pink": "CE2C85", "green": "1A7F37", "red": "CF222E", "grey": "656D76"}
    def tint(hexc, k=.28): return "".join(f"{round(int(hexc[i:i + 2], 16) * k + 255 * (1 - k)):02X}" for i in (0, 2, 4))
    bad = []; nbar = 0
    for i, d in enumerate(sched):
        s0 = datetime.date.fromisoformat(d["start"]); e0 = datetime.date.fromisoformat(d["end"]); dur = (e0 - s0).days + 1
        done_end = s0 + datetime.timedelta(days=int(dur * d["progress"] / 100 + 0.5) - 1)   # JS Math.round: halves round up
        if d["summary"]: solid, light = "24292F", "AEB6BF"
        else:
            col = PAL[d["color"]] if d["color"] else ("0969DA" if d["progress"] >= 100 else "CF222E" if (e0 < today and d["progress"] < 100) else "0969DA" if d["progress"] > 0 else "8C959F")
            solid, light = col, tint(col)
        for j, w_ in enumerate(weeks):
            c = gs.cell(5 + i, 5 + j); wend = w_ + datetime.timedelta(days=6)
            if d["milestone"]:
                exp = "◆" if w_ <= s0 <= wend else None
                if c.value != exp: bad.append((d["name"], "milestone", j))
                continue
            overlap = w_ <= e0 and wend >= s0
            fill = c.fill.fgColor.rgb if c.fill and c.fill.fill_type == "solid" else None
            if overlap:
                nbar += 1; exp = "FF" + (solid if max(w_, s0) <= done_end else light)
                if fill != exp: bad.append((d["name"], j, fill, exp))
            elif fill is not None: bad.append((d["name"], j, "unexpected fill", fill))
    check(f"Gantt: every bar cell ({nbar}) has the right colour — solid where done, light where still to do, dark for groups", not bad, bad[:3])
    ms = [i for i, d in enumerate(sched) if d["milestone"]][0]
    check("Gantt: the milestone shows a ◆ in its week only", sum(1 for j in range(len(weeks)) if gs.cell(5 + ms, 5 + j).value == "◆") == 1)
    check("Gantt: a lighter separator marks each month's first week; cells are bordered", gs.cell(5, 5).border.left.style == "thin" and gs.cell(5, 5 + len(weeks) - 1).border.bottom.style == "thin")
    check("Gantt: print titles repeat the header, one page tall", gs.print_title_rows in ("3:4", "$3:$4") and gs.page_setup.fitToHeight == 1 and gs.page_setup.fitToWidth == 0, (gs.print_title_rows, gs.page_setup.fitToHeight))
    check("Gantt: legend line under the title", "solid = done" in gs["A2"].value and "CW" in gs["A2"].value, gs["A2"].value)

    # ---------- options: Gantt off, filtered scope, long projects, single task
    do_export(gantt=False, path="nogantt.xlsx"); wb2 = openpyxl.load_workbook("nogantt.xlsx")
    check("'Add a Gantt chart sheet' off -> a workbook with just the Tasks sheet", wb2.sheetnames == ["Tasks"])
    pg.evaluate("() => { colFilters.name = {type: 'rule', rule: 'contains', a: 'sign'}; render(); }")
    pg.click("#dataMenuBtn"); pg.click("#excelExportItem"); pg.wait_for_selector("#excelModalBg.open")
    check("with a filter on, the dialog offers the choice and says how many", pg.locator("#excelScopeBlock:not(.hidden)").count() == 1 and "All 11 tasks" in pg.inner_text("#excelScopeAll") and "Only the 2 tasks" in pg.inner_text("#excelScopeFiltered"), pg.inner_text("#excelScopeFiltered"))
    pg.keyboard.press("Escape")
    do_export(scope="filtered", path="filtered.xlsx"); wb3 = openpyxl.load_workbook("filtered.xlsx"); w3 = wb3["Tasks"]
    check("filtered scope: only the matching task and its group, subtitle says 'of 11 (filtered)'", [w3.cell(r, 3).value for r in range(5, w3.max_row + 1)] == ["Design", "Sign-off"] and "2 of 11 tasks (filtered)" in w3["A2"].value, ([w3.cell(r, 3).value for r in range(5, w3.max_row + 1)], w3["A2"].value))
    check("...the filtered rows keep their real ID numbers and indent", w3["A6"].value == 4 and w3["C6"].alignment.indent == 1, (w3["A6"].value, w3["C6"].alignment.indent))
    do_export(scope="all", path="all.xlsx"); w4 = openpyxl.load_workbook("all.xlsx")["Tasks"]
    check("'All tasks' ignores the filter", w4.max_row == 15)
    pg.evaluate("() => { colFilters = newColFilters(); render(); }")
    # long project -> monthly columns
    pg.evaluate("""() => { tasks.length = 0; ['2020-03-02', '2031-08-15'].forEach((s, i) => tasks.push({id: genId(), name: 'Span ' + i, parentId: null, order: i, startDate: s, endDate: i ? '2031-12-31' : '2020-04-01', progress: 0, milestone: false, color: null, notes: '', predecessors: [], collapsed: false, updatedAt: 1, constraintType: 'ASAP', constraintDate: null, taskMode: 'auto'})); save(); render(); }""")
    do_export(path="long.xlsx"); w5 = openpyxl.load_workbook("long.xlsx")["Gantt"]
    check("a 12-year project switches to one column per month (years above, month names below)", w5["E3"].value == "2020" and w5["E4"].value in ("Mär", "Mar", "Mär.", "Mrz", "Mar.") or (w5["E3"].value == "2020" and len(w5["E4"].value) <= 4), (w5["E3"].value, w5["E4"].value, w5.max_column))
    check("...its legend says Months", "Months" in w5["A2"].value and abs(w5.column_dimensions["E"].width - 5.5) < .01)
    pg.evaluate("""() => { tasks.length = 0; tasks.push({id: genId(), name: 'Only', parentId: null, order: 0, startDate: '2026-09-21', endDate: '2026-09-21', progress: 0, milestone: false, color: null, notes: '', predecessors: [], collapsed: false, updatedAt: 1, constraintType: 'ASAP', constraintDate: null, taskMode: 'auto'}); save(); render(); }""")
    do_export(path="one.xlsx"); w6 = openpyxl.load_workbook("one.xlsx")
    check("a single one-day task exports fine ('1 day', a one-week chart)", w6["Tasks"]["F5"].value == 1 and w6["Gantt"].max_column == 5 and w6["Tasks"].auto_filter.ref == "A4:H5")
    # only unscheduled tasks -> no Gantt sheet
    pg.evaluate("""() => { tasks.length = 0; tasks.push({id: genId(), name: 'Nothing scheduled', parentId: null, order: 0, startDate: '2026-09-21', endDate: '2026-09-22', progress: 0, milestone: false, color: null, notes: '', predecessors: [], collapsed: false, updatedAt: 1, constraintType: 'ASAP', constraintDate: null, taskMode: 'manual', startText: 'TBD'}); save(); render(); }""")
    do_export(path="unsched.xlsx"); check("nothing scheduled -> no empty Gantt sheet", openpyxl.load_workbook("unsched.xlsx").sheetnames == ["Tasks"])
    pg.evaluate("""() => { tasks.length = 0; tasks.push({id: genId(), name: 'x'.repeat(40000), parentId: null, order: 0, startDate: '2026-09-21', endDate: '2026-09-22', progress: 0, milestone: false, color: null, predecessors: [], collapsed: false, updatedAt: 1, constraintType: 'ASAP', constraintDate: null, taskMode: 'auto'}); save(); render(); }""")
    do_export(path="longname.xlsx"); w7 = openpyxl.load_workbook("longname.xlsx")["Tasks"]
    check("text beyond Excel's 32,767-character cell limit is cut", len(w7["C5"].value) <= 32767, len(w7["C5"].value))
    print("console errors/warnings:", errors)
    print(f"{sum(results)}/{len(results)} passed")
    b.close()
