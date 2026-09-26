# -*- coding: utf-8 -*-
"""Dialog UX round: (1) minute mode finished in the dialogs — Remaining Duration in the plan's unit (dialog, list, copied rows,
Excel), Actual Start/Finish time boxes blank and disabled until a date is recorded, Reschedule's time + an exact "would move"
preview (a dry run that leaves every task as it was) + an undo that restores times, Edit-tasks "Move dates by" in working time,
the import preview showing times; (2) the Scheduling precision dialog: no intro clutter, time boxes wide enough for a 12-hour
"08:00 AM", the per-weekday section folded away until needed, one break format; (3) the task dialog's predecessor rows: a heading,
the predecessor's NAME, labels; (4) accessibility: every dialog is a named modal dialog, close buttons have names, Tab stays
inside the top dialog (a confirmation above the task dialog too), focus returns to what opened it; (5) polish: what a JSON import
file contains, what a backup contains, the Baselines wording, the folder dialog's points, one 'mid' dialog width."""
import os, re
from playwright.sync_api import sync_playwright
import openpyxl, tempfile
URL = os.environ.get("MILESTONE_URL", "http://127.0.0.1:8937/milestone.html")
errors, results = [], []
def check(name, cond, detail=""):
    results.append(bool(cond)); print(("PASS  " if cond else "FAIL  ") + name + (f"   [{str(detail)[:400]}]" if not cond and detail else ""))
SEED = re.search(r'SEED = """(.*?)"""', open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'verify_clone.py')).read(), re.S).group(1)
TMP = tempfile.mkdtemp()

with sync_playwright() as p:
    b = p.chromium.launch(headless=True)
    ctx = b.new_context(viewport={"width": 1400, "height": 860}, accept_downloads=True); ctx.add_init_script("delete window.showOpenFilePicker; delete window.showSaveFilePicker; delete window.showDirectoryPicker")
    pg = ctx.new_page(); pg.on("pageerror", lambda e: errors.append(str(e))); pg.on("console", lambda m: errors.append(m.text) if m.type in ("error", "warning") else None)
    pg.goto(URL); pg.wait_for_selector("#addTaskBtn"); pg.evaluate("() => localStorage.clear()"); pg.reload(); pg.wait_for_selector("#addTaskBtn")
    ev = pg.evaluate
    tid = lambda n: ev("n => tasks.find(t => t.name === n).id", n)
    task = lambda n: ev("n => { const t = tasks.find(x => x.name === n); return [t.startDate, t.startTime, t.endDate, t.endTime]; }", n)
    toast = lambda: ev("() => document.getElementById('toastMsg').textContent")
    def plan(specs, minute=True):
        ev("m => { if (m) { project.timeUnit = 'minute'; normalizeData(); } else { delete project.timeUnit; delete project.workHours; delete project.lagUnit; delete project.workHoursByDay; }; setView('tasks'); }", minute)
        ev(SEED, specs)
        ev("() => { normalizeData(); save(); render(); }")
    def close_all():
        ev("() => { document.querySelectorAll('.modal-bg.open').forEach(m => m.classList.remove('open')); }")
    A = {"name": "A", "s": "2026-09-07", "e": "2026-09-07", "extra": {"startTime": "09:00", "endTime": "11:30"}}
    B = {"name": "B", "s": "2026-09-08", "e": "2026-09-09", "preds": [["A", "FS", 0]]}

    # ================================================================ 1. minute mode in the dialogs
    plan([A, B])
    ev("() => { colHidden.delete('remaining'); render(); }")
    def remaining_col(name):
        idx = ev("() => [...document.querySelectorAll('#gridHeader .col-filter-btn')].map(b => b.dataset.col).indexOf('remaining')")
        return pg.locator(f".grid-row[data-id='{tid(name)}']").locator(":scope > div").nth(idx + 1).inner_text()
    check("the list's Remaining column reads in working time in a minute-mode plan: A (2.5 hrs, 0%) = '2.5 hrs', B (two working days) = '2 days'", remaining_col("A") == "2.5 hrs" and remaining_col("B") == "2 days", (remaining_col("A"), remaining_col("B")))
    ev("() => { tasks.find(t => t.name === 'A').progress = 40; save(); render(); }")
    check("...40% done leaves 1.5 hrs", remaining_col("A") == "1.5 hrs", remaining_col("A"))
    ev("() => { openTaskModal(tasks.find(t => t.name === 'A').id); }"); pg.wait_for_timeout(250)
    check("the task dialog's Remaining follows (1.5 hrs, not '1.5 days' for a 2.5-hour task)", pg.inner_text("#taskRemainingInfo") == "1.5 hrs", pg.inner_text("#taskRemainingInfo"))
    pg.fill("#taskProgressInput", "100"); pg.dispatch_event("#taskProgressInput", "input"); pg.dispatch_event("#taskProgressInput", "change")
    check("...and live as % complete is typed (100% = 0 min)", pg.inner_text("#taskRemainingInfo") == "0 min", pg.inner_text("#taskRemainingInfo"))
    pg.keyboard.press("Escape"); close_all()
    ev("() => { setSelection(tasks.map(t => t.id)); }")
    check("copied rows carry it", "1.5 hrs" in ev("() => buildClip().tsv"))
    pg.click("#dataMenuBtn"); pg.click("#excelExportItem"); pg.wait_for_selector("#excelModalBg.open")
    pg.check("input[name='excelCols'][value='all']")
    with pg.expect_download() as d: pg.click("#excelExportBtn")
    path = os.path.join(TMP, "rem.xlsx"); d.value.save_as(path)
    ws = openpyxl.load_workbook(path)["Tasks"]; hdr = {c.value: c.column for c in ws[4]}
    check("Excel's Remaining cell is minute-unit text ('1.5 hrs')", ws.cell(row=5, column=hdr["Remaining Duration"]).value == "1.5 hrs", ws.cell(row=5, column=hdr["Remaining Duration"]).value)
    plan([{"name": "G", "s": "2026-09-07", "e": "2026-09-08"}, dict(A, parent="G"), dict(B, parent="G")])
    ev("() => { colHidden.delete('remaining'); render(); }")
    check("a group's Remaining is worked out in minutes too, from its rolled-up span (Mon 09:00 to Wed 17:00 = 23 working hours)", remaining_col("G") == "23 hrs", remaining_col("G"))
    plan([{"name": "A", "s": "2026-09-07", "e": "2026-09-09"}], minute=False)
    ev("() => { colHidden.delete('remaining'); render(); }")
    check("a day-mode plan is unchanged: '3 days'", remaining_col("A") == "3 days", remaining_col("A"))

    # ---- Actual Start / Finish time boxes
    plan([A, B])
    ev("() => openTaskModal(tasks.find(t => t.name === 'A').id)"); pg.wait_for_timeout(250)
    check("with no actual dates the two time boxes are empty and disabled", ev("() => document.getElementById('taskActualStartTimeInput').value") == "" and ev("() => document.getElementById('taskActualFinishTimeInput').value") == "" and pg.is_disabled("#taskActualStartTimeInput") and pg.is_disabled("#taskActualFinishTimeInput"))
    pg.fill("#taskActualStartInput", "2026-09-07"); pg.dispatch_event("#taskActualStartInput", "change")
    check("recording an Actual Start enables its time and fills the planned start time (09:00)", pg.is_enabled("#taskActualStartTimeInput") and ev("() => document.getElementById('taskActualStartTimeInput').value") == "09:00" and pg.is_disabled("#taskActualFinishTimeInput"), (ev("() => document.getElementById('taskActualStartTimeInput').value"), pg.is_enabled("#taskActualStartTimeInput")))
    pg.fill("#taskActualStartInput", ""); pg.dispatch_event("#taskActualStartInput", "change")
    check("clearing the date clears and disables the time again", ev("() => document.getElementById('taskActualStartTimeInput').value") == "" and pg.is_disabled("#taskActualStartTimeInput"))
    pg.click("#taskModalBg .modal-footer button.btn-primary"); pg.wait_for_timeout(200)
    check("saving with no actual date stores no actual time", ev("() => { const t = tasks.find(x => x.name === 'A'); return [t.actualStart, t.actualStartTime, t.actualFinish, t.actualFinishTime]; }") == [None, None, None, None])
    ev("() => openTaskModal(tasks.find(t => t.name === 'A').id)"); pg.wait_for_timeout(250)
    check("the dialog says what an actual date does", "Recording an Actual" in pg.inner_text("#taskActualsNote") and pg.is_visible("#taskActualsNote"))
    pg.keyboard.press("Escape"); close_all()
    plan([{"name": "G", "s": "2026-09-07", "e": "2026-09-08"}, dict(A, parent="G")])
    ev("() => openTaskModal(tasks.find(t => t.name === 'G').id)"); pg.wait_for_timeout(250)
    check("a group's actual time boxes stay disabled", pg.is_disabled("#taskActualStartTimeInput") and pg.is_disabled("#taskActualFinishTimeInput"))
    pg.keyboard.press("Escape"); close_all()

    # ---- Reschedule
    plan([{"name": "Late", "s": "2026-09-01", "e": "2026-09-02", "extra": {"startTime": "09:00", "endTime": "11:00"}}, {"name": "Done", "s": "2026-09-01", "e": "2026-09-01", "extra": {"progress": 100}}])
    before = ev("() => JSON.stringify(tasks)")
    ev("() => openRescheduleModal()"); pg.wait_for_timeout(200)
    now = ev("() => { const n = new Date(); return [n.getHours() * 60 + n.getMinutes()]; }")[0]
    tv = ev("() => document.getElementById('rescheduleTimeInput').value"); tm = int(tv[:2]) * 60 + int(tv[3:])
    check("a minute-mode plan's Reschedule has a time box, defaulting to the time now (today)", pg.is_visible("#rescheduleTimeInput") and abs(tm - now) <= 2, (tv, now))
    pg.fill("#rescheduleDateInput", "2026-10-05"); pg.dispatch_event("#rescheduleDateInput", "change")
    check("...another date defaults to that day's own opening (08:00)", ev("() => document.getElementById('rescheduleTimeInput').value") == "08:00", ev("() => document.getElementById('rescheduleTimeInput').value"))
    pg.fill("#rescheduleTimeInput", "13:30"); pg.dispatch_event("#rescheduleTimeInput", "change")
    pg.fill("#rescheduleDateInput", "2026-10-06"); pg.dispatch_event("#rescheduleDateInput", "change")
    check("a time you typed is kept when the date changes", ev("() => document.getElementById('rescheduleTimeInput').value") == "13:30")
    check("the hint says how many would move ('The 1 unfinished task would move')", "would move" in pg.inner_text("#rescheduleHint") and "1 unfinished task" in pg.inner_text("#rescheduleHint"), pg.inner_text("#rescheduleHint"))
    check("...and the preview (a dry run) changed nothing: every task is exactly as it was", ev("() => JSON.stringify(tasks)") == before)
    pg.click("#rescheduleBtn"); pg.wait_for_timeout(300)
    moved = task("Late")
    check("Reschedule moves the late task to the status MOMENT (2026-10-06 13:30 or the next working minute)", moved[0] == "2026-10-06" and moved[1] == "13:30", moved)
    check("the toast names the moment", "06.10.2026 13:30" in toast(), toast())
    pg.click("#toastUndoBtn"); pg.wait_for_timeout(300)
    check("its Undo puts the times back too (09:00–11:00 on 01–02.09.)", task("Late") == ["2026-09-01", "09:00", "2026-09-02", "11:00"], task("Late"))
    plan([{"name": "Late", "s": "2026-09-01", "e": "2026-09-02"}], minute=False)
    ev("() => openRescheduleModal()"); pg.wait_for_timeout(200)
    check("a day-mode plan has no time box", not pg.is_visible("#rescheduleTimeInput"))
    close_all()

    # ---- Edit tasks: Move dates by, in working time
    plan([A, {"name": "C", "s": "2026-09-08", "e": "2026-09-08", "extra": {"startTime": "13:00", "endTime": "15:00"}}])
    ev("() => { setSelection([tasks.find(t => t.name === 'A').id]); openBulkModal(); }"); pg.wait_for_timeout(200)
    check("in a minute-mode plan 'Move dates by' takes working time (a text box with the units as its hint)", pg.get_attribute("#bulkVal-shift", "type") == "text" and "2h" in pg.inner_text("#bulkShiftHint"), pg.inner_text("#bulkShiftHint"))
    pg.check("#bulkOn-shift"); pg.fill("#bulkVal-shift", "nonsense"); pg.click("#bulkModalBg .modal-footer button.btn-primary"); pg.wait_for_timeout(150)
    check("nonsense is refused with a message", "working time" in toast() and pg.locator("#bulkModalBg.open").count() == 1, toast())
    pg.fill("#bulkVal-shift", "2h"); pg.click("#bulkModalBg .modal-footer button.btn-primary"); pg.wait_for_timeout(250)
    check("+2h moves A (09:00–11:30, 2.5h) to 11:00 and keeps its 2.5 working hours across lunch: it ends 14:30", task("A") == ["2026-09-07", "11:00", "2026-09-07", "14:30"], task("A"))
    ev("() => { setSelection([tasks.find(t => t.name === 'C').id]); openBulkModal(); }"); pg.wait_for_timeout(200)
    pg.check("#bulkOn-shift"); pg.fill("#bulkVal-shift", "-1d"); pg.click("#bulkModalBg .modal-footer button.btn-primary"); pg.wait_for_timeout(250)
    check("-1d moves a task one working day earlier, time kept and length kept (Tue 13:00–15:00 -> Mon 13:00–15:00)", task("C") == ["2026-09-07", "13:00", "2026-09-07", "15:00"], task("C"))
    plan([{"name": "A", "s": "2026-09-07", "e": "2026-09-07"}], minute=False)
    ev("() => { setSelection([tasks[0].id]); openBulkModal(); }"); pg.wait_for_timeout(200)
    check("a day-mode plan keeps the number box and 'working days'", pg.get_attribute("#bulkVal-shift", "type") == "number" and "working days" in pg.inner_text("#bulkShiftHint"))
    pg.check("#bulkOn-shift"); pg.fill("#bulkVal-shift", "2"); pg.click("#bulkModalBg .modal-footer button.btn-primary"); pg.wait_for_timeout(250)
    check("...and moves by working days as before (Mon -> Wed)", ev("() => tasks[0].startDate") == "2026-09-09")

    # ---- import preview
    plan([A])
    ev("() => openForeignImport('Task Name,Start,Finish\\nA,07.09.2026 09:00,07.09.2026 11:30\\nB,08.09.2026 13:00,08.09.2026 15:00', 'plan.csv')"); pg.wait_for_timeout(250)
    txt = pg.inner_text("#foreignImportSummary")
    check("the CSV preview shows the times ('07.09.2026 09:00 – 07.09.2026 11:30')", "07.09.2026 09:00" in txt and "11:30" in txt, txt)
    close_all()

    # ================================================================ 2. the Scheduling precision dialog
    plan([A])
    pg.click("#scheduleMenuBtn"); pg.click("#planPrecisionItem"); pg.wait_for_selector("#precisionModalBg.open"); pg.wait_for_timeout(200)
    check("it needs no scrolling at 1400x860 with the per-weekday section folded away", ev("() => { const b = document.querySelector('#precisionModalBg .modal-body'); return b.scrollHeight <= b.clientHeight + 1; }"))
    check("...and that section is a folded 'Different hours on some days' with no overrides yet", not ev("() => document.getElementById('whDayDetails').open") and "Different hours" in pg.inner_text("#whDayDetails summary"))
    w = ev("() => [...document.querySelectorAll('#precisionModalBg input[type=time]')].map(e => Math.round(e.getBoundingClientRect().width))")
    check("every time box is wide enough for '08:00 AM' and its clock icon (>= 120px)", all(x >= 120 for x in w[:5]), w)
    check("the dialog is the wide (620px) size", ev("() => Math.round(document.querySelector('#precisionModalBg .modal').getBoundingClientRect().width)") == 620)
    check("the breaks list and the day rows write a break the same way ('12:00-13:00')", "12:00-13:00" in pg.inner_text("#whBreakList") and ev("() => breaksText([{start:'12:00',end:'13:00'}])") == "12:00-13:00")
    pg.click("#whDayDetails summary"); pg.wait_for_timeout(100)
    rows = pg.locator("#whDayRows .whd-row")
    h0 = rows.first.bounding_box()["height"]
    check("each weekday row is ONE line (start, end and breaks side by side)", h0 < 44, h0)
    pg.locator("#whDayRows .whd-row[data-wd='5'] .whd-on").check()
    check("ticking a day updates the summary count ('1 day differs')", "1 day differs" in pg.inner_text("#whDayDetails summary"), pg.inner_text("#whDayDetails summary"))
    pg.keyboard.press("Escape"); pg.click("#confirmModalBg button:has-text('Discard')"); pg.wait_for_timeout(150)
    ev("() => { project.workHoursByDay = { 5: { start: '08:00', end: '13:00', breaks: [] } }; normalizeData(); save(); }")
    pg.click("#scheduleMenuBtn"); pg.click("#planPrecisionItem"); pg.wait_for_selector("#precisionModalBg.open"); pg.wait_for_timeout(200)
    check("with a different day saved the section opens by itself and says so", ev("() => document.getElementById('whDayDetails').open") and "1 day differs" in pg.inner_text("#whDayDetails summary"))
    pg.keyboard.press("Escape")

    # ================================================================ 3. predecessor rows in the task dialog
    plan([A, B])
    ev("() => openTaskModal(tasks.find(t => t.name === 'B').id)"); pg.wait_for_timeout(250)
    check("the predecessor rows have a heading (Task #, Task name, Type, Lead / lag)", all(x in pg.inner_text("#predecessorRows .pred-head") for x in ("TASK #", "TASK NAME", "TYPE", "LEAD / LAG")) or all(x.lower() in pg.inner_text("#predecessorRows .pred-head").lower() for x in ("Task #", "Task name", "Type", "Lead / lag")))
    check("...each row shows the predecessor's NAME beside its number (A)", pg.inner_text("#predecessorRows .pred-row .pred-name") == "A")
    check("...and its boxes have accessible names", pg.get_attribute("#predecessorRows .pred-row select", "aria-label") == "Link type" and pg.get_attribute("#predecessorRows .pred-id-input", "aria-label") == "Predecessor task number")
    pg.fill("#predecessorRows .pred-id-input", "99"); pg.dispatch_event("#predecessorRows .pred-id-input", "change"); pg.wait_for_timeout(100)
    check("a number that isn't a task is refused, and the row asks for one", "type the number" in pg.inner_text("#predecessorRows .pred-row .pred-name"), pg.inner_text("#predecessorRows .pred-row .pred-name"))
    pg.fill("#predecessorRows .pred-id-input", "1"); pg.dispatch_event("#predecessorRows .pred-id-input", "change"); pg.wait_for_timeout(100)
    check("typing a valid number shows that task's name", pg.inner_text("#predecessorRows .pred-row .pred-name") == "A")
    lag_box = pg.locator("#predecessorRows .pred-lag-input").bounding_box()
    check("the lag box is wide enough for '+90m' (>= 80px)", lag_box["width"] >= 80, lag_box)
    pg.keyboard.press("Escape"); close_all()

    # ================================================================ 4. accessibility
    info = ev("""() => [...document.querySelectorAll('.modal-bg')].map(bg => { const m = bg.querySelector('.modal'), lb = m.getAttribute('aria-labelledby'), h = lb && document.getElementById(lb);
        return { id: bg.id, role: m.getAttribute('role'), modal: m.getAttribute('aria-modal'), name: (h ? h.textContent.trim() : m.getAttribute('aria-label')) || '',
          unnamed: [...bg.querySelectorAll('.modal-header > button')].filter(x => !x.getAttribute('aria-label') && !x.textContent.replace(/[×\\s]/g, '')).length }; })""")
    check("every one of the %d dialogs is a modal dialog (role dialog / alertdialog, aria-modal)" % len(info), all(i["role"] in ("dialog", "alertdialog") and i["modal"] == "true" for i in info), [i["id"] for i in info if i["role"] not in ("dialog", "alertdialog") or i["modal"] != "true"])
    check("...each has a name (its title)", all(i["name"] for i in info), [i["id"] for i in info if not i["name"]])
    check("...the confirmation is an alertdialog", [i["role"] for i in info if i["id"] == "confirmModalBg"] == ["alertdialog"])
    check("...and no header × is left without a name", all(i["unnamed"] == 0 for i in info), [i["id"] for i in info if i["unnamed"]])
    check("the toast is a polite live region", ev("() => [document.getElementById('toast').getAttribute('role'), document.getElementById('toast').getAttribute('aria-live')]") == ["status", "polite"])
    plan([A, B])
    ev("() => openTaskModal(tasks.find(t => t.name === 'B').id)"); pg.wait_for_timeout(300)
    inside = lambda: ev("() => !!document.activeElement.closest('#taskModalBg')")
    ev("() => { const f = [...document.querySelectorAll('#taskModalBg .modal-footer button')]; f[f.length - 1].focus(); }")
    pg.keyboard.press("Tab")
    check("Tab from the last control goes to the first inside the dialog (not out to the page)", inside() and ev("() => document.activeElement.closest('.modal-header') !== null || document.activeElement.getAttribute('aria-label') === 'Close'"), ev("() => document.activeElement.outerHTML.slice(0, 80)"))
    pg.keyboard.press("Shift+Tab")
    check("Shift+Tab from the first goes to the last (Save)", inside() and ev("() => document.activeElement.textContent.trim()") == "Save", ev("() => document.activeElement.textContent"))
    ev("() => document.activeElement.blur()")
    pg.keyboard.press("Tab")
    check("with focus on the page behind, Tab pulls it into the dialog", inside())
    ev("() => openConfirmModal({ title: 'Sure?', body: 'x', confirmLabel: 'Yes', cancelLabel: 'No', action() {} })"); pg.wait_for_timeout(200)
    ev("() => { const f = [...document.querySelectorAll('#confirmModalBg button:not([disabled])')]; f[f.length - 1].focus(); }")
    pg.keyboard.press("Tab")
    check("a confirmation above the task dialog keeps Tab inside itself", ev("() => !!document.activeElement.closest('#confirmModalBg')"))
    close_all()
    ev("() => document.querySelector('.icon-btn[title=Help]').focus()")
    pg.click(".icon-btn[title=Help]"); pg.wait_for_selector("#helpModalBg.open"); pg.wait_for_timeout(150)
    pg.keyboard.press("Escape"); pg.wait_for_timeout(200)
    check("closing a dialog puts focus back on the button that opened it", ev("() => document.activeElement.getAttribute('title')") == "Help", ev("() => document.activeElement.outerHTML.slice(0, 60)"))

    # ================================================================ 5. polish
    plan([A, B])
    ev("() => { pendingImportData = { project: { name: 'Website relaunch', timeUnit: 'minute' }, tasks: [{ name: 'a' }, { name: 'b' }, { name: 's', spacer: true }] }; pendingImportName = 'relaunch.json'; openImportModal(); }"); pg.wait_for_timeout(150)
    t = pg.inner_text("#importInfo")
    check("the JSON import dialog says what is in the file: name, task count, plan, mode", "relaunch.json" in t and "2 tasks" in t and "Website relaunch" in t and "Hours" in t, t)
    close_all()
    ev("() => openBackupsModal()"); pg.wait_for_timeout(150)
    t = pg.inner_text("#backupsModalBg .modal-body")
    check("Local Backups explains itself and each row says how many tasks it holds", "once a day" in t and re.search(r"\d+ tasks?", t), t)
    close_all()
    ev("() => openBaselineModal()"); pg.wait_for_timeout(200)
    t = pg.inner_text("#baselineModalBg .modal-body").upper()
    check("Baselines: 'Save a baseline' says it applies on Set baseline, the comparison says it applies right away", "SAVE A BASELINE" in t and "APPLIES WHEN YOU CLICK SET BASELINE" in t and "APPLIES RIGHT AWAY" in t, t[:300])
    check("...and in a minute-mode plan the footer notes the time of day is kept", "time of day" in pg.inner_text("#baselineFootNote"))
    close_all()
    ev("() => { fileSyncStatus = 'unlinked'; openFileSyncModal(); }"); pg.wait_for_timeout(150)
    check("the folder dialog is a short line and three points, not one paragraph", pg.locator("#fileSyncModalText .fs-points li").count() == 3, pg.inner_text("#fileSyncModalText"))
    close_all()
    check("'mid' dialogs share one width (560px)", ev("() => { document.getElementById('rescheduleModalBg').classList.add('open'); const w = Math.round(document.querySelector('#rescheduleModalBg .modal').getBoundingClientRect().width); document.getElementById('rescheduleModalBg').classList.remove('open'); return w; }") == 560)

    check("no console errors or page errors across the whole run", not errors, errors[:5])
    n_ok, n_all = sum(results), len(results)
    print(f"\n{n_ok}/{n_all} checks passed")
    b.close()
    raise SystemExit(0 if n_ok == n_all else 1)
