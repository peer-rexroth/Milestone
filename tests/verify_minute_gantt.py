# -*- coding: utf-8 -*-
"""Minute-mode scheduling precision on the CHART: the Gantt's "Hours" scale (only offered in a minute-mode plan), where a bar
sits at its real start/finish instant on a 24-hour day (480px a day, 20px an hour) and a drag/resize snaps to 15 minutes —
an Auto task lands on working time (start on a working minute, finish where its working duration ends, so a move across the
lunch break comes out after it), a Manual task goes exactly where it is dropped, and successors follow at minute precision.
Also: the hour ticks and the shaded hours off (nights, the lunch break, days off), the day range's one-day padding, group
rollups that carry a time, the coarser scales staying day-based, a day-mode plan never seeing the scale (and a saved 'hour'
preference falling back to the standard one there), and the Excel Gantt sheet's one-column-per-working-hour layout (falling
back to weeks past 400 columns, and staying weekly for a day-mode plan)."""
import os, re, tempfile
from playwright.sync_api import sync_playwright
import openpyxl
URL = os.environ.get("MILESTONE_URL", "http://127.0.0.1:8937/milestone.html")
errors, results = [], []
def check(name, cond, detail=""):
    results.append(bool(cond)); print(("PASS  " if cond else "FAIL  ") + name + (f"   [{str(detail)[:400]}]" if not cond and detail else ""))
SEED = re.search(r'SEED = """(.*?)"""', open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'verify_clone.py')).read(), re.S).group(1)
TMP = tempfile.mkdtemp()

with sync_playwright() as p:
    b = p.chromium.launch(headless=True)
    ctx = b.new_context(viewport={"width": 1500, "height": 900}, accept_downloads=True); ctx.add_init_script("delete window.showOpenFilePicker; delete window.showSaveFilePicker; delete window.showDirectoryPicker")
    pg = ctx.new_page(); pg.on("pageerror", lambda e: errors.append(str(e))); pg.on("console", lambda m: errors.append(m.text) if m.type in ("error", "warning") else None)
    pg.goto(URL); pg.wait_for_selector("#addTaskBtn"); pg.evaluate("() => localStorage.clear()"); pg.reload(); pg.wait_for_selector("#addTaskBtn")
    ev = pg.evaluate
    seed = lambda specs: ev(SEED, specs)
    tid = lambda n: ev("n => tasks.find(t => t.name === n).id", n)
    task = lambda n: ev("n => { const t = tasks.find(x => x.name === n); return [t.startDate, t.startTime, t.endDate, t.endTime, t.constraintType]; }", n)
    day0 = ev("() => dayNumber('2026-09-07')")   # Monday

    def minute_plan(specs):
        ev("() => { project.timeUnit = 'minute'; }")   # (before the seed: saving a day-mode plan drops every time of day)
        seed(specs)
        ev("() => { project.timeUnit = 'minute'; normalizeData(); save(); render(); }")
    def gantt(zoom):
        ev("z => { setView('gantt'); setZoom(z); }", zoom); pg.wait_for_timeout(200)
    def drag(name, dx, grab="mid"):
        """Drag a bar by dx pixels (mid = its body; left/right = the resize handles, whose mousedown is dispatched since a hover link-handle covers a narrow bar's right edge)."""
        i = tid(name)
        if grab == "mid":
            box = pg.locator(f".gantt-bar[data-id='{i}']").bounding_box()
            x, y = box["x"] + min(20, box["width"] / 2), box["y"] + 10
            pg.mouse.move(x, y); pg.mouse.down(); pg.mouse.move(x + dx / 2, y, steps=4); pg.mouse.move(x + dx, y, steps=4); pg.mouse.up()
        else:
            h = pg.locator(f".gantt-bar[data-id='{i}'] .resize-handle.{grab}"); box = h.bounding_box(); x, y = box["x"] + 3, box["y"] + 8
            h.dispatch_event("mousedown", {"button": 0, "clientX": x, "clientY": y, "bubbles": True})
            pg.mouse.move(x + dx / 2, y, steps=4); pg.mouse.move(x + dx, y, steps=4); pg.mouse.up()
        pg.wait_for_timeout(200)

    # ---------------------------------------------------------------- the scale exists only in a minute-mode plan
    seed([{"name": "A", "s": "2026-09-07", "e": "2026-09-07"}])
    tabs = lambda: ev("() => [...document.querySelectorAll('#zoomTabs .view-tab')].map(b => b.textContent)")
    check("a day-mode plan offers Week / Month / Year only", tabs() == ["Week", "Month", "Year"], tabs())
    ev("() => { project.timeUnit = 'minute'; normalizeData(); save(); render(); }")
    check("switching the plan to Hours & minutes adds the Hours scale, first", tabs() == ["Hours", "Week", "Month", "Year"], tabs())

    # ---------------------------------------------------------------- bar geometry at 20px an hour
    minute_plan([{"name": "A", "s": "2026-09-07", "e": "2026-09-07", "extra": {"startTime": "09:00", "endTime": "11:30"}},
                 {"name": "B", "s": "2026-09-08", "e": "2026-09-08"},
                 {"name": "M", "s": "2026-09-08", "e": "2026-09-08", "extra": {"milestone": True, "startTime": "14:00", "endTime": "14:00"}}])
    gantt("hour")
    check("Hours is the active scale", ev("() => currentZoom().id") == "hour" and pg.locator("#zoomTabs .view-tab.active").inner_text() == "Hours")
    rng = ev("() => dateRange")
    check("the range is one day before the first task to two after the last (a day is 480px: no week of padding)", rng["start"] == day0 - 1 and rng["end"] == day0 + 1 + 2, rng)
    g = ev("() => barGeom")
    a, bb, m = g[tid("A")], g[tid("B")], g[tid("M")]
    check("A (09:00–11:30) starts 9 hours into its day and is 150 minutes = 50px wide", a["left"] == (1 + 9 / 24) * 480 and a["width"] == 50, a)
    check("B, with no time of its own, spans the working day 08:00–17:00 = 9 hours = 180px", bb["left"] == (2 + 8 / 24) * 480 and bb["width"] == 180, bb)
    check("the milestone (14:00) sits at its own instant: its centre is 14 hours into the day", abs((m["left"] + 8) - (2 + 14 / 24) * 480) < 0.01, m)
    check("the bar's tooltip names when it starts and finishes", "07.09.2026 09:00 – 11:30" in pg.get_attribute(f".gantt-bar[data-id='{tid('A')}']", "title"), pg.get_attribute(f".gantt-bar[data-id='{tid('A')}']", "title"))

    # ---------------------------------------------------------------- the header: the day over its 24 hours
    days = rng["end"] - rng["start"]
    check("one day tick per day and 24 hourly ticks under each", pg.locator("#ganttHeader .zoom-hour-day").count() == days and pg.locator("#ganttHeader .zoom-hour-h").count() == 24 * days)
    check("the day tick reads weekday and date, the hour ticks 00…23", "Mon 07.09.2026" in pg.locator("#ganttHeader .zoom-hour-day").nth(1).inner_text() and pg.locator("#ganttHeader .zoom-hour-h").nth(24 + 9).inner_text() == "09")
    check("the two header rows share the standard header height (the rows still line up with the list)", ev("() => document.getElementById('ganttHeader').getBoundingClientRect().height") == 42)
    check("a day off's tick is dimmed (Sunday before the plan)", "dayoff" in (pg.locator("#ganttHeader .zoom-hour-day").nth(0).get_attribute("class") or ""))

    # ---------------------------------------------------------------- the hours off are shaded
    nonwork = ev("() => [...document.querySelectorAll('#ganttRows .gantt-nonwork')].map(e => [parseFloat(e.style.left), parseFloat(e.style.width)])")
    def has(day_off, from_h, to_h): return any(abs(l - (day_off + from_h / 24) * 480) < 0.01 and abs(w - (to_h - from_h) / 24 * 480) < 0.01 for l, w in nonwork)
    check("Monday 00:00–08:00, the 12:00–13:00 lunch break and 17:00–24:00 are shaded", has(1, 0, 8) and has(1, 12, 13) and has(1, 17, 24), nonwork[:8])
    check("Sunday (a day off) is shaded whole", has(0, 0, 24))
    tl = ev("() => parseFloat(document.querySelector('.gantt-today-line').style.left)")
    exp = ev("() => { const n = new Date(); return (dayNumber(todayStr()) - dateRange.start + (n.getHours() * 60 + n.getMinutes()) / 1440) * 480; }")
    check("the today line is at the current time of day, not the start of today", abs(tl - exp) < 8, (tl, exp))

    # ---------------------------------------------------------------- the coarser scales stay day-based in a minute-mode plan
    gantt("week")
    ga = ev("() => barGeom")[tid("A")]
    check("at Week zoom A is still a whole-day cell (its day, 14px less the 2px gap), whatever its times", ga["width"] == 14 - 2 and ga["left"] % 14 == 0, ga)
    check("...with the usual three days of padding on the range, not the Hours scale's one", ev("() => dateRange.start") == day0 - 3, ev("() => dateRange"))

    # ---------------------------------------------------------------- dragging on the Hours scale: 15-minute steps, working time, cascade
    minute_plan([{"name": "A", "s": "2026-09-07", "e": "2026-09-07", "extra": {"startTime": "09:00", "endTime": "11:30"}},
                 {"name": "B", "s": "2026-09-07", "e": "2026-09-07", "preds": [["A", "FS", 0]], "extra": {"startTime": "11:30", "endTime": "12:00"}}])
    gantt("hour")
    drag("A", 5)
    check("5px = 15 minutes: A moves to 09:15–11:45", task("A")[:4] == ["2026-09-07", "09:15", "2026-09-07", "11:45"], task("A"))
    check("...and B, its FS successor, followed at minute precision (starts at A's new finish)", task("B")[1] == "11:45", task("B"))
    drag("A", 60)   # +3 hours: 09:15 -> 12:15, inside the lunch break -> 13:00
    check("+3h lands the start in the lunch break, so an Auto task starts after it (13:00) and keeps its 2.5h working length (finish 15:30)", task("A")[:4] == ["2026-09-07", "13:00", "2026-09-07", "15:30"], task("A"))
    check("...B follows again", task("B")[:2] == ["2026-09-07", "15:30"], task("B"))
    ev("() => { const t = tasks.find(x => x.name === 'A'); t.startTime = '09:00'; t.endTime = '11:30'; save(); render(); }")
    drag("A", 36)   # 36px = 108 minutes -> snaps to 105
    check("a drag snaps to 15-minute steps (36px = 108 min -> 105 min: 10:45)", task("A")[1] == "10:45", task("A"))
    ev("() => { const t = tasks.find(x => x.name === 'A'); t.startTime = '09:00'; t.endTime = '11:30'; save(); render(); }")
    drag("A", -60, "left")
    check("resizing the left edge 3h earlier moves only the start (to 08:00 — the day's open, an Auto task does not start before it)", task("A")[1] == "08:00" and task("A")[3] == "11:30", task("A"))
    drag("A", 40, "right")
    check("resizing the right edge 2h later moves only the finish: 11:30 + 2h of the chart = 13:30 (its length is then 4h30 of working time, the lunch break not counted)", task("A")[1] == "08:00" and task("A")[3] == "13:30", task("A"))
    minute_plan([{"name": "A", "s": "2026-09-07", "e": "2026-09-07", "extra": {"startTime": "09:00", "endTime": "11:00"}},
                 {"name": "B", "s": "2026-09-07", "e": "2026-09-07", "preds": [["A", "FS", 0]], "extra": {"startTime": "11:00", "endTime": "12:00"}}])
    gantt("hour")
    drag("B", 20)
    bt = ev("() => { const t = tasks.find(x => x.name === 'B'); return [t.startTime, t.constraintType, t.constraintDate, t.constraintTime]; }")
    check("dragging a linked Auto task an hour later (11:00 -> 12:00, the lunch break, so 13:00) pins it like any typed date: Start No Earlier Than that moment, time included", bt == ["13:00", "SNET", "2026-09-07", "13:00"], bt)

    # a Manual task goes exactly where it is dropped
    minute_plan([{"name": "Mn", "s": "2026-09-07", "e": "2026-09-07", "extra": {"taskMode": "manual", "startTime": "09:00", "endTime": "10:00"}}])
    gantt("hour")
    drag("Mn", 60)
    check("a Manual task is not held to working time: 09:00 + 3h = 12:00 (in the lunch break) stays there", task("Mn")[:4] == ["2026-09-07", "12:00", "2026-09-07", "13:00"], task("Mn"))

    # a milestone
    minute_plan([{"name": "Ms", "s": "2026-09-07", "e": "2026-09-07", "extra": {"milestone": True, "startTime": "10:00", "endTime": "10:00"}}])
    gantt("hour")
    box = pg.locator(".gantt-milestone").bounding_box()
    pg.mouse.move(box["x"] + 8, box["y"] + 8); pg.mouse.down(); pg.mouse.move(box["x"] + 30, box["y"] + 8, steps=4); pg.mouse.move(box["x"] + 48, box["y"] + 8, steps=4); pg.mouse.up(); pg.wait_for_timeout(200)
    ms = task("Ms")
    check("a milestone drags to a time of its own (10:00 + 40px = 2h -> 12:00 is lunch, so 13:00) and keeps a single instant", ms[1] == "13:00" and ms[3] == "13:00" and ms[0] == ms[2], ms)
    check("no drag tip is left behind", pg.locator("#ganttDragTip").count() == 0)

    # a click that doesn't move still opens the task dialog
    minute_plan([{"name": "A", "s": "2026-09-07", "e": "2026-09-07", "extra": {"startTime": "09:00", "endTime": "11:30"}}])
    gantt("hour")
    box = pg.locator(f".gantt-bar[data-id='{tid('A')}']").bounding_box(); pg.mouse.click(box["x"] + 20, box["y"] + 10); pg.wait_for_timeout(200)
    check("a click without a drag opens the task dialog (unchanged)", pg.locator("#taskModalBg.open").count() == 1)
    pg.keyboard.press("Escape")

    # ---------------------------------------------------------------- groups roll up their children's times
    minute_plan([{"name": "G", "s": "2026-09-07", "e": "2026-09-07"},
                 {"name": "C1", "s": "2026-09-07", "e": "2026-09-07", "parent": "G", "extra": {"startTime": "13:00", "endTime": "15:00"}},
                 {"name": "C2", "s": "2026-09-07", "e": "2026-09-07", "parent": "G", "extra": {"startTime": "09:00", "endTime": "10:30"}}])
    eff = ev("() => { const e = effectiveDates(tasks.find(t => t.name === 'G').id); return [e.start, e.startTime, e.end, e.endTime]; }")
    check("a group's rollup carries the earliest child start (09:00) and the latest child finish (15:00)", eff == ["2026-09-07", "09:00", "2026-09-07", "15:00"], eff)
    gantt("hour")
    ga = ev("() => barGeom")[tid("G")]
    check("...and its bar on the Hours scale spans exactly that: 09:00–15:00 = 6h = 120px", ga["left"] == (1 + 9 / 24) * 480 and ga["width"] == 120, ga)
    ev("() => { project.timeUnit = 'day'; normalizeData(); save(); render(); }")
    check("in day mode effectiveDates carries no time fields at all (byte-identical to before)", ev("() => { const e = effectiveDates(tasks.find(t => t.name === 'G').id); return 'startTime' in e || 'endTime' in e; }") is False)

    # ---------------------------------------------------------------- baseline bars in the Hours scale
    minute_plan([{"name": "A", "s": "2026-09-07", "e": "2026-09-07", "extra": {"startTime": "09:00", "endTime": "11:00"}}])
    ev("() => { applyBaselineChange(0, 'all', false); const t = tasks.find(x => x.name === 'A'); t.startTime = '10:00'; t.endTime = '12:00'; save(); showBaseline = true; render(); }")
    gantt("hour")
    bl = ev("() => { const e = document.querySelector('.gantt-base'); return e ? [parseFloat(e.style.left), parseFloat(e.style.width)] : null; }")
    check("the baseline bar keeps the baseline's own 09:00–11:00 (not the moved task's 10:00–12:00)", bl == [(1 + 9 / 24) * 480, 40], bl)

    # ---------------------------------------------------------------- a day-mode plan never draws the scale, and a saved preference falls back
    ev("() => { project.timeUnit = 'day'; normalizeData(); save(); render(); }")
    check("with the saved scale still 'hour', a day-mode plan draws the standard one (Year) and offers no Hours tab", ev("() => currentZoom().id") == "year" and tabs() == ["Week", "Month", "Year"], (ev("() => currentZoom().id"), tabs()))
    ev("() => { project.timeUnit = 'minute'; normalizeData(); save(); render(); }")
    check("...and back in a minute-mode plan the Hours scale is still the chosen one", ev("() => currentZoom().id") == "hour")

    # ---------------------------------------------------------------- Excel Gantt sheet
    def export_gantt(name):
        pg.click("#dataMenuBtn"); pg.click("#excelExportItem"); pg.wait_for_selector("#excelModalBg.open")
        with pg.expect_download() as d: pg.click("#excelExportBtn")
        path = os.path.join(TMP, name); d.value.save_as(path)
        return openpyxl.load_workbook(path)["Gantt"]
    minute_plan([{"name": "A", "s": "2026-09-07", "e": "2026-09-07", "extra": {"startTime": "09:00", "endTime": "11:30"}},
                 {"name": "B", "s": "2026-09-08", "e": "2026-09-08", "extra": {"startTime": "13:00", "endTime": "16:00"}},
                 {"name": "M", "s": "2026-09-08", "e": "2026-09-08", "extra": {"milestone": True, "startTime": "10:00", "endTime": "10:00"}}])
    ws = export_gantt("hourly.xlsx")
    hdr = {c.column: c.value for c in ws[4] if c.value is not None}
    first = min(c for c, v in hdr.items() if isinstance(v, str) and re.fullmatch(r"\d\d", v))
    hours = [hdr[c] for c in sorted(hdr) if c >= first]
    check("the Gantt sheet has one column per WORKING hour: 08–11 and 13–16 for each of the two days (the lunch hour is left out)", hours == ["08", "09", "10", "11", "13", "14", "15", "16"] * 2, hours)
    check("the day is written over its hours ('Mon 07.09.2026')", "Mon 07.09.2026" in (ws.cell(row=3, column=first).value or "") and "Tue 08.09.2026" in (ws.cell(row=3, column=first + 8).value or ""), (ws.cell(row=3, column=first).value, ws.cell(row=3, column=first + 8).value))
    name_col = [c.column for c in ws[3] if c.value == "Task Name"][0]
    row_a = next(r for r in range(5, 12) if ws.cell(row=r, column=name_col).value == "A")
    filled = lambda r: [ws.cell(row=r, column=first + i).fill.fill_type == "solid" for i in range(16)]
    check("A (09:00–11:30) fills the 09, 10 and 11 columns of Monday — and nothing else", filled(row_a) == [False, True, True, True] + [False] * 12, filled(row_a))
    row_b = next(r for r in range(5, 12) if ws.cell(row=r, column=name_col).value == "B"); row_m = next(r for r in range(5, 12) if ws.cell(row=r, column=name_col).value == "◆ M")
    check("B (Tue 13:00–16:00) fills Tuesday's 13, 14 and 15 columns", filled(row_b) == [False] * 8 + [False, False, False, False, True, True, True, False], filled(row_b))
    check("the milestone (Tue 10:00) is a diamond in Tuesday's 10 column", ws.cell(row=row_m, column=first + 8 + 2).value == "◆" and sum(1 for i in range(16) if ws.cell(row=row_m, column=first + i).value == "◆") == 1)
    check("the sub-title says what a column is", "one column per hour" in (ws.cell(row=2, column=1).value or ""), ws.cell(row=2, column=1).value)

    # past 400 columns it falls back to the calendar weeks
    seed([{"name": "Long", "s": "2026-01-05", "e": "2027-06-30"}])
    ev("() => { project.timeUnit = 'minute'; normalizeData(); save(); render(); }")
    ws = export_gantt("long.xlsx")
    lab = [c.value for c in ws[4] if isinstance(c.value, str) and c.value.startswith("CW")]
    check("a plan too long for hourly columns falls back to calendar weeks (day-based, as before)", len(lab) > 30 and not any(re.fullmatch(r"\d\d", str(c.value)) for c in ws[4] if c.value), lab[:3])
    # ... and a day-mode plan is weekly
    minute_plan([{"name": "A", "s": "2026-09-07", "e": "2026-09-09"}])
    ev("() => { delete project.timeUnit; delete project.workHours; normalizeData(); save(); render(); }")
    ws = export_gantt("day.xlsx")
    check("a day-mode plan's Gantt sheet stays one column per calendar week", any(isinstance(c.value, str) and c.value.startswith("CW") for c in ws[4]))

    check("no console errors or page errors across the whole run", not errors, errors[:5])
    n_ok, n_all = sum(results), len(results)
    print(f"\n{n_ok}/{n_all} checks passed")
    b.close()
    raise SystemExit(0 if n_ok == n_all else 1)
